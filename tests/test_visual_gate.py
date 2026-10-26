import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from PIL import Image
import pytest

from app.assets.resolve import AssetResolver, AssetSettings
from app.core.deck_schema import Deck
from app.jobs.runner import JobRunner
from app.jobs.store import JobStore
from app.main import create_app
from app.qa.gate import build_checked, QualityGateError
from app.qa.lint import lint_pptx
from app.core.deck_schema import StorySlide
from app.render.builder import build_deck


@pytest.fixture
def visual_deck(monkeypatch, tmp_path):
    picture = io.BytesIO()
    Image.new("RGB", (120, 80), "#bb9988").save(picture, format="PNG")
    monkeypatch.setattr(
        "app.assets.resolve.download", lambda *a, **k: (picture.getvalue(), "image/png")
    )
    resolver = AssetResolver(AssetSettings(cache_dir=tmp_path / "assets"))
    monkeypatch.setattr("app.qa.gate.prepare_assets", resolver.resolve)
    return resolver.resolve(
        Deck.model_validate_json(Path("tests/fixtures/decks/blush_story.json").read_text()),
        "Blush beauty",
    )


def fake_render(path, folder, **kwargs):
    from pptx import Presentation

    paths = []
    for i in range(len(Presentation(path).slides)):
        png = folder / f"slide-{i + 1:02d}.png"
        Image.new("RGB", (160, 90), "white").save(png)
        paths.append(png)
    return paths


@pytest.mark.parametrize(
    "archetype,count",
    [
        ("split_image", 3),
        ("three_cards", 3),
        ("four_cards", 4),
        ("grid_2x2", 4),
        ("comparison", 2),
        ("process_steps", 3),
        ("timeline_horizontal", 5),
        ("big_stat", 1),
        ("quote", 1),
    ],
)
def test_every_archetype_exports_readable_editable_regions(visual_deck, tmp_path, archetype, count):
    slide = StorySlide(
        archetype=archetype,
        title="A personal approach",
        items=[
            {
                "title": "One idea" if count == 1 else f"Focus {i + 1}",
                "description": "Share customer stories through a focused content program.",
            }
            for i in range(count)
        ],
    )
    deck = visual_deck.model_copy(
        update={"slides": [visual_deck.slides[0], slide, visual_deck.slides[-1]]}
    )
    report = lint_pptx(build_deck(deck, tmp_path / (archetype + ".pptx"), strict=True))
    assert report["errors"] == []
    assert report["minimum_font_pt"] >= 14


def test_gate_requires_actual_saved_lint_and_all_slide_scores(visual_deck, tmp_path):
    out = tmp_path / "approved.pptx"
    build_checked(
        visual_deck, out, render=fake_render, reviewer=lambda *a, **k: {"score": 9, "fixes": []}
    )
    report = json.loads((tmp_path / "approved.qa/report.json").read_text())
    assert out.is_file() and report["passed"]
    assert report["attempts"][0]["lint"]["minimum_font_pt"] == 14
    assert len(report["attempts"][0]["vision"]) == 9
    assert report["attempts"][0]["lint"]["errors"] == []


@pytest.mark.parametrize("archetype", ["quote", "big_stat"])
def test_statement_body_follows_heading_and_long_source_copy_still_exports(
    visual_deck, tmp_path, archetype
):
    from pptx import Presentation
    from pptx.util import Inches

    description = (
        "Foundation, lipstick, and skincare are blended exclusively for your skin tone, "
        "undertone, and individual style, with the supplied launch centered on a quiz, "
        "waitlist, and creator-led demos."
    )
    statement = StorySlide(
        archetype=archetype,
        title="The Shade Dilemma",
        items=[{"title": "Never Wrong Shade", "description": description}],
    )
    deck = visual_deck.model_copy(
        update={"slides": [visual_deck.slides[0], statement, visual_deck.slides[-1]]}
    )
    path = build_deck(deck, tmp_path / f"{archetype}-flow.pptx", strict=True)
    shapes = Presentation(path).slides[1].shapes
    heading = next(shape for shape in shapes if getattr(shape, "text", "") == "Never Wrong Shade")
    body = next(shape for shape in shapes if getattr(shape, "text", "") == description)
    gap = body.top - heading.top - heading.height
    assert Inches(0.15) <= gap <= Inches(0.35)
    assert body.top + body.height <= Inches(6.31)
    assert lint_pptx(path)["errors"] == []


def test_low_score_rerenders_and_maximum_three_attempts_never_delivers(visual_deck, tmp_path):
    out = tmp_path / "rejected.pptx"

    def reject(*args, **kwargs):
        return {"score": 5, "fixes": [{"action": "recrop_image", "detail": "Improve subject crop"}]}

    with pytest.raises(QualityGateError):
        build_checked(visual_deck, out, render=fake_render, reviewer=reject)
    report = json.loads((tmp_path / "rejected.qa/report.json").read_text())
    assert not out.exists() and not report["passed"]
    assert len(report["attempts"]) == 3


def test_debug_reads_known_job_only_and_rejects_path_traversal(tmp_path):
    store = JobStore(tmp_path)
    job = store.create("Blush test prompt")
    folder = store.output_dir / (job.id + ".qa")
    (folder / "1").mkdir(parents=True)
    (folder / "report.json").write_text(
        json.dumps({"passed": False, "attempts": [{"images": ["1/slide-01.png"]}]})
    )
    Image.new("RGB", (20, 20)).save(folder / "1/slide-01.png")
    with JobRunner(store, generate=lambda _: None) as runner:
        with TestClient(create_app(runner=runner)) as client:
            response = client.get("/debug/" + job.id)
            assert response.status_code == 200
            assert response.json()["attempts"][0]["images"] == [f"/debug/{job.id}/1/slide-01.png"]
            assert (
                client.get(f"/debug/{job.id}/1/slide-01.png").headers["content-type"] == "image/png"
            )
            assert client.get("/debug/unknown").status_code == 404
            assert client.get(f"/debug/{job.id}/4/slide-01.png").status_code == 404
            assert client.get(f"/debug/{job.id}/1/.env").status_code == 404


def test_vision_retries_only_transient_unavailability_and_sends_pixels(tmp_path, monkeypatch):
    import base64
    from app.llm.vision import GroqVisionReviewer
    from app.core.errors import LLMProviderError

    calls = []

    def complete(**kwargs):
        calls.append(kwargs)
        if len(calls) == 1:
            raise LLMProviderError("service unavailable", status_code=503)
        return json.dumps(
            {
                "observed_title": "Test",
                "composition": "A green slide",
                "score": 8,
                "hierarchy": 8,
                "whitespace_balance": 8,
                "alignment": 8,
                "legibility": 8,
                "visual_interest": 8,
                "on_brand_palette": 8,
                "fixes": [],
            }
        )

    class Client:
        pass

    client = Client()
    client.complete = complete
    monkeypatch.setattr("app.llm.vision.time.sleep", lambda _: None)
    png = tmp_path / "slide.png"
    Image.new("RGB", (160, 90), "green").save(png)
    grade = GroqVisionReviewer(client=client)(png, theme="climate", slide_index=0)
    assert grade.score == 8 and len(calls) == 2
    encoded = calls[-1]["messages"][1]["content"][1]["image_url"]["url"].split(",")[1]
    with Image.open(io.BytesIO(base64.b64decode(encoded))) as sent:
        assert sent.getpixel((500, 500))[1] > 100


def test_vision_daily_quota_is_not_a_stylistic_retry(tmp_path):
    from app.llm.vision import GroqVisionReviewer
    from app.core.errors import ProviderLimitError

    calls = []

    class Client:
        def complete(self, **kwargs):
            calls.append(kwargs)
            raise ProviderLimitError(diagnostics={"limit_type": "tokens_per_day"})

    png = tmp_path / "slide.png"
    Image.new("RGB", (160, 90), "green").save(png)
    with pytest.raises(ProviderLimitError):
        GroqVisionReviewer(client=Client())(png, theme="climate", slide_index=0)
    assert len(calls) == 1
