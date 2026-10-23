"""Render already-resolved assets only. This module never fetches URLs."""

from contextlib import contextmanager
from contextvars import ContextVar
import hashlib
from pathlib import Path

from PIL import Image
from pptx.util import Inches


_assets = ContextVar("slide_assets", default=())


@contextmanager
def asset_scope(assets):
    token = _assets.set(tuple(assets))
    try:
        yield
    finally:
        _assets.reset(token)


def find_asset(key):
    return next((asset for asset in _assets.get() if asset.key == key), None)


def add_asset(slide, key, x, y, w, h, *, crop=False):
    asset = find_asset(key)
    if asset is None:
        return False
    try:
        path = Path(asset.path)
        if hashlib.sha256(path.read_bytes()).hexdigest() != asset.sha256:
            return False
        with Image.open(path) as image:
            aspect = image.width / image.height
        target = w / h
        if not crop:
            width, height = (w, w / aspect) if aspect > target else (h * aspect, h)
            x, y = x + (w - width) / 2, y + (h - height) / 2
            picture = slide.shapes.add_picture(
                str(path), Inches(x), Inches(y), Inches(width), Inches(height)
            )
        else:
            picture = slide.shapes.add_picture(
                str(path), Inches(x), Inches(y), Inches(w), Inches(h)
            )
            if aspect > target:
                picture.crop_left = picture.crop_right = (1 - target / aspect) / 2
            else:
                picture.crop_top = picture.crop_bottom = (1 - aspect / target) / 2
        picture.name = f"Resolved {asset.role}: {key}"
        return True
    except Exception:
        return False
