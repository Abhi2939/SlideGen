"""Regression coverage for the owner-supplied Blush deck and source prompt."""

import json

import pytest
from pptx import Presentation

from app.config import Settings
from app.core.deck_schema import ClosingSlide
from app.pipeline.analyze import analyze_prompt
from app.pipeline.briefs import beauty_brand
from app.pipeline.models import AnalysisOutput, OutlineOutput
from app.pipeline.orchestrator import generate_deck
from app.pipeline.write_slides import _slide_text_content, _write_one
from app.render.builder import build_deck


PROMPT = (
    "Blush is a custom beauty brand offering foundation, lipstick, and skincare blended "
    "exclusively for your skin tone, undertone, and style, so you never get the wrong shade. "
    "Our launch will focus on a shade-matching quiz, waitlist, and creator-led demos, "
    "followed by quarterly growth through UGC, subscriptions, new product lines, pop-ups, "
    "and festive scaling."
)
FACT_TEXTS = (
    "Blush",
    "foundation",
    "lipstick",
    "skincare",
    "skin tone",
    "undertone",
    "style",
    "shade-matching quiz",
    "waitlist",
    "creator-led demos",
    "quarterly growth",
    "UGC",
    "subscriptions",
    "new product lines",
    "pop-ups",
    "festive scaling",
    "you never get the wrong shade",
    "blended exclusively",
)
SLIDES = (
    {"layout": "title", "title": "Blush", "subtitle": "Custom beauty"},
    {
        "layout": "two_column",
        "title": "Custom beauty, matched to you",
        "left": {
            "heading": "The collection",
            "bullets": [{"text": text} for text in ("Foundation", "Lipstick", "Skincare")],
        },
        "right": {
            "heading": "Personalization",
            "bullets": [
                {
                    "text": "Blended exclusively",
                    "detail": "for your skin tone, undertone, and style",
                },
            ],
        },
    },
    {
        "layout": "bullets",
        "title": "Launch focus",
        "bullets": [{"text": text[0].upper() + text[1:]} for text in FACT_TEXTS[7:10]],
    },
    {
        "layout": "flow_steps",
        "title": "Quarterly growth plans",
        "steps": [{"title": text[0].upper() + text[1:]} for text in FACT_TEXTS[11:16]],
    },
    {
        "layout": "closing",
        "title": "The Blush promise",
        "subtitle": "You never get the wrong shade",
    },
)


class BlushGenerator:
    """Typed fake results; the real analyzer/outline/writer/review still run."""

    def __init__(self):
        self.calls = []

    def generate(self, output_model, **kwargs):
        self.calls.append((output_model, kwargs))
        if kwargs["stage"] == "analyze":
            return output_model(
                deck_title="Offering Foundation",  # Reproduce the incorrect source-name guess.
                deck_type="product launch and growth plan",
                slide_count=len(SLIDES),
                required_facts=[
                    {"id": index, "text": text, "verbatim": index == 0}
                    for index, text in enumerate(FACT_TEXTS)
                ],
            )
        if kwargs["stage"] == "outline":
            brief = json.loads(kwargs["user_prompt"].split("ANALYZED BRIEF:\n")[1])
            assert brief["deck_title"] == "Blush"
            assert [fact["text"] for fact in brief["required_facts"]] == list(FACT_TEXTS)
            assert brief["slide_count"] == len(SLIDES)
            fact_ids = ([0], [1, 2, 3, 4, 5, 6, 17], [7, 8, 9], [10, 11, 12, 13, 14, 15], [16])
            return output_model(
                deck_title="Blush",
                slides=[
                    {
                        "index": index,
                        "title": slide["title"],
                        "layout": slide["layout"],
                        "purpose": "Present these supplied details visibly, as plans where applicable.",
                        "required_fact_ids": assigned,
                    }
                    for index, (slide, assigned) in enumerate(zip(SLIDES, fact_ids))
                ],
            )
        if kwargs["stage"] == "write":
            plan = json.loads(
                kwargs["user_prompt"].split("SLIDE PLAN:\n")[1].split("\n\nFACTS ASSIGNED")[0]
            )
            return output_model.model_validate(SLIDES[plan["index"]])
        if kwargs["stage"] == "review":
            return output_model(passed=True)
        raise AssertionError(kwargs["stage"])


def offline_deck():
    generator = BlushGenerator()
    settings = Settings("groq", None, "fake", None, None, None, None, 2)
    return generate_deck(PROMPT, settings=settings, generator=generator), generator


@pytest.mark.parametrize(
    "prompt, expected",
    [
        (PROMPT, "Blush"),
        (PROMPT.replace("Blush", "Rose & Co", 1), "Rose & Co"),
        ("Brand name blush, foundation, lipstick, skincare", "Blush"),
        ("Brand: Blush. Foundation, lipstick, skincare. Make a 6-slide deck.", "Blush"),
        ("Startup name blush, foundation, lipstick, skincare", "Blush"),
        ("My startup makes foundation, lipstick and skincare", None),
        ("A custom beauty brand offering foundation, lipstick and skincare", None),
        ("Our company is a beauty brand offering foundation and skincare", None),
        ("An electronics brand offering computers", None),
    ],
)
def test_brand_extraction_requires_explicit_name_or_named_subject(prompt, expected):
    assert beauty_brand(prompt) == expected


def test_exact_blush_prompt_keeps_all_source_details_in_visible_powerpoint(tmp_path):
    deck, generator = offline_deck()
    assert deck.title == "Blush"
    assert deck.theme == "blush"
    assert len(deck.slides) == len(SLIDES)
    assert any(kwargs["stage"] == "outline" for _, kwargs in generator.calls)
    assert "shade-matching quiz" in _slide_text_content(deck.slides[2]).casefold()
    assert "festive scaling" in _slide_text_content(deck.slides[3]).casefold()

    output = build_deck(deck, tmp_path / "blush.pptx", strict=True)
    presentation = Presentation(output)
    visible = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ).casefold()
    assert all(text.casefold() in visible for text in FACT_TEXTS)
    assert "offering foundation" not in visible
    assert "custom beauty  ," not in visible
    assert not any(slide.has_notes_slide for slide in presentation.slides)
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert 0 <= shape.left <= shape.left + shape.width <= presentation.slide_width
            assert 0 <= shape.top <= shape.top + shape.height <= presentation.slide_height


@pytest.mark.parametrize("count", [3, 6, 8])
def test_beauty_analysis_preserves_explicit_slide_count_and_supplied_plans(count):
    class Generator(BlushGenerator):
        def generate(self, output_model, **kwargs):
            result = super().generate(output_model, **kwargs)
            return result.model_copy(update={"slide_count": count})

    analysis = analyze_prompt(f"{PROMPT} Make a {count}-slide deck.", Generator(), "fake")
    assert analysis.deck_title == "Blush"
    assert analysis.slide_count == count
    assert [fact.text for fact in analysis.required_facts] == list(FACT_TEXTS)


def test_launch_plan_in_closing_notes_does_not_satisfy_visible_fact_requirement():
    launch_plan = PROMPT.split(". ", 1)[1]
    assert 140 < len(launch_plan) <= 200
    analysis = AnalysisOutput(
        deck_title="Blush",
        deck_type="launch",
        slide_count=3,
        required_facts=[{"id": 0, "text": launch_plan, "verbatim": True}],
    )
    outline = OutlineOutput(
        deck_title="Blush",
        slides=[
            {"index": 0, "title": "Blush", "purpose": "Introduce", "layout": "title"},
            {"index": 1, "title": "Products", "purpose": "Explain", "layout": "bullets"},
            {
                "index": 2,
                "title": "The Blush promise",
                "purpose": "Close",
                "layout": "closing",
                "required_fact_ids": [0],
            },
        ],
    )

    class NotesOnly:
        calls = 0

        def generate(self, output_model, **kwargs):
            self.calls += 1
            return ClosingSlide(
                layout="closing", title="The Blush promise", speaker_notes=launch_plan
            )

    generator = NotesOnly()
    with pytest.raises(ValueError, match="omitted verbatim fact IDs"):
        _write_one(
            PROMPT, analysis, outline, 2, generator, "fake", max_tokens=1000, semantic_attempts=2
        )
    assert generator.calls == 2
