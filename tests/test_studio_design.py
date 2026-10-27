from pathlib import Path

import pytest
from pptx import Presentation

from app.core.deck_schema import Deck
from app.pipeline.draft import planned_draft_model
from app.qa.draft import build_draft
from app.qa.lint import lint_pptx
from app.render.data_charts import extract_chart
from app.render.studio import Canvas, brief_groups
from app.render.theme import STUDIO, STUDIO_BEAUTY, contrast, draft_theme_for_topic, rgb
from test_draft_delivery import content
from test_editorial_pipeline import PROMPT


FIXTURE = Path("tests/fixtures/decks/blush_reference_draft.json")


def leaves(slide):
    for shape in slide.shapes:
        if shape.shape_type == 6:
            yield from leaves(shape)
        else:
            yield shape


def test_fuller_beauty_design_keeps_source_and_embeds_an_editable_source_count_chart(tmp_path):
    deck = Deck.model_validate_json(FIXTURE.read_text()).model_copy(
        update={"theme": "studio-beauty"}
    )
    path = build_draft(deck, tmp_path / "rich.pptx")
    prs = Presentation(path)
    assert [slide._element.cSld.get("name") for slide in prs.slides] == [
        "studio_cover",
        "studio_collection",
        "studio_personalization",
        "studio_launch",
        "studio_growth",
        "studio_closing",
    ]
    assert [len(labels) for _, labels in brief_groups(deck)] == [3, 3, 5]
    for data, slide in zip(deck.slides, prs.slides):
        visible = " ".join(s.text for s in leaves(slide) if s.has_text_frame)
        for point in getattr(data, "points", []):
            assert point in visible
    chart_shape = next(s for s in prs.slides[-1].shapes if s.has_chart)
    assert chart_shape.name == "Chart:brief-counts"
    assert tuple(chart_shape.chart.series[0].values) == (3.0, 3.0, 5.0)
    colors = {
        str(s.fill.fore_color.rgb)
        for s in leaves(prs.slides[4])
        if getattr(s, "fill", None) is not None and s.fill.type == 1
    }
    assert set(map(str, STUDIO_BEAUTY.chart_colors)) <= colors
    assert lint_pptx(path, whitespace=False)["errors"] == []


def test_compound_fact_fills_personalization_and_repairs_mismatched_headings_idempotently():
    data = content()
    data["facts"] = [
        PROMPT.split(". ")[0] + ".",
        "Our launch will focus on a shade-matching quiz, waitlist, and creator-led demos",
        "quarterly growth through UGC, subscriptions, new product lines, pop-ups, and festive scaling.",
    ]
    data["sections"] = [
        {
            "title": title,
            "points": ["Proposed next step: discuss the supplied details with your team."],
        }
        for title in ("Brand Overview", "Product Range", "Launch Strategy", "Growth Plan")
    ]
    model = planned_draft_model(PROMPT)
    result = model.model_validate(data)
    assert [section.title for section in result.sections] == [
        "Product Range",
        "Personalization",
        "Launch Strategy",
        "Growth Plan",
    ]
    assert "foundation, lipstick, and skincare" in " ".join(result.sections[0].points)
    assert "skin tone, undertone, and style" in " ".join(result.sections[1].points)
    assert "shade-matching quiz" in " ".join(result.sections[2].points)
    assert "festive scaling" in " ".join(result.sections[3].points)
    assert not any(
        "Add verified" in point for section in result.sections for point in section.points
    )
    assert model.model_validate(result.model_dump()).model_dump() == result.model_dump()


@pytest.mark.parametrize(
    "points,values,unit",
    [
        (["Marketing: 40%, Product: 35%, Operations: 25%"], (40, 35, 25), "%"),
        (["Q1: 1,000 users", "Q2: 2,400 users"], (1000, 2400), "users"),
        (["East: -12 units; West: 18 units"], (-12, 18), "units"),
    ],
)
def test_only_explicit_uniform_values_become_native_chart_data(points, values, unit):
    result = extract_chart(points)
    assert result.values == values
    assert result.unit == unit


@pytest.mark.parametrize(
    "points",
    [
        ["UGC, subscriptions, new product lines, pop-ups, and festive scaling."],
        ["Quiz completion: 70%; Customers: 1000 users"],
        ["Revenue rose by 20% and the team plans another launch."],
        ["Marketing: 30%; Marketing: 50%"],
        ["No supplied values"],
    ],
)
def test_missing_ambiguous_or_mixed_metrics_are_not_turned_into_charts(points):
    assert extract_chart(points) is None


def test_supplied_percentage_values_survive_the_saved_editable_chart(tmp_path):
    deck = Deck.model_validate_json(FIXTURE.read_text())
    data = {
        "layout": "draft_content",
        "title": "Supplied budget allocation",
        "points": ["Marketing: 40%, Product: 35%, Operations: 25%"],
    }
    deck = Deck.model_validate(
        {
            **deck.model_dump(),
            "theme": "studio",
            "slides": [
                deck.slides[0].model_dump(),
                data,
                deck.slides[-1].model_dump(),
            ],
        }
    )
    path = build_draft(deck, tmp_path / "data.pptx")
    chart = next(s.chart for s in Presentation(path).slides[1].shapes if s.has_chart)
    assert tuple(chart.series[0].values) == (40.0, 35.0, 25.0)
    assert [category.label for category in chart.plots[0].categories] == [
        "Marketing",
        "Product",
        "Operations",
    ]
    assert lint_pptx(path, whitespace=False)["errors"] == []


@pytest.mark.parametrize("count", [1, 2, 3, 4])
def test_general_topic_layouts_remain_filled_and_readable_without_beauty_source(tmp_path, count):
    points = [
        "Discuss the product concept and agree on the next planning priorities with the team."
    ] * count
    deck = Deck.model_validate(
        {
            "title": "Planning",
            "theme": "studio",
            "design_version": "draft-v1",
            "slides": [
                {"layout": "title", "title": "Planning", "subtitle": "A focused discussion"},
                {"layout": "draft_content", "title": "Priorities", "points": points},
                {
                    "layout": "closing",
                    "title": "Next steps",
                    "subtitle": "Review the planning priorities.",
                },
            ],
        }
    )
    path = build_draft(deck, tmp_path / "general.pptx")
    prs = Presentation(path)
    assert any(s.name == "Source topic" for s in prs.slides[0].shapes)
    assert not any(s.has_chart for slide in prs.slides for s in slide.shapes)
    assert lint_pptx(path, whitespace=False)["errors"] == []


def test_expanded_palettes_and_figure_text_have_readable_contrast():
    for theme in (STUDIO, STUDIO_BEAUTY):
        for foreground in (theme.primary, theme.text, theme.muted, theme.accent):
            assert contrast(foreground, theme.bg) >= 4.5
            assert contrast(foreground, theme.surface) >= 4.5
        canvas = Canvas(None, theme)
        assert all(contrast(canvas.ink(color), color) >= 4.5 for color in theme.chart_colors)
    assert draft_theme_for_topic("Blush skincare") == "studio-beauty"
    assert draft_theme_for_topic("Education planning") == "studio"
    assert contrast(rgb("FFFFFF"), STUDIO_BEAUTY.primary) >= 4.5
