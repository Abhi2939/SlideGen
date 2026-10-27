import hashlib
import json
import sqlite3
import threading
import time
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pydantic import ValidationError

from app.core.errors import DraftContentError, ProviderLimitError
from app.jobs.runner import JobRunner
from app.jobs.store import JobStore
from app.llm.client import LLMClient, prompt_tokens
from app.llm.structured import StructuredGenerator
from app.main import create_app
from app.pipeline.draft import generate_draft_deck, planned_draft_model
from app.qa.draft import build_draft
from test_editorial_pipeline import PROMPT
from test_llm_structured import FakeProvider


def content():
    return {
        "title": "Blush",
        "subtitle": "Custom beauty for your individual style",
        "facts": [
            "foundation, lipstick, and skincare",
            "skin tone, undertone, and style",
            "a shade-matching quiz, waitlist, and creator-led demos",
            "UGC, subscriptions, new product lines, pop-ups, and festive scaling",
        ],
        "sections": [
            {
                "title": "Custom beauty",
                "points": [
                    "Foundation, lipstick, and skincare blended exclusively for your skin tone, undertone, and style.",
                    "So you never get the wrong shade.",
                ],
            },
            {
                "title": "Launch touchpoints",
                "points": [
                    "Our launch will focus on a shade-matching quiz, waitlist, and creator-led demos.",
                    "Proposed next step: prepare a concise introduction to the product concept.",
                ],
            },
            {
                "title": "Quarterly growth",
                "points": [
                    "Quarterly growth through UGC, subscriptions, new product lines, pop-ups, and festive scaling.",
                    "Proposed next step: discuss the supplied growth activities with the team.",
                ],
            },
            {
                "title": "Planning priorities",
                "points": [
                    "Proposed next step: organize the supplied launch touchpoints.",
                    "Proposed next step: prepare the product introduction and growth discussion.",
                ],
            },
        ],
        "closing_title": "Plan the next step",
        "closing_text": "Discuss the proposed launch priorities.",
    }


def generator(values=None):
    provider = FakeProvider(values or [json.dumps(content())])
    return StructuredGenerator(LLMClient(provider)), provider


def wait_status(client, job_id, wanted):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] == wanted:
            return job
        if job["status"] == "failed":
            pytest.fail(str(job["failure"]))
        time.sleep(0.01)
    pytest.fail(f"job did not reach {wanted}")


def test_default_web_flow_is_one_content_request_and_download_without_visual_review(
    monkeypatch, tmp_path
):
    monkeypatch.setenv("SLIDEGEN_DELIVERY_MODE", "draft")
    monkeypatch.setenv("SLIDEGEN_DESIGN_PIPELINE", "editorial")  # Older Render setting.
    base, provider = generator()
    monkeypatch.setattr("app.pipeline.draft.create_llm_client", lambda *args: base.client)
    monkeypatch.setattr(
        "app.qa.gate.build_checked", lambda *args, **kwargs: pytest.fail("blocking visual review")
    )
    monkeypatch.setattr(
        "app.assets.resolve.prepare_assets", lambda *args, **kwargs: pytest.fail("asset download")
    )
    store = JobStore(tmp_path)
    with TestClient(create_app(store=store)) as client:
        response = client.post("/api/jobs", json={"prompt": PROMPT})
        assert response.status_code == 202
        job_id = response.json()["id"]
        wait_status(client, job_id, "done")
        response = client.get(f"/api/jobs/{job_id}/download")
        assert response.status_code == 200
        deck = Presentation(BytesIO(response.content))
        assert len(deck.slides) == 6
        assert all(slide.shapes for slide in deck.slides)
        assert client.get(f"/api/jobs/{job_id}/review").json()["status"] == "not_requested"
        report = client.get(f"/debug/{job_id}").json()
        assert report["delivery"] == "draft"
        assert report["file_checks"]["errors"] == []
        assert report["file_checks"]["minimum_font_pt"] >= 14
        assert "visual_review" in report and "passed" not in report["visual_review"]
    assert len(provider.calls) == 1
    assert provider.calls[0][1]["max_tokens"] == 3600


@pytest.mark.parametrize("change", ["invented_fact"])
def test_local_source_rules_reject_unsupported_or_missing_content(change):
    data = content()
    data["facts"].append("PRIVATE invented company claim")
    with pytest.raises(ValidationError):
        planned_draft_model(PROMPT).model_validate(data)


def test_company_copy_uses_source_excerpts_instead_of_invented_promises():
    data = content()
    data["sections"][0]["points"][0] = "PRIVATE: sales reached 100%."
    data["sections"][1]["points"][1] = "The waitlist provides exclusive early access."
    data["sections"][2]["points"][0] = "The model omitted the roadmap."
    grounded = planned_draft_model(PROMPT).model_validate(data)
    visible = " ".join(point for section in grounded.sections for point in section.points)
    assert "PRIVATE" not in visible and "100%" not in visible
    assert "early access" not in visible
    assert "UGC, subscriptions, new product lines, pop-ups, and festive scaling" in visible


def test_source_correction_is_bounded_and_does_not_repeat_the_whole_rejected_deck():
    bad = content()
    bad["facts"].append("PRIVATE_FIELD_MARKER: sales reached 100%.")
    fake, provider = generator([json.dumps(bad), json.dumps(content())])
    result = generate_draft_deck(PROMPT, generator=fake)
    assert len(result.slides) == 6
    assert len(provider.calls) == 2
    correction = provider.calls[1][0]
    assert "PRIVATE_FIELD_MARKER" not in json.dumps(correction)
    assert not any(message["role"] == "assistant" for message in correction)
    assert prompt_tokens(correction) + 3600 < 6800


def test_exhausted_company_validation_rebuilds_only_source_and_quota_still_fails():
    bad = content()
    bad["facts"].append("PRIVATE invented evidence")
    fake, provider = generator([json.dumps(bad)] * 2)
    deck = generate_draft_deck(PROMPT, generator=fake)
    assert len(provider.calls) == 2
    assert len(deck.slides) == 6
    visible = " ".join(point for slide in deck.slides for point in getattr(slide, "points", []))
    assert "PRIVATE invented evidence" not in visible
    assert "festive scaling" in visible
    assert all("Source-only draft" in slide.speaker_notes for slide in deck.slides)
    fake, provider = generator([ProviderLimitError()])
    with pytest.raises(ProviderLimitError):
        generate_draft_deck(PROMPT, generator=fake)
    assert len(provider.calls) == 1


def test_content_limit_failure_has_actionable_safe_details():
    from app.jobs.failures import describe_failure

    failure = describe_failure(DraftContentError(2))
    assert failure.code == "draft_content_invalid"
    assert failure.stage == "write"
    assert "could not be validated" in failure.message
    assert "job reference" in failure.message


@pytest.mark.parametrize("count", [3, 6, 10])
def test_requested_slide_count_uses_measured_editable_templates(tmp_path, count):
    data = content()
    data["facts"] = ["foundation, lipstick, and skincare"]
    data["sections"] = [data["sections"][0]] * (count - 2)
    fake, _ = generator([json.dumps(data)])
    deck = generate_draft_deck(PROMPT + f" Create {count}-slide presentation.", generator=fake)
    path = build_draft(deck, tmp_path / "draft.pptx")
    assert len(Presentation(path).slides) == count


def test_optional_review_failure_keeps_completed_job_and_download_intact(monkeypatch, tmp_path):
    fake, _ = generator()
    deck = generate_draft_deck(PROMPT, generator=fake)
    started, release = threading.Event(), threading.Event()

    def review(path, **kwargs):
        from app.qa.draft import write_json

        started.set()
        assert release.wait(5)
        write_json(
            path.parent / (path.stem + ".qa") / "review.json",
            {"status": "failed", "message": "Review unavailable."},
        )

    monkeypatch.setattr("app.qa.draft.review_draft", review)
    store = JobStore(tmp_path)
    with JobRunner(store, generate=lambda _: deck) as runner:
        with TestClient(create_app(runner=runner)) as client:
            job_id = client.post("/api/jobs", json={"prompt": PROMPT}).json()["id"]
            wait_status(client, job_id, "done")
            before = client.get(f"/api/jobs/{job_id}/download").content
            assert client.post(f"/api/jobs/{job_id}/review").status_code == 202
            assert started.wait(5)
            assert client.post(f"/api/jobs/{job_id}/review").json()["status"] == "queued"
            assert client.get(f"/api/jobs/{job_id}").json()["status"] == "done"
            assert client.get(f"/api/jobs/{job_id}/download").content == before
            next_id = client.post("/api/jobs", json={"prompt": "Another brief"}).json()["id"]
            wait_status(client, next_id, "done")
            assert client.post(f"/api/jobs/{next_id}/review").status_code == 409
            release.set()
    assert store.get(job_id).status == "done"
    from pathlib import Path

    assert (
        hashlib.sha256(before).hexdigest()
        == hashlib.sha256(Path(store.get(job_id).output_path).read_bytes()).hexdigest()
    )


def test_expired_queued_job_does_not_reset_its_deadline_on_restart(tmp_path):
    store = JobStore(tmp_path)
    old = store.create("An old queued brief")
    with sqlite3.connect(store.database_path) as connection:
        connection.execute(
            "UPDATE jobs SET created_at=? WHERE id=?",
            (
                (datetime.now(timezone.utc) - timedelta(minutes=11)).isoformat(),
                old.id,
            ),
        )
    calls = []
    with JobRunner(store, generate=lambda prompt: calls.append(prompt)):
        pass
    assert store.get(old.id).status == "failed"
    assert store.get(old.id).failure.code == "generation_timeout"
    assert calls == []


@pytest.mark.parametrize("quota_failure", [False, True])
def test_actual_optional_review_keeps_the_draft_even_for_low_scores_or_quota(
    tmp_path, quota_failure
):
    from app.qa.draft import review_draft
    from test_visual_gate import fake_render

    fake, _ = generator()
    deck = generate_draft_deck(PROMPT, generator=fake)
    path = build_draft(deck, tmp_path / "source-draft.pptx")
    before = path.read_bytes()
    calls = []

    def reviewer(*args, **kwargs):
        calls.append(kwargs["slide_index"])
        if quota_failure:
            raise ProviderLimitError()
        return {"score": 3, "fixes": [{"action": "rebalance_spacing", "detail": "Review spacing."}]}

    report = review_draft(path, reviewer=reviewer, render=fake_render)
    assert path.read_bytes() == before
    assert "passed" not in report
    if quota_failure:
        assert report["status"] == "failed"
        assert report["failure_type"] == "ProviderLimitError"
        assert calls == [0]
    else:
        assert report["status"] == "completed"
        assert len(report["grades"]) == 6
        assert all(grade["score"] == 3 for grade in report["grades"])
