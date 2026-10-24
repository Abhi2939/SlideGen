import json
from pathlib import Path
import sqlite3
import threading
import time
from datetime import datetime, timedelta, timezone

import pytest

from app.core.deck_schema import Deck
from app.core.errors import DeckReviewError, SlideWriteError
from app.jobs.runner import USER_ERROR, JobRunner
from app.jobs.store import JobStore
from app.pipeline.models import ReviewOutput


FIXTURE = Path(__file__).parent / "fixtures" / "decks" / "all_layouts.json"


def fixture_deck() -> Deck:
    return Deck.model_validate(json.loads(FIXTURE.read_text(encoding="utf-8")))


def test_expiration_removes_successful_and_failed_visual_reports(tmp_path):
    store = JobStore(tmp_path, retention_hours=1)
    for failed in (False, True):
        job = store.create("A test brief")
        folder = store.output_dir / (job.id + ".qa")
        (folder / "1").mkdir(parents=True)
        (folder / "1/slide-01.png").write_bytes(b"preview")
        if failed:
            store.mark_failed(job.id, "visual check failed")
        else:
            output = store.output_dir / (job.id + ".pptx")
            output.write_bytes(b"deck")
            store.mark_running(job.id)
            store.mark_done(job.id, output)
    assert store.cleanup_expired(now=datetime.now(timezone.utc) + timedelta(hours=2)) == 2
    assert list(store.output_dir.iterdir()) == []


def wait_for_status(store: JobStore, job_id: str, expected: str):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = store.get(job_id)
        if job and job.status == expected:
            return job
        time.sleep(0.01)
    pytest.fail(f"job {job_id} did not reach status {expected!r}")


def test_job_moves_from_queued_to_running_to_done(tmp_path):
    store = JobStore(tmp_path)
    generation_started = threading.Event()
    allow_generation_to_finish = threading.Event()

    def generate(prompt: str) -> Deck:
        assert prompt == "Make a Blush presentation"
        generation_started.set()
        assert allow_generation_to_finish.wait(timeout=5)
        return fixture_deck()

    runner = JobRunner(store, generate=generate)
    try:
        submitted = runner.submit("Make a Blush presentation")
        assert submitted.status == "queued"
        assert generation_started.wait(timeout=5)
        assert wait_for_status(store, submitted.id, "running").status == "running"

        allow_generation_to_finish.set()
        completed = wait_for_status(store, submitted.id, "done")

        assert completed.output_path is not None
        assert Path(completed.output_path).is_file()
        assert completed.error is None
    finally:
        allow_generation_to_finish.set()
        runner.close()


def test_job_failure_stores_clean_user_message(tmp_path):
    store = JobStore(tmp_path)

    def fail_generation(prompt: str) -> Deck:
        raise RuntimeError("private provider detail")

    runner = JobRunner(store, generate=fail_generation)
    try:
        submitted = runner.submit("Make a presentation")
        failed = wait_for_status(store, submitted.id, "failed")
        assert failed.error == USER_ERROR
        assert "private provider detail" not in failed.error
        assert failed.output_path is None
    finally:
        runner.close()


def test_exhausted_parse_failure_logs_only_sanitized_provider_error(tmp_path, caplog):
    from app.llm.client import LLMClient
    from app.llm.structured import StructuredGenerator
    from app.pipeline.models import AnalysisOutput

    calls = []

    class ParseFailure(Exception):
        status_code = 400
        body = {"error": {"code": "output_parse_failed", "failed_generation": "PRIVATE_REASONING"}}

    class Provider:
        def complete(self, *args, **kwargs):
            calls.append(1)
            raise ParseFailure("PRIVATE_REASONING PRIVATE_CREDENTIAL")

    generator = StructuredGenerator(LLMClient(Provider()))

    def generate(prompt):
        return generator.generate(
            AnalysisOutput,
            stage="analyze",
            model="openai/gpt-oss-20b",
            system_prompt="Return JSON",
            user_prompt=prompt,
        )

    store = JobStore(tmp_path)
    with JobRunner(store, generate=generate) as runner:
        job = runner.submit("Make a presentation")
    failed = store.get(job.id)
    assert failed.status == "failed"
    assert failed.error == USER_ERROR
    assert calls == [1, 1]
    assert "output_parse_failed" in caplog.text
    assert "PRIVATE_REASONING" not in caplog.text
    assert "PRIVATE_CREDENTIAL" not in caplog.text
    assert "failed_generation" not in caplog.text


def test_slide_validation_failure_logs_safe_reason_and_index(tmp_path, caplog):
    def generate(prompt):
        raise SlideWriteError(
            2, "PRIVATE_TITLE", "unprovided numeric claim: PRIVATE_COMPLETION", stage="repair"
        )

    store = JobStore(tmp_path)
    with JobRunner(store, generate=generate) as runner:
        job = runner.submit("PRIVATE_BRIEF")
    assert store.get(job.id).error == USER_ERROR
    assert "reason=unprovided_numeric_data stage=repair slide_index=2" in caplog.text
    assert "PRIVATE" not in caplog.text


def test_plain_value_error_logs_application_location_without_error_text(tmp_path, caplog):
    from app.pipeline.orchestrator import generate_deck

    def generate(prompt):
        return generate_deck(prompt, design="PRIVATE_SETTING")

    store = JobStore(tmp_path)
    with JobRunner(store, generate=generate) as runner:
        job = runner.submit("PRIVATE_BRIEF")
    assert store.get(job.id).error == USER_ERROR
    assert "type=ValueError" in caplog.text
    assert "location=pipeline/orchestrator.py:" in caplog.text
    assert "PRIVATE" not in caplog.text
    assert str(Path.cwd()) not in caplog.text


def test_job_recovers_outline_source_mismatch_and_saves_powerpoint(tmp_path, monkeypatch):
    from pptx import Presentation
    from app.pipeline.orchestrator import generate_deck
    from app.render.builder import build_deck
    from test_editorial_pipeline import PROMPT, SourceCorrectionGenerator

    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    # Exercise generation and real PPTX rendering; pixel QA is tested separately.
    monkeypatch.setattr(
        "app.qa.gate.build_checked", lambda deck, path, **kwargs: build_deck(deck, path)
    )
    fake = SourceCorrectionGenerator("PRIVATE_COMPLETION unsupported revenue")
    store = JobStore(tmp_path)
    with JobRunner(store, generate=lambda prompt: generate_deck(prompt, generator=fake)) as runner:
        job = runner.submit(PROMPT)
    completed = store.get(job.id)
    assert completed.status == "done"
    assert completed.error is None
    assert len(Presentation(completed.output_path).slides) == 9
    assert fake.outline_calls == 2


def test_exhausted_outline_source_failure_logs_reason_without_generated_text(tmp_path, caplog):
    from app.pipeline.orchestrator import generate_deck
    from test_editorial_pipeline import PROMPT, SourceCorrectionGenerator

    fake = SourceCorrectionGenerator("PRIVATE_COMPLETION unsupported revenue", recover=False)
    store = JobStore(tmp_path)
    with JobRunner(store, generate=lambda prompt: generate_deck(prompt, generator=fake)) as runner:
        job = runner.submit(PROMPT)
    failed = store.get(job.id)
    assert failed.status == "failed"
    assert failed.error == USER_ERROR
    assert failed.output_path is None
    assert fake.stages == ["outline", "outline"]
    assert "reason=outline_fact_not_in_source stage=outline" in caplog.text
    assert "PRIVATE" not in caplog.text


def test_job_recovers_roadmap_activity_mismatch_and_saves_powerpoint(tmp_path, monkeypatch):
    from pptx import Presentation
    from app.pipeline.orchestrator import generate_deck
    from app.render.builder import build_deck
    from test_editorial_pipeline import PROMPT, MismatchedRoadmapGenerator

    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    monkeypatch.setattr(
        "app.qa.gate.build_checked", lambda deck, path, **kwargs: build_deck(deck, path)
    )
    fake = MismatchedRoadmapGenerator()
    store = JobStore(tmp_path)
    with JobRunner(
        store, generate=lambda prompt: generate_deck(prompt, generator=fake, semantic_attempts=1)
    ) as runner:
        job = runner.submit(PROMPT)
    completed = store.get(job.id)
    assert completed.status == "done"
    assert completed.failure is None
    assert len(Presentation(completed.output_path).slides) == 9
    assert fake.stages.count("write") == 9


def test_job_recovers_analysis_parse_failure_and_saves_deck(tmp_path):
    from app.config import Settings
    from app.llm.client import LLMClient
    from app.llm.structured import StructuredGenerator
    from app.pipeline.orchestrator import generate_legacy_deck as generate_deck
    from test_llm_structured import parse_error
    from test_pipeline_orchestrator import FakePipelineGenerator

    from app.core.deck_schema import BulletsSlide, ClosingSlide, TitleSlide
    from app.pipeline.models import AnalysisOutput, OutlineOutput, ReviewOutput

    stage_responses = FakePipelineGenerator()
    output_models = {
        model.__name__: model
        for model in (
            AnalysisOutput,
            OutlineOutput,
            TitleSlide,
            BulletsSlide,
            ClosingSlide,
            ReviewOutput,
        )
    }
    calls = []

    class Provider:
        def complete(self, messages, **kwargs):
            schema = json.loads(messages[0]["content"].split("JSON Schema:\n")[1])
            calls.append(kwargs["response_format"])
            if len(calls) == 1:
                raise parse_error()
            return stage_responses.generate(output_models[schema["title"]]).model_dump_json()

    generator = StructuredGenerator(LLMClient(Provider()))
    settings = Settings("groq", None, "openai/gpt-oss-20b", None, None, None, None, 2)
    store = JobStore(tmp_path)
    with JobRunner(
        store, generate=lambda prompt: generate_deck(prompt, settings=settings, generator=generator)
    ) as runner:
        job = runner.submit("Make a 3-slide launch deck. Launch is June 1.")

    completed = store.get(job.id)
    assert completed.status == "done"
    assert completed.error is None
    assert Path(completed.output_path).is_file()
    assert len(calls) == 7
    assert all(call["type"] == "json_schema" for call in calls)


def test_job_corrects_out_of_range_review_references_and_saves_powerpoint(tmp_path):
    from app.config import Settings
    from app.pipeline.orchestrator import generate_legacy_deck as generate_deck
    from pptx import Presentation
    from test_pipeline_orchestrator import FakePipelineGenerator

    class Generator(FakePipelineGenerator):
        review_calls = 0

        def generate(self, output_model, **kwargs):
            if output_model is ReviewOutput:
                self.review_calls += 1
                if self.review_calls == 1:
                    return ReviewOutput(
                        passed=False,
                        issues=[
                            {
                                "category": "overflow",
                                "severity": "high",
                                "detail": "Check slides",
                                "slide_indices": [4, 5],
                            }
                        ],
                        regenerate_slide_indices=[4, 5],
                    )
                assert "0, 1, 2" in kwargs["user_prompt"]
            return super().generate(output_model, **kwargs)

    generator = Generator()
    store = JobStore(tmp_path)
    settings = Settings("groq", None, "fake", None, None, None, None, 2)
    with JobRunner(
        store, generate=lambda prompt: generate_deck(prompt, settings=settings, generator=generator)
    ) as runner:
        job = runner.submit("Make a 3-slide launch deck. Launch is June 1.")
    completed = store.get(job.id)
    assert completed.status == "done"
    assert generator.review_calls == 2
    assert len(Presentation(completed.output_path).slides) == 3


def test_job_timeout_fails_with_clean_user_message(tmp_path):
    store = JobStore(tmp_path)
    started = threading.Event()
    release = threading.Event()

    def slow_generate(prompt: str) -> Deck:
        started.set()
        assert release.wait(timeout=5)
        return fixture_deck()

    runner = JobRunner(
        store,
        generate=slow_generate,
        max_runtime_seconds=0.05,
    )
    try:
        submitted = runner.submit("Make a presentation")
        assert started.wait(timeout=5)
        failed = wait_for_status(store, submitted.id, "failed")
        assert failed.error == USER_ERROR
        assert failed.output_path is None
    finally:
        release.set()
        runner.close()


def test_review_failure_logs_findings_but_keeps_user_error_generic(tmp_path, caplog):
    store = JobStore(tmp_path)
    review = ReviewOutput(
        passed=False,
        issues=[
            {
                "category": "unsupported_claim",
                "severity": "medium",
                "detail": "Unsupported AI capability",
                "slide_indices": [1],
            }
        ],
        regenerate_slide_indices=[1],
    )

    def fail_review(prompt: str) -> Deck:
        raise DeckReviewError(review, repair_rounds=2)

    runner = JobRunner(store, generate=fail_review)
    try:
        submitted = runner.submit("Make a product deck")
        failed = wait_for_status(store, submitted.id, "failed")
        assert failed.error == USER_ERROR
        assert "Unsupported AI capability" not in caplog.text
        assert "issues=1" in caplog.text
        assert submitted.id in caplog.text
    finally:
        runner.close()


def test_startup_recovery_marks_interrupted_jobs_failed(tmp_path):
    store_before_restart = JobStore(tmp_path)
    queued = store_before_restart.create("Interrupted prompt")
    store_before_restart.mark_running(queued.id)

    store_after_restart = JobStore(tmp_path)

    recovered = store_after_restart.get(queued.id)
    assert recovered is not None
    assert recovered.status == "failed"
    assert (
        recovered.error
        == "Server restarted while this job was waiting or running. Please submit it again."
    )
    assert recovered.output_path is None


def test_startup_recovery_marks_queued_jobs_failed(tmp_path):
    store_before_restart = JobStore(tmp_path)
    queued = store_before_restart.create("Waiting prompt")

    store_after_restart = JobStore(tmp_path)

    recovered = store_after_restart.get(queued.id)
    assert recovered is not None
    assert recovered.status == "failed"
    assert (
        recovered.error
        == "Server restarted while this job was waiting or running. Please submit it again."
    )
    assert store_after_restart.count_active() == 0


def test_job_store_migrates_existing_database_for_stage_tracking(tmp_path):
    database = tmp_path / "jobs.sqlite3"
    with sqlite3.connect(database) as connection:
        connection.execute(
            """CREATE TABLE jobs (
                id TEXT PRIMARY KEY,
                prompt TEXT NOT NULL,
                status TEXT NOT NULL CHECK (status IN ('queued', 'running', 'done', 'failed')),
                output_path TEXT,
                error TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )"""
        )

    store = JobStore(tmp_path)
    queued = store.create("Migrated job")

    assert queued.status == "queued"
    assert queued.stage == "Waiting in queue"
