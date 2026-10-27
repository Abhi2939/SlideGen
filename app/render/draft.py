"""Reference-inspired editable diagrams, with source text and no network dependency."""

import re

from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.render.common import add_rect, add_text, new_slide
from app.render.measure import fitted_size, measure_text, TextNeedsSplit
from app.render.theme import rgb

GREEN = rgb("8CC614")
LINE = rgb("D6DBE1")


def _list(text, pattern):
    """Only extract an explicit, bounded source list; never generate descriptions."""
    match = re.search(pattern, text, re.I)
    if not match:
        return []
    values = re.split(r",\s*(?:and\s+)?|\s+and\s+", match.group(1))
    values = [value.strip(" ,.;:") for value in values if value.strip(" ,.;:")]
    return values if 3 <= len(values) <= 6 and all(len(v) <= 45 for v in values) else []


def composition(points):
    """Select a visual from supplied content, independently of a generated heading."""
    source = " ".join(points)
    if len(source.split()) > 45:
        return "points", []
    patterns = (
        ("growth", r"\bgrowth\s+through\s+([^.!?]+)"),
        ("launch", r"\b(?:focus\s+on|touchpoints?\s*:)\s+(.+?)(?:,?\s+followed\s+by|[.!?]|$)"),
        ("collection", r"\b(?:offering|collection\s*:|range\s*:)\s+(.+?)(?:\s+blended\b|[.!?]|$)"),
        (
            "personalization",
            r"\b(?:for\s+your|around\s+your|personalization\s*:)\s+(.+?)(?:,?\s+so\b|[.!?]|$)",
        ),
    )
    for kind, pattern in patterns:
        labels = _list(source, pattern)
        if labels:
            return kind, labels
    # Source excerpts may contain just a list without its introductory clause.
    labels = _list(source, r"^([^.!?]+)[.!?]?$")
    if labels:
        if any("quiz" in value.casefold() for value in labels):
            return "launch", labels
        if any("subscription" in value.casefold() for value in labels):
            return "growth", labels
        if {v.casefold() for v in labels} == {"foundation", "lipstick", "skincare"}:
            return "collection", labels
        if {v.casefold() for v in labels} == {"skin tone", "undertone", "style"}:
            return "personalization", labels
    return "points", []


class Canvas:
    def __init__(self, slide, theme):
        self.slide, self.theme = slide, theme

    def panel(self, x, y, w, h, color, name="panel", kind=MSO_SHAPE.RECTANGLE):
        shape = add_rect(self.slide, *map(Inches, (x, y, w, h)), color, kind)
        shape.name = "Layer:background:" + name
        return shape

    def text(
        self, value, x, y, w, h, size=22, *, bold=False, color=None, name="Text", center=False
    ):
        if not value:
            return None
        point = fitted_size(value, w, h, size, self.theme.font_body, bold)
        shape = add_text(
            self.slide,
            *map(Inches, (x, y, w, h)),
            value,
            size_pt=point,
            color=color or self.theme.text,
            font=self.theme.font_body,
            bold=bold,
            align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
        )
        shape.name = name
        return shape

    def line(self, x1, y1, x2, y2, *, color=LINE, width=1, dashed=False, name="rule"):
        line = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, *map(Inches, (x1, y1, x2, y2))
        )
        line.name = "Layer:background:" + name
        line.line.color.rgb, line.line.width = color, Pt(width)
        line.shadow.inherit = False
        for effect in line._element.iter(qn("a:effectRef")):
            effect.set("idx", "0")
        if dashed:
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return line

    def icon(self, label, cx, cy, size=0.45, color=None):
        """Small editable line illustrations rather than a network-fetched icon."""
        color = color or self.theme.primary
        x, y, u = cx - size / 2, cy - size / 2, size

        def line(a, b, c, d):
            self.line(x + a * u, y + b * u, x + c * u, y + d * u, color=color, width=2)

        def outline(a, b, w, h, kind=MSO_SHAPE.RECTANGLE):
            shape = self.panel(x + a * u, y + b * u, w * u, h * u, color, "icon", kind)
            shape.fill.background()
            shape.line.color.rgb, shape.line.width = color, Pt(2)

        key = label.casefold()
        if "quiz" in key or "plan" in key:
            outline(0.1, 0.03, 0.8, 0.94, MSO_SHAPE.ROUNDED_RECTANGLE)
            for yy in (0.3, 0.54, 0.78):
                line(0.23, yy, 0.34, yy + 0.07)
                line(0.34, yy + 0.07, 0.45, yy - 0.05)
                line(0.56, yy, 0.77, yy)
        elif "creator" in key or "demo" in key:
            outline(0.03, 0.2, 0.94, 0.64, MSO_SHAPE.ROUNDED_RECTANGLE)
            for values in (
                (0.4, 0.36, 0.4, 0.69),
                (0.4, 0.36, 0.67, 0.52),
                (0.4, 0.69, 0.67, 0.52),
            ):
                line(*values)
        elif "waitlist" in key or "ugc" in key:
            outline(0.34, 0.05, 0.32, 0.32, MSO_SHAPE.OVAL)
            outline(0.18, 0.51, 0.64, 0.4, MSO_SHAPE.ARC)
            line(0.06, 0.18, 0.06, 0.43)
            line(0.0, 0.3, 0.15, 0.3)
        elif "subscription" in key:
            outline(0.13, 0.13, 0.74, 0.74, MSO_SHAPE.OVAL)
            line(0.49, 0.27, 0.49, 0.53)
            line(0.49, 0.53, 0.68, 0.63)
        elif "pop" in key:
            outline(0.15, 0.35, 0.7, 0.58)
            outline(0.06, 0.1, 0.88, 0.25, MSO_SHAPE.TRAPEZOID)
            line(0.48, 0.58, 0.48, 0.94)
        elif "festiv" in key:
            outline(0.1, 0.4, 0.8, 0.52)
            line(0.49, 0.14, 0.49, 0.91)
            line(0.05, 0.4, 0.95, 0.4)
            line(0.22, 0.12, 0.49, 0.31)
            line(0.49, 0.31, 0.8, 0.12)
        else:
            for yy in (0.12, 0.48):
                for xx in (0.12, 0.55):
                    outline(xx, yy, 0.3, 0.3, MSO_SHAPE.ROUNDED_RECTANGLE)

    def product(self, label, cx, y, scale=1.0, inverse=False):
        """Editable illustrative silhouettes, not photographs of a brand's products."""
        ink = self.theme.on_primary if inverse else self.theme.primary
        fill = rgb("CFD9E1") if inverse else rgb("DDE6EE")
        u = scale

        def rect(dx, dy, w, h, color, kind=MSO_SHAPE.RECTANGLE):
            return self.panel(cx + dx * u, y + dy * u, w * u, h * u, color, "product", kind)

        key = label.casefold()
        if "foundation" in key:
            rect(-0.4, 0.5, 0.8, 1.5, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
            rect(-0.24, 0.2, 0.48, 0.32, ink)
            rect(-0.24, 0.1, 0.65, 0.13, ink)
            rect(-0.35, 1.18, 0.7, 0.45, self.theme.primary if inverse else self.theme.surface)
        elif "lipstick" in key:
            rect(-0.25, 0.2, 0.5, 0.95, GREEN, MSO_SHAPE.ROUNDED_RECTANGLE)
            rect(-0.34, 1.06, 0.68, 0.94, ink)
            rect(-0.34, 1.06, 0.68, 0.14, GREEN)
        else:
            rect(-0.51, 0.94, 1.02, 1.06, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
            rect(-0.55, 0.76, 1.1, 0.26, ink)
            rect(-0.43, 1.38, 0.86, 0.32, self.theme.primary if inverse else self.theme.surface)


def _source(c, points, *, bottom=2.7):
    value = " ".join(points)
    for size in (19, 18, 16):
        height = measure_text(value, 12.0, size, c.theme.font_body)[1] + 0.07
        if height <= bottom - 1.55:
            c.text(
                value, 0.65, 1.55, 12.0, height, size, color=c.theme.muted, name="Source context"
            )
            return
    raise TextNeedsSplit("source context needs another slide")


def _flow(c, labels, *, growth=False):
    n = len(labels)
    step = 12.0 / n
    centers = [0.65 + (i + 0.5) * step for i in range(n)]
    size = 21 if n < 5 else 18
    label_h = max(
        0.72,
        *(measure_text(label, step - 0.4, size, c.theme.font_body, True)[1] for label in labels),
    )
    cy = 3.45
    c.line(centers[0], cy, centers[-1], cy, color=GREEN, width=2.5, name="workflow")
    for i, (label, cx) in enumerate(zip(labels, centers)):
        c.panel(cx - 0.44, cy - 0.44, 0.88, 0.88, GREEN, "workflow-node", MSO_SHAPE.OVAL)
        c.icon(label, cx, cy, 0.4)
        c.panel(cx + 0.24, cy - 0.43, 0.31, 0.31, c.theme.primary, "sequence", MSO_SHAPE.OVAL)
        c.text(
            str(i + 1),
            cx + 0.245,
            cy - 0.405,
            0.3,
            0.25,
            14,
            color=c.theme.on_primary,
            center=True,
            name="Sequence number",
        )
        # Staggered stems and labels echo the supplied process diagram.
        label_y = 4.5 if growth and i % 2 else 4.16
        c.line(cx, cy + 0.48, cx, label_y - 0.12, dashed=True)
        w = step - 0.2
        c.panel(
            cx - w / 2,
            label_y,
            w,
            label_h + 0.4,
            c.theme.surface,
            "workflow-label",
            MSO_SHAPE.ROUNDED_RECTANGLE,
        )
        display = re.sub(r"^(?:a|an|the)\s+", "", label, flags=re.I)
        c.text(
            display[0].upper() + display[1:],
            cx - w / 2 + 0.1,
            label_y + 0.2,
            w - 0.2,
            label_h,
            size,
            bold=True,
            color=c.theme.primary,
            center=True,
            name="Workflow activity",
        )


def _collection(c, labels):
    step = 12.0 / len(labels)
    for i, label in enumerate(labels):
        x, w = 0.65 + i * step, step - 0.25
        c.panel(x, 2.9, w, 3.05, c.theme.surface, "collection-panel", MSO_SHAPE.ROUNDED_RECTANGLE)
        c.panel(x, 2.9, w, 0.08, GREEN, "collection-accent")
        if label.casefold() in {"foundation", "lipstick", "skincare"}:
            c.product(label, x + w / 2, 3.16, 0.93)
        else:
            c.icon(label, x + w / 2, 4.1, 0.7)
        c.text(
            label[0].upper() + label[1:],
            x + 0.15,
            5.28,
            w - 0.3,
            0.5,
            24,
            bold=True,
            color=c.theme.primary,
            center=True,
            name="Collection label",
        )


def _personalization(c, labels):
    step = 12.0 / len(labels)
    for i, label in enumerate(labels):
        cx = 0.65 + (i + 0.5) * step
        c.panel(cx - 0.72, 3.12, 1.44, 1.44, c.theme.surface, "personalization", MSO_SHAPE.OVAL)
        # These are category illustrations, not a claimed diagnostic or matching method.
        if "style" in label.casefold():
            for angle, tone in ((-25, "DABDA4"), (0, "263FB8"), (25, "8CC614")):
                swatch = c.panel(
                    cx - 0.14, 3.45, 0.28, 0.72, rgb(tone), "swatch", MSO_SHAPE.ROUNDED_RECTANGLE
                )
                swatch.rotation = angle
        else:
            tones = (
                ("DABDA4", "BB937A", "8C6958")
                if "tone" == label.casefold().split()[-1]
                else ("D2AA8E", "D4C7B1", "C1A7BB")
            )
            for j, tone in enumerate(tones):
                c.panel(cx - 0.43 + j * 0.29, 3.61, 0.35, 0.5, rgb(tone), "swatch", MSO_SHAPE.OVAL)
        c.text(
            label[0].upper() + label[1:],
            cx - step / 2 + 0.12,
            4.94,
            step - 0.24,
            0.8,
            26,
            bold=True,
            color=c.theme.primary,
            center=True,
            name="Personalization label",
        )


def _points(c, points):
    if len(points) == 1:
        c.panel(0.65, 2.45, 0.08, 2.85, GREEN, "statement-rule")
        value = points[0]
        for size in (36, 30, 26, 22, 18):
            if measure_text(value, 10.9, size, c.theme.font_body, True)[1] <= 2.6:
                c.text(
                    value,
                    1.12,
                    2.55,
                    10.9,
                    2.6,
                    size,
                    bold=True,
                    color=c.theme.primary,
                    name="Statement",
                )
                return
        raise TextNeedsSplit("statement needs another slide")
    columns = 2 if len(points) in {2, 4} else 3
    rows = (len(points) + columns - 1) // columns
    step, row_h = 12.0 / columns, 3.95 / rows
    for i, value in enumerate(points):
        x, y = 0.65 + i % columns * step, 2.15 + i // columns * row_h
        w = step - 0.35
        c.panel(x, y, w, 0.07, GREEN, "point-rule")
        c.text(f"{i + 1:02d}", x, y + 0.28, 0.7, 0.5, 24, bold=True, color=c.theme.accent)
        h = row_h - 1.02
        for size in (26, 24, 22, 20, 18):
            if measure_text(value, w, size, c.theme.font_body)[1] <= h:
                c.text(value, x, y + 0.95, w, h, size, name="Source point")
                break
        else:
            raise TextNeedsSplit("point grid needs another slide")


def render(prs, data, theme, *, brand, number, deck=None):
    if theme.name.startswith("studio"):
        from app.render.studio import render as render_studio

        return render_studio(prs, data, theme, brand=brand, number=number, deck=deck)
    slide = new_slide(prs, theme)
    c = Canvas(slide, theme)
    cover = data.layout == "title"
    closing = data.layout == "closing"
    source = " ".join(
        str(value)
        for part in (deck.slides if deck else [data])
        for value in [getattr(part, "subtitle", "") or "", *getattr(part, "points", [])]
    )
    beauty = all(
        re.search(r"\b" + label + r"\b", source, re.I)
        for label in ("foundation", "lipstick", "skincare")
    )
    footer_color = theme.muted
    if closing:
        slide.background.fill.fore_color.rgb = theme.primary
        c.panel(0.65, 1.55, 1.0, 0.08, GREEN, "closing-rule")
        c.text(
            data.title,
            0.65,
            2.0,
            11.7,
            1.65,
            46,
            bold=True,
            color=theme.on_primary,
            name="Slide title",
        )
        c.text(getattr(data, "subtitle", None), 0.65, 4.0, 10.8, 1.4, 30, color=theme.on_primary)
        footer_color = theme.on_primary
        kind = "closing"
    elif cover:
        slide.background.fill.fore_color.rgb = theme.surface
        c.panel(7.7, 0.6, 5.0, 5.85, theme.primary, "brand-panel", MSO_SHAPE.ROUNDED_RECTANGLE)
        c.panel(0.65, 1.9, 0.85, 0.08, GREEN, "cover-rule")
        title_size = next(
            (
                size
                for size in (60, 52, 44, 36, 30)
                if measure_text(data.title, 6.65, size, theme.font_body, True)[1] <= 1.9
            ),
            30,
        )
        title_h = measure_text(data.title, 6.65, title_size, theme.font_body, True)[1] + 0.05
        c.text(
            data.title,
            0.65,
            2.35,
            6.65,
            title_h,
            title_size,
            bold=True,
            color=theme.primary,
            name="Slide title",
        )
        subtitle_y = 2.35 + title_h + 0.4
        c.text(getattr(data, "subtitle", None), 0.65, subtitle_y, 6.4, 1.25, 24, color=theme.muted)
        if beauty:
            for label, cx, yy in (
                ("Foundation", 8.62, 2.22),
                ("Lipstick", 10.1, 2.55),
                ("Skincare", 11.62, 2.25),
            ):
                c.product(label, cx, yy, 1.05, inverse=True)
                c.text(
                    label,
                    cx - 0.7,
                    5.2,
                    1.4,
                    0.4,
                    16,
                    color=theme.on_primary,
                    center=True,
                    name="Collection label",
                )
        else:
            sections = [part.title for part in (deck.slides[1:-1] if deck else [])][:4]
            for i, title in enumerate(sections):
                yy = 1.2 + i * 1.17
                c.panel(8.1, yy + 0.1, 0.12, 0.42, GREEN, "cover-topic-rule")
                c.text(title, 8.5, yy, 3.8, 0.9, 22, color=theme.on_primary, name="Cover topic")
        kind = "cover"
    else:
        kind, labels = composition(data.points)
        title = data.title
        if title.casefold() == "cover":
            title = "The collection" if kind == "collection" else "Overview"
        c.text(
            title, 0.65, 0.62, 12.0, 0.85, 32, bold=True, color=theme.primary, name="Slide title"
        )
        if kind == "points":
            _points(c, data.points)
        else:
            _source(c, data.points)
            if kind in {"launch", "growth"}:
                _flow(c, labels, growth=kind == "growth")
            elif kind == "collection":
                _collection(c, labels)
            elif kind == "personalization":
                _personalization(c, labels)
    slide._element.cSld.set("name", "draft_" + kind)
    c.line(0.65, 6.45, 12.65, 6.45)
    c.text(brand[:45], 0.65, 6.6, 6.5, 0.26, 14, color=footer_color, name="Footer:brand")
    c.text(
        f"Draft  |  {number:02d}",
        10.75,
        6.6,
        1.9,
        0.26,
        14,
        color=footer_color,
        name="Footer:number",
    )
    return slide
