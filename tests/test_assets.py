import io
import json
import socket
import time

import httpx
from PIL import Image
from pptx import Presentation
import pytest

from app.assets.fetch import download, image_png, public_address
from app.assets.resolve import AssetResolver, AssetSettings, icon_name, topic_query
from app.render.builder import build_deck


SVG = b'<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24"><circle cx="12" cy="12" r="8" stroke="navy" fill="none"/></svg>'


def png():
    out = io.BytesIO()
    Image.new("RGB", (80, 60), "navy").save(out, format="PNG")
    return out.getvalue()


@pytest.fixture
def fake_http(monkeypatch):
    calls = []

    def respond(request):
        calls.append(request)
        assert request.extensions["sni_hostname"] == request.headers["host"]
        if request.headers["host"] == "api.pexels.com":
            return httpx.Response(
                200,
                json={
                    "photos": [
                        {
                            "src": {"large": "https://images.pexels.com/photos/123/photo.jpeg"},
                            "url": "https://www.pexels.com/photo/test-123/",
                            "photographer": "Fixture Photographer",
                        }
                    ]
                },
            )
        if request.headers["host"] == "images.pexels.com":
            return httpx.Response(200, content=png(), headers={"content-type": "image/png"})
        return httpx.Response(200, content=SVG, headers={"content-type": "image/svg+xml"})

    original = httpx.Client
    monkeypatch.setattr("app.assets.fetch.public_address", lambda host: "93.184.216.34")
    monkeypatch.setattr(
        httpx, "Client", lambda **kw: original(transport=httpx.MockTransport(respond), **kw)
    )
    return calls


def test_full_fake_pipeline_embeds_images_and_sources_without_render_network(
    tmp_path, monkeypatch, fake_http
):
    from test_blush_prompt import BlushGenerator, PROMPT
    from app.config import Settings
    from app.pipeline.orchestrator import generate_deck

    resolver = AssetResolver(AssetSettings(cache_dir=tmp_path / "cache", pexels_key="FAKE_KEY"))
    monkeypatch.setattr("app.pipeline.orchestrator.prepare_assets", resolver.resolve)
    deck = generate_deck(
        PROMPT,
        settings=Settings("groq", None, "fake", None, None, None, None, 2),
        generator=BlushGenerator(),
    )
    assert deck.assets
    assert any(asset.role == "photo" for asset in deck.assets)
    monkeypatch.setattr(
        "app.assets.resolve.download", lambda *a, **kw: pytest.fail("renderer fetched an asset")
    )
    out = build_deck(deck, tmp_path / "assets.pptx", strict=True)
    prs = Presentation(out)
    assert any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes)
    assert "Fixture Photographer" in prs.slides[0].notes_slide.notes_text_frame.text
    assert "Pexels" in prs.slides[0].notes_slide.notes_text_frame.text
    assert any("Pexels" in shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert all(request.url.host == "93.184.216.34" for request in fake_http)
    assert list((tmp_path / "cache/manifests").glob("*.json"))


def test_cache_hit_does_not_fetch_and_corrupt_cache_is_replaced(tmp_path, fake_http):
    resolver = AssetResolver(AssetSettings(cache_dir=tmp_path))
    first = resolver._resolve("icon", "circle", "1b365d")
    count = len(fake_http)
    assert resolver._resolve("icon", "circle", "1b365d") == first
    assert len(fake_http) == count
    from pathlib import Path

    Path(first["path"]).write_bytes(b"corrupt")
    assert resolver._resolve("icon", "circle", "1b365d")["sha256"] == first["sha256"]
    assert len(fake_http) == count + 1


def test_missing_photo_key_does_not_make_http_request(tmp_path, fake_http):
    assert AssetResolver(AssetSettings(cache_dir=tmp_path))._resolve("photo", "dogs", 0) is None
    assert fake_http == []


@pytest.mark.parametrize(
    "url",
    [
        "http://api.iconify.design/lucide/circle.svg",
        "https://127.0.0.1/a.png",
        "https://169.254.169.254/latest/meta-data/",
        "https://example.com/a.png",
        "https://api.iconify.design.evil.test/a.png",
        "https://user:pass@api.iconify.design/a.svg",
        "https://api.iconify.design:444/a.svg",
    ],
)
def test_blocked_hosts_and_non_https_never_connect(url, fake_http):
    with pytest.raises(ValueError):
        download(url)
    assert fake_http == []


@pytest.mark.parametrize("address", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "::1", "fc00::1"])
def test_dns_private_destinations_rejected(monkeypatch, address):
    monkeypatch.setattr(
        socket, "getaddrinfo", lambda *a, **kw: [(None, None, None, None, (address, 443))]
    )
    with pytest.raises(ValueError):
        public_address("api.iconify.design")


@pytest.mark.parametrize(
    "response",
    [
        httpx.Response(302, headers={"location": "http://127.0.0.1/"}),
        httpx.Response(200, content=b"x", headers={"content-length": "9000000"}),
        httpx.Response(200, content=b"123456789"),
    ],
)
def test_redirect_and_oversized_responses_rejected(monkeypatch, response):
    original = httpx.Client
    monkeypatch.setattr("app.assets.fetch.public_address", lambda host: "93.184.216.34")
    monkeypatch.setattr(
        httpx,
        "Client",
        lambda **kw: original(transport=httpx.MockTransport(lambda req: response), **kw),
    )
    with pytest.raises((ValueError, httpx.HTTPStatusError)):
        download("https://api.iconify.design/lucide/circle.svg", max_bytes=5)


@pytest.mark.parametrize(
    "data,kind",
    [
        (b"<html>error</html>", "text/html"),
        (b"not a png", "image/png"),
        (b'<svg><image href="http://localhost/secret"/></svg>', "image/svg+xml"),
        (b"<!DOCTYPE svg><svg/>", "image/svg+xml"),
    ],
)
def test_wrong_content_and_unsafe_svg_rejected(data, kind):
    with pytest.raises(Exception):
        image_png(data, kind)


def test_timeout_is_optional_and_no_logo_falls_back_to_text(tmp_path, monkeypatch):
    from app.core.deck_schema import Deck

    def fail(*args, **kwargs):
        raise httpx.ReadTimeout("PRIVATE_KEY")

    monkeypatch.setattr("app.assets.resolve.download", fail)
    deck = Deck.model_validate(
        {
            "title": "Competitors",
            "slides": [
                {"layout": "title", "title": "Competitors"},
                {
                    "layout": "quadrants",
                    "title": "Competitive landscape",
                    "quadrants": [
                        {"heading": "Cards", "names": ["Visa", "InventedCo"]},
                        *[{"heading": "Other", "names": []} for _ in range(3)],
                    ],
                },
                {"layout": "closing", "title": "Summary"},
            ],
        }
    )
    resolved = AssetResolver(AssetSettings(cache_dir=tmp_path / "cache")).resolve(
        deck, "Competitors: Visa"
    )
    assert not resolved.assets
    prs = Presentation(build_deck(resolved, tmp_path / "fallback.pptx", strict=True))
    assert any(s.has_text_frame and s.text == "Visa" for s in prs.slides[1].shapes)


def test_logos_only_fetched_for_names_in_prompt(tmp_path, monkeypatch):
    from app.core.deck_schema import Deck
    from pathlib import Path

    deck = Deck.model_validate(
        json.loads(Path("tests/fixtures/decks/reference_layouts.json").read_text(encoding="utf-8"))
    )
    names = []

    def collect(role, query, variant):
        if role == "logo":
            names.append(query)

    monkeypatch.setattr(AssetResolver, "_resolve", lambda self, *args: collect(*args))
    AssetResolver(AssetSettings(cache_dir=tmp_path)).resolve(
        deck, "Competitive landscape: Visa and Polygon"
    )
    assert set(names) == {"visa", "polygon"}


def test_total_asset_budget_does_not_wait_for_slow_workers(tmp_path, monkeypatch):
    from test_blush_prompt import offline_deck

    deck, _ = offline_deck()

    def slow(*args):
        time.sleep(0.15)

    monkeypatch.setattr(AssetResolver, "_resolve", slow)
    start = time.monotonic()
    resolved = AssetResolver(AssetSettings(cache_dir=tmp_path, budget_seconds=0.01)).resolve(
        deck, "Blush"
    )
    assert time.monotonic() - start < 0.12
    assert resolved.assets == []


def test_keyword_mapping_and_private_prompt_not_sent_as_photo_query():
    assert icon_name("Verifies KYC") == "shield-check"
    assert icon_name("Invoice") == "file-text"
    assert topic_query("Blush private@example.com") == "cosmetics makeup"
