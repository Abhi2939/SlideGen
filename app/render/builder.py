import logging
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from app.core.deck_schema import Deck
from app.render.common import add_text, add_title, new_slide
from app.render.assets import asset_scope
from app.render.layout import LAYOUT_RENDERERS
from app.render.theme import BODY_TOP, CONTENT_W, MARGIN_X, SLIDE_H, SLIDE_W, get_theme

log = logging.getLogger(__name__)


def _fallback(prs, data, theme):
    slide = new_slide(prs, theme)
    add_title(slide, theme, data.title)
    add_text(
        slide,
        MARGIN_X,
        BODY_TOP,
        CONTENT_W,
        Inches(0.5),
        f"[{data.layout} layout not implemented yet]",
        size_pt=16,
        color=theme.muted,
        font=theme.font_body,
    )
    return slide


def _drop_slides_after(prs, n):
    """Remove any slides beyond the first n (cleans up a half-drawn slide)."""
    ids = prs.slides._sldIdLst
    for sld in list(ids)[n:]:
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def build_deck(deck: Deck, output_path, strict: bool = False) -> Path:
    theme = get_theme(deck.theme)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for index, data in enumerate(deck.slides):
        assets = [a for a in deck.assets if a.slide_index == index]
        renderer = LAYOUT_RENDERERS.get(data.layout, _fallback)
        count = len(prs.slides)
        try:
            with asset_scope(assets):
                slide = renderer(prs, data, theme)
        except Exception:
            if strict:
                raise
            log.exception("layout %r failed, using fallback", data.layout)
            _drop_slides_after(prs, count)
            slide = _fallback(prs, data, theme)

        notes = [data.speaker_notes] if data.speaker_notes else []
        for asset in assets:
            notes.append(
                f"Visual source ({asset.key}): {asset.source_url}\n"
                f"Credit: {asset.photographer or asset.query}. Terms: {asset.licence}"
            )
        if any(a.role == "icon" for a in assets):
            notice = Path(__file__).resolve().parents[1] / "assets" / "LUCIDE-LICENSE"
            try:
                notes.append(notice.read_text(encoding="utf-8"))
            except OSError:
                pass
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
        photos = [a for a in assets if a.role == "photo"]
        if photos:
            credit = "Photos: " + "; ".join(f"{a.photographer} / Pexels" for a in photos)
            box = add_text(
                slide,
                Inches(0.55),
                Inches(7.15),
                Inches(11.8),
                Inches(0.22),
                credit[:155],
                size_pt=8,
                color=theme.on_primary
                if data.layout in {"title", "section", "closing"}
                else theme.muted,
                font=theme.font_body,
            )
            box.text_frame.paragraphs[0].runs[0].hyperlink.address = photos[0].source_url

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path
