"""Inspect actual exported shapes, fonts and images, not only source JSON."""

from io import BytesIO
from itertools import combinations
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE

from app.render.measure import measure_text
from app.render.theme import contrast

EMU = 914400


def rect(shape):
    return tuple(float(v) / EMU for v in (shape.left, shape.top, shape.width, shape.height))


def intersection(a, b):
    x, y, w, h = a
    xx, yy, ww, hh = b
    return max(0, min(x + w, xx + ww) - max(x, xx)) * max(0, min(y + h, yy + hh) - max(y, yy))


def contains(a, b):
    return intersection(a, b) >= b[2] * b[3] - 0.005


def union_area(boxes):
    xs = sorted({p for x, _, w, _ in boxes for p in (x, x + w)})
    area = 0
    for x1, x2 in zip(xs, xs[1:]):
        spans = sorted(
            (y, y + h) for x, y, w, h in boxes if x <= x1 + 0.001 and x + w >= x2 - 0.001
        )
        end = -1
        for start, stop in spans:
            area += max(0, stop - max(start, end)) * (x2 - x1)
            end = max(end, stop)
    return area


def _color(color):
    return color.rgb if color.type == MSO_COLOR_TYPE.RGB else None


def lint_pptx(path: str | Path, *, whitespace=True):
    prs = Presentation(path)
    sw, sh = prs.slide_width / EMU, prs.slide_height / EMU
    report = {
        "passed": True,
        "errors": [],
        "slides": [],
        "minimum_font_pt": None,
        "whitespace_method": "union of semantic text/image/card regions within the safe content area; backgrounds excluded",
    }
    for index, slide in enumerate(prs.slides):
        errors = []

        def issue(code, shape=None):
            entry = {"slide": index + 1, "code": code}
            if shape is not None:
                entry["shape"] = shape.name
            errors.append(entry)

        boxes = []
        groups = [s for s in slide.shapes if s.shape_type == 6]
        shapes = [s for s in slide.shapes if s.shape_type != 6]
        for group in groups:
            shapes.extend(list(group.shapes))
        bg = _color(slide.background.fill.fore_color) if slide.background.fill.type == 1 else None
        for shape in shapes:
            box = rect(shape)
            x, y, w, h = box
            bleed = shape.name.startswith(("Layer:hero", "Layer:overlay", "Layer:background"))
            if not bleed and (x < 0.599 or y < 0.599 or x + w > sw - 0.599 or y + h > sh - 0.599):
                issue("outside_safe_margin", shape)
            if bleed and (x < -0.005 or y < -0.005 or x + w > sw + 0.005 or y + h > sh + 0.005):
                issue("outside_canvas", shape)
            if shape.is_placeholder and not shape.text.strip():
                issue("empty_placeholder", shape)
            if shape.shape_type == 13:
                try:
                    with Image.open(BytesIO(shape.image.blob)) as image:
                        image.verify()
                except Exception:
                    issue("unresolved_picture", shape)
            if "corner" in shape.name.casefold():
                issue("decorative_corner_icon", shape)
            if shape.has_text_frame and shape.text.strip():
                if any(
                    word in shape.text.casefold()
                    for word in (
                        "confirm before use",
                        "conceptual workflow details",
                        "[internal",
                        "prompt:",
                        "archetype:",
                    )
                ):
                    issue("internal_note_leak", shape)
                sizes = []
                total_height = 0
                for paragraph in shape.text_frame.paragraphs:
                    if not paragraph.text.strip():
                        continue
                    runs = [r for r in paragraph.runs if r.text]
                    point = max((r.font.size.pt if r.font.size else 18 for r in runs), default=18)
                    sizes += [r.font.size.pt if r.font.size else point for r in runs]
                    font = next((r.font.name for r in runs if r.font.name), "Calibri")
                    width = w - (shape.text_frame.margin_left + shape.text_frame.margin_right) / EMU
                    measured_w, measured_h, _ = measure_text(
                        paragraph.text, width, point, font, any(r.font.bold for r in runs)
                    )
                    if measured_w > width + 0.03:
                        issue("text_width_overflow", shape)
                    total_height += measured_h
                    background = bg
                    for outer in shapes[: shapes.index(shape)]:
                        if (
                            contains(rect(outer), box)
                            and getattr(outer, "fill", None) is not None
                            and outer.fill.type == 1
                        ):
                            background = _color(outer.fill.fore_color) or background
                    for run in runs:
                        fg = _color(run.font.color)
                        if (
                            fg is not None
                            and background is not None
                            and contrast(fg, background) < 4.5
                        ):
                            issue("contrast_below_4_5", shape)
                if sizes:
                    smallest = min(sizes)
                    report["minimum_font_pt"] = min(report["minimum_font_pt"] or smallest, smallest)
                    if smallest < 14:
                        issue("font_below_14", shape)
                if (
                    total_height
                    > h
                    - (shape.text_frame.margin_top + shape.text_frame.margin_bottom) / EMU
                    + 0.035
                ):
                    issue("text_height_overflow", shape)
                if len(shape.text.split()) > 45:
                    issue("region_over_45_words", shape)
            meaningful = (
                (shape.has_text_frame and bool(shape.text.strip()))
                or shape.shape_type == 13
                or shape.name.startswith("Layer:card")
            )
            if meaningful and not shape.name.startswith("Footer:"):
                # Clip full-bleed hero regions to the assessed safe content area.
                xx, yy = max(0.6, x), max(2.0, y)
                ww, hh = min(sw - 0.6, x + w) - xx, min(sh - 1.0, y + h) - yy
                if ww > 0 and hh > 0:
                    boxes.append((xx, yy, ww, hh))
        for group in groups:
            if group.name.startswith("Diagram:"):
                boxes.append(rect(group))
        cards = []
        for card in (s for s in shapes if s.name.startswith("Layer:card")):
            inner_height = max(0.01, rect(card)[3] - 0.5)
            filled_height = 0
            for child in shapes:
                if child == card or not contains(rect(card), rect(child)):
                    continue
                if child.shape_type == 13:
                    filled_height += rect(child)[3]
                elif child.has_text_frame and child.text.strip():
                    runs = [r for p in child.text_frame.paragraphs for r in p.runs if r.text]
                    point = max((r.font.size.pt if r.font.size else 18 for r in runs), default=18)
                    font = next((r.font.name for r in runs if r.font.name), "Calibri")
                    filled_height += measure_text(
                        child.text, rect(child)[2], point, font, any(r.font.bold for r in runs)
                    )[1]
            ratio = min(1, filled_height / inner_height)
            cards.append({"shape": card.name, "vertical_content_fill_ratio": round(ratio, 3)})
            if ratio < 0.7:
                issue("card_content_under_70_percent", card)
        archetype = slide._element.cSld.get("name")
        if archetype in {"cover", "split_image", "big_stat", "quote", "closing_cta"} and not any(
            s.shape_type == 13 and s.name.startswith("Layer:hero") for s in shapes
        ):
            issue("missing_hero_picture")
        for a, b in combinations(shapes, 2):
            ar, br = rect(a), rect(b)
            if intersection(ar, br) < 0.015:
                continue
            # Only explicitly intentional layer/containment relationships are exempt.
            if a.name.startswith(
                ("Layer:hero", "Layer:overlay", "Layer:background")
            ) or b.name.startswith(("Layer:hero", "Layer:overlay", "Layer:background")):
                continue
            if (a.name.startswith("Layer:card") and contains(ar, br)) or (
                b.name.startswith("Layer:card") and contains(br, ar)
            ):
                continue
            issue("shape_overlap", a)
        empty = max(0, 1 - union_area(boxes) / ((sw - 1.2) * (sh - 3.0)))
        if whitespace and empty > 0.45:
            issue("blank_content_over_45_percent")
        report["slides"].append(
            {
                "slide": index + 1,
                "blank_content_ratio": round(empty, 3),
                "errors": errors,
                "cards": cards,
            }
        )
        report["errors"].extend(errors)
    report["passed"] = not report["errors"]
    return report
