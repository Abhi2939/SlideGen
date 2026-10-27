"""Recompose the saved Blush draft through source checks without provider requests."""

import argparse
from pathlib import Path
import re
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation  # noqa: E402

from app.core.deck_schema import Deck  # noqa: E402
from app.pipeline.draft import planned_draft_model, _source_points  # noqa: E402
from app.qa.draft import build_draft  # noqa: E402
from scripts.blush_smoke import PROMPT  # noqa: E402


def replay_content(path, prompt):
    source = Presentation(path)
    sections, facts = [], []
    titles, subtitles = [], []
    for i, slide in enumerate(source.slides):
        title = next(s.text for s in slide.shapes if s.name == "Slide title")
        if i in {0, len(source.slides) - 1}:
            titles.append(title)
            subtitles.append(
                next(
                    s.text for s in slide.shapes if s.name == "Text" and s.has_text_frame and s.text
                )
            )
            continue
        points = [s.text for s in slide.shapes if s.name in {"Source context", "Source point"}]
        sections.append(
            {
                "title": title,
                "points": [piece for point in points for piece in _source_points(point)],
            }
        )
        factual = [
            point for point in points if not re.match(r"(?i)proposed\b|add verified\b", point)
        ]
        if factual:
            facts.append(" ".join(factual))
    return planned_draft_model(prompt).model_validate(
        {
            "title": titles[0],
            "subtitle": subtitles[0],
            "facts": facts,
            "sections": sections,
            "closing_title": titles[-1],
            "closing_text": subtitles[-1],
        }
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-pptx", type=Path, default=Path("output.pptx"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/qa/vivid-refresh"))
    parser.add_argument("--brief-file", type=Path)
    args = parser.parse_args()
    prompt = args.brief_file.read_text(encoding="utf-8") if args.brief_file else PROMPT
    content = replay_content(args.input_pptx, prompt)
    slides = [{"layout": "title", "title": content.title, "subtitle": content.subtitle}]
    slides += [
        {"layout": "draft_content", "title": section.title, "points": section.points}
        for section in content.sections
    ]
    slides += [
        {"layout": "closing", "title": content.closing_title, "subtitle": content.closing_text}
    ]
    deck = Deck(
        title=content.title, theme="studio-beauty", design_version="draft-v1", slides=slides
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "replayed-content.json").write_text(
        content.model_dump_json(indent=2), encoding="utf-8"
    )
    print(build_draft(deck, args.output_dir / "blush-rich.pptx"))
    print("Saved output replayed through source checks and new renderer; no provider requests.")


if __name__ == "__main__":
    main()
