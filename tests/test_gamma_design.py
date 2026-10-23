import json

from pptx import Presentation
from pptx.util import Inches

from app.qa.lint import lint_pptx
from app.render.common import add_text
from app.render.grid import GRID
from app.render.theme import THEMES, contrast, rgb, theme_for_topic


def test_curated_themes_are_distinct_and_accessible():
    names = ["beauty", "fintech", "health", "edtech", "climate", "bold-dark"]
    assert len({str(THEMES[name].bg) for name in names}) == 6
    for name in names:
        theme = THEMES[name]
        for text in (theme.text, theme.muted, theme.primary, theme.accent):
            assert contrast(text, theme.bg) >= 4.5
            assert contrast(text, theme.surface) >= 4.5
        assert contrast(theme.on_primary, theme.primary) >= 4.5
    assert theme_for_topic("Blush custom skincare") == "beauty"
    assert theme_for_topic("edtech product for students") == "edtech"
    assert theme_for_topic("payments fintech pitch") == "fintech"


def test_grid_safe_edges_and_card_columns():
    assert abs(GRID.x(0) - 0.6) < 0.001
    assert abs(GRID.x(0) + GRID.span(12) - (13.333 - 0.6)) < 0.001
    assert GRID.span(3) >= 2.4


def test_lint_reads_actual_bad_export(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    for x in (0.05, 0.1):
        add_text(
            slide,
            Inches(x),
            Inches(2),
            Inches(3),
            Inches(0.1),
            "Conceptual workflow details; confirm before use",
            size_pt=10,
            color=rgb("AAAAAA"),
            font="Calibri",
        )
    path = tmp_path / "bad.pptx"
    prs.save(path)
    result = lint_pptx(path)
    codes = {item["code"] for item in result["errors"]}
    assert {
        "outside_safe_margin",
        "empty_placeholder",
        "font_below_14",
        "text_height_overflow",
        "shape_overlap",
        "internal_note_leak",
        "blank_content_over_45_percent",
    } <= codes
    json.dumps(result)
