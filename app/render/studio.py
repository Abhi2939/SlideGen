"""Fuller, colorful draft compositions with editable figures and source-based charts."""

import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.render.common import new_slide
from app.render.data_charts import SourceChart, add_chart, extract_chart
from app.render.draft import Canvas as BaseCanvas, composition, _list
from app.render.measure import measure_text, TextNeedsSplit
from app.render.theme import contrast, rgb

WHITE = rgb("FFFFFF")


def mix(color, white=0.86):
    return rgb("".join(f"{round(channel * (1 - white) + 255 * white):02X}" for channel in color))


class Canvas(BaseCanvas):
    def color(self, index):
        return self.theme.chart_colors[index % len(self.theme.chart_colors)]

    def ink(self, background):
        return WHITE if contrast(WHITE, background) >= 4.5 else rgb("0D0A10")

    def fit(self, value, x, y, w, h, size=24, **kwargs):
        for point in range(size, 13, -2):
            measured_w, measured_h, _ = measure_text(
                value, w, point, self.theme.font_body, kwargs.get("bold", False)
            )
            if measured_w <= w + 0.01 and measured_h <= h:
                return self.text(value, x, y, w, h, point, **kwargs)
        raise TextNeedsSplit("source text needs a larger visual region")

    def tile(self, x, y, w, h, color, name):
        return self.panel(x, y, w, h, color, name, MSO_SHAPE.ROUNDED_RECTANGLE)

    def circle(self, cx, cy, diameter, color, name="figure"):
        return self.panel(
            cx - diameter / 2, cy - diameter / 2, diameter, diameter, color, name, MSO_SHAPE.OVAL
        )

    def gradient(self, x, y, w, h, first, last, name):
        shape = self.tile(x, y, w, h, first, name)
        shape.fill.gradient()
        shape.fill.gradient_angle = 35
        shape.fill.gradient_stops[0].color.rgb = first
        shape.fill.gradient_stops[1].color.rgb = last
        return shape

    def group(self, start, name):
        shape = self.slide.shapes.add_group_shape(list(self.slide.shapes)[start:])
        shape.name = "Diagram:" + name
        return shape

    def product(self, label, cx, y, scale=1.0, color=None):
        """Illustrative packaging built from editable shapes and custom geometry."""
        color = color or self.color(0)
        u = scale
        ink = self.theme.primary
        start = len(self.slide.shapes)

        def part(dx, dy, w, h, fill, name="product", kind=MSO_SHAPE.RECTANGLE):
            return self.panel(cx + dx * u, y + dy * u, w * u, h * u, fill, name, kind)

        key = label.casefold()
        if "foundation" in key:
            part(-0.48, 0.62, 0.96, 1.53, mix(color, 0.48), kind=MSO_SHAPE.ROUNDED_RECTANGLE)
            part(-0.3, 0.31, 0.6, 0.36, ink)
            part(-0.3, 0.17, 0.84, 0.14, ink)
            part(-0.15, 0.09, 0.36, 0.09, rgb("D6AD79"))
            part(-0.36, 1.19, 0.72, 0.5, mix(color, 0.94))
            part(-0.34, 0.78, 0.06, 1.13, mix(color, 0.8), "glass-highlight")
            self.line(
                cx - 0.22 * u, y + 1.37 * u, cx + 0.22 * u, y + 1.37 * u, color=color, width=2
            )
        elif "lipstick" in key:
            part(-0.3, 1.18, 0.6, 1.0, ink, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
            part(-0.27, 0.95, 0.54, 0.35, rgb("D6AD79"))
            part(-0.29, 1.25, 0.58, 0.08, color)
            # Slanted lipstick tip is native freeform geometry, not a raster.
            x, yy = cx - 0.21 * u, y + 0.12 * u
            b = self.slide.shapes.build_freeform(Inches(x), Inches(yy + 0.23 * u))
            b.add_line_segments(
                [
                    (Inches(x + 0.42 * u), Inches(yy)),
                    (Inches(x + 0.42 * u), Inches(y + 1.03 * u)),
                    (Inches(x), Inches(y + 1.03 * u)),
                ],
                close=True,
            )
            wax = b.convert_to_shape()
            wax.name = "Layer:background:lipstick-wax"
            wax.fill.solid()
            wax.fill.fore_color.rgb = color
            wax.line.fill.background()
            self.line(
                cx - 0.12 * u,
                y + 0.53 * u,
                cx - 0.12 * u,
                y + 0.95 * u,
                color=mix(color, 0.55),
                width=2,
            )
        else:
            part(-0.61, 1.24, 1.22, 0.91, mix(color, 0.52), kind=MSO_SHAPE.ROUNDED_RECTANGLE)
            part(-0.65, 1.04, 1.3, 0.27, ink, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
            part(-0.65, 0.99, 1.3, 0.17, mix(ink, 0.12), kind=MSO_SHAPE.OVAL)
            part(-0.51, 1.5, 1.02, 0.35, mix(color, 0.95))
            self.line(
                cx - 0.28 * u, y + 1.65 * u, cx + 0.28 * u, y + 1.65 * u, color=color, width=2
            )
        return self.group(start, "product:" + key)


def brief_groups(deck):
    if deck is None:
        return []
    source = " ".join(value for slide in deck.slides for value in getattr(slide, "points", []))
    groups = []
    for heading, pattern in (
        ("Products", r"\boffering\s+(.+?)(?:\s+blended\b|[.!?]|$)"),
        ("Launch touchpoints", r"\bfocus\s+on\s+(.+?)(?:,?\s+followed\s+by|[.!?]|$)"),
        ("Growth activities", r"\bgrowth\s+through\s+([^.!?]+)"),
    ):
        labels = _list(source, pattern)
        if labels:
            groups.append((heading, labels))
    # Short exact excerpts do not always include the introductory clause.
    for slide in deck.slides:
        if not hasattr(slide, "points"):
            continue
        kind, labels = composition(slide.points)
        heading = {
            "collection": "Products",
            "launch": "Launch touchpoints",
            "growth": "Growth activities",
        }.get(kind)
        if heading and labels and heading not in {group[0] for group in groups}:
            groups.append((heading, labels))
    return groups


def _source(c, points):
    value = " ".join(points)
    for size in (18, 17, 16):
        h = measure_text(value, 12.0, size, c.theme.font_body)[1] + 0.06
        if h <= 0.95 and len(value.split()) <= 45:
            c.text(value, 0.65, 1.48, 12.0, h, size, color=c.theme.muted, name="Source context")
            return max(2.25, 1.48 + h + 0.2)
    raise TextNeedsSplit("source context needs its own region")


def _header(c, title):
    c.fit(title, 0.65, 0.6, 12.0, 0.75, 34, bold=True, color=c.theme.primary, name="Slide title")


def _cover(c, data, groups, brand, topics):
    c.gradient(7.0, 0.6, 5.67, 5.7, mix(c.color(0), 0.79), mix(c.color(1), 0.76), "cover-stage")
    c.circle(10.0, 3.4, 4.15, mix(c.color(0), 0.9), "product-backdrop")
    c.panel(0.65, 1.25, 0.75, 0.08, c.color(0), "cover-accent")
    c.fit(
        data.title, 0.65, 1.65, 6.0, 1.65, 64, bold=True, color=c.theme.primary, name="Slide title"
    )
    c.fit(data.subtitle or "", 0.65, 3.45, 5.95, 1.35, 26, color=c.theme.muted)
    products = next((labels for name, labels in groups if name == "Products"), [])
    if {label.casefold() for label in products} == {"foundation", "lipstick", "skincare"}:
        c.product("Foundation", 9.1, 1.9, 1.37, c.color(3))
        c.product("Lipstick", 10.9, 2.2, 1.2, c.color(0))
        c.product("Skincare", 11.0, 3.25, 0.95, c.color(2))
        c.text(
            brand[:30],
            7.5,
            0.98,
            4.6,
            0.5,
            24,
            bold=True,
            color=c.theme.primary,
            name="Figure brand",
        )
        c.fit(
            "Foundation, lipstick and skincare",
            7.5,
            5.58,
            4.7,
            0.45,
            18,
            color=c.theme.primary,
            name="Collection label",
        )
    elif groups:
        for i, (name, labels) in enumerate(groups[:3]):
            yy = 1.25 + i * 1.5
            c.circle(7.95, yy + 0.35, 0.62, c.color(i), "cover-topic")
            c.text(
                str(len(labels)),
                7.7,
                yy + 0.12,
                0.5,
                0.4,
                22,
                bold=True,
                color=c.ink(c.color(i)),
                center=True,
            )
            c.fit(name, 8.6, yy, 3.4, 0.5, 22, bold=True, color=c.theme.primary)
            c.fit(", ".join(labels), 8.6, yy + 0.6, 3.5, 0.7, 16, color=c.theme.muted)
    else:
        for i, title in enumerate(topics[:3]):
            yy, color = 1.35 + i * 1.52, c.color(i)
            c.tile(7.5, yy, 4.7, 1.24, mix(color, 0.89), "topic-figure")
            c.circle(8.12, yy + 0.62, 0.75, color, "topic-node")
            c.icon(title, 8.12, yy + 0.62, 0.37, color=c.ink(color))
            c.fit(
                title,
                8.85,
                yy + 0.22,
                3.08,
                0.85,
                23,
                bold=True,
                color=c.theme.primary,
                name="Source topic",
            )
    for i, (heading, labels) in enumerate(groups[:3]):
        x = 0.65 + i * 2.0
        c.tile(x, 5.15, 1.82, 1.03, mix(c.color(i), 0.87), "cover-scope")
        c.text(
            str(len(labels)),
            x + 0.15,
            5.28,
            0.5,
            0.5,
            30,
            bold=True,
            color=c.theme.primary,
            name="Source list count",
        )
        c.fit(heading, x + 0.72, 5.35, 0.98, 0.65, 16, color=c.theme.primary, name="Count category")


def _collection(c, labels, top):
    n = len(labels)
    if n > 3:
        for i, label in enumerate(labels):
            x, yy, color = 0.65 + (i % 3) * 4.0, top + (i // 3) * 1.94, c.color(i)
            c.tile(x, yy, 3.76, 1.72, mix(color, 0.84), "collection-panel")
            c.circle(x + 0.67, yy + 0.85, 0.85, WHITE, "collection-icon")
            c.icon(label, x + 0.67, yy + 0.85, 0.48, color=color)
            c.fit(
                label[0].upper() + label[1:],
                x + 1.3,
                yy + 0.35,
                2.2,
                1.15,
                23,
                bold=True,
                color=c.theme.primary,
                name="Collection label",
            )
        return
    step = 12.0 / n
    h = 6.1 - top
    for i, label in enumerate(labels):
        x, w, color = 0.65 + i * step, step - 0.22, c.color(i)
        c.tile(x, top, w, h, mix(color, 0.83), "collection-panel")
        c.circle(x + 0.48, top + 0.43, 0.5, color, "collection-index")
        c.text(
            str(i + 1),
            x + 0.25,
            top + 0.26,
            0.46,
            0.32,
            18,
            bold=True,
            color=c.ink(color),
            center=True,
        )
        if label.casefold() in {"foundation", "lipstick", "skincare"}:
            c.product(label, x + w / 2, top + 0.5, min(1.13, (h - 1.18) / 2.18), color)
        else:
            c.icon(label, x + w / 2, top + 1.55, 0.95, color=color)
        c.fit(
            label[0].upper() + label[1:],
            x + 0.18,
            5.39,
            w - 0.36,
            0.65,
            26,
            bold=True,
            color=c.theme.primary,
            center=True,
            name="Collection label",
        )


def _dimension_icon(c, index, cx, cy, scale=1.0):
    if index == 2:
        for angle, color in ((-25, c.color(3)), (0, c.color(1)), (25, c.color(0))):
            shape = c.tile(
                cx - 0.14 * scale,
                cy - 0.35 * scale,
                0.28 * scale,
                0.7 * scale,
                color,
                "style-swatch",
            )
            shape.rotation = angle
    else:
        colors = ("E0C1A4", "B78970", "795241") if index == 0 else ("DBA280", "D8CDB5", "BFA2BC")
        for j, value in enumerate(colors):
            c.circle(cx + (j - 1) * 0.25 * scale, cy, 0.4 * scale, rgb(value), "tone-swatch")


def _personalization(c, labels, top):
    start = len(c.slide.shapes)
    c.tile(0.65, top, 6.15, 6.1 - top, mix(c.color(1), 0.9), "blend-diagram")
    cx, cy = 3.7, 4.15
    nodes = ((1.55, 3.15), (5.85, 3.15), (3.7, 5.55))
    for i, (xx, yy) in enumerate(nodes):
        c.line(cx, cy, xx, yy, color=mix(c.color(i), 0.45), width=3)
    c.circle(cx, cy, 1.68, c.theme.primary, "custom-blend")
    c.text(
        "Your blend",
        cx - 0.75,
        cy - 0.38,
        1.5,
        0.85,
        24,
        bold=True,
        color=c.theme.on_primary,
        center=True,
        name="Blend hub",
    )
    for i, ((xx, yy), label) in enumerate(zip(nodes, labels)):
        c.circle(xx, yy, 0.96, WHITE, "dimension-node")
        _dimension_icon(c, i, xx, yy, 1.0)
    c.group(start, "personalization-map")
    for i, label in enumerate(labels):
        x, yy = 7.1, top + i * 1.3
        c.tile(x, yy, 5.55, 1.1, mix(c.color(i), 0.84), "dimension-panel")
        c.circle(x + 0.58, yy + 0.55, 0.7, WHITE, "dimension-icon")
        _dimension_icon(c, i, x + 0.58, yy + 0.55, 0.65)
        c.fit(
            label[0].upper() + label[1:],
            x + 1.2,
            yy + 0.31,
            4.0,
            0.65,
            27,
            bold=True,
            color=c.theme.primary,
            name="Personalization label",
        )


def _launch_illustration(c, label, x, y, w, color):
    c.tile(x, y, w, 1.55, WHITE, "launch-illustration")
    key = label.casefold()
    if "quiz" in key:
        for i in range(3):
            yy = y + 0.32 + i * 0.4
            c.circle(x + 0.35, yy + 0.03, 0.16, mix(color, 0.55), "quiz-choice")
            c.line(x + 0.62, yy, x + w - 0.28 - (i % 2) * 0.35, yy, color=mix(color, 0.4), width=4)
    elif "waitlist" in key:
        for i in range(3):
            yy = y + 0.3 + i * 0.42
            c.circle(x + 0.38, yy + 0.08, 0.24, mix(color, 0.6), "waitlist-avatar")
            c.line(x + 0.73, yy + 0.08, x + w - 0.3, yy + 0.08, color=mix(color, 0.4), width=4)
    else:
        c.tile(x + 0.16, y + 0.15, w - 0.32, 1.25, mix(color, 0.88), "demo-screen")
        c.circle(x + w / 2, y + 0.78, 0.7, color, "demo-play")
        triangle = c.panel(
            x + w / 2 - 0.12,
            y + 0.61,
            0.3,
            0.33,
            c.ink(color),
            "play-symbol",
            MSO_SHAPE.ISOSCELES_TRIANGLE,
        )
        triangle.rotation = 90


def _launch(c, labels, top):
    step = 12.0 / len(labels)
    for i, label in enumerate(labels):
        x, w, color = 0.65 + i * step, step - 0.24, c.color(i)
        c.tile(x, top, w, 6.1 - top, mix(color, 0.83), "launch-panel")
        c.circle(x + 0.5, top + 0.46, 0.55, color, "launch-stage")
        c.text(
            str(i + 1),
            x + 0.25,
            top + 0.27,
            0.5,
            0.38,
            20,
            bold=True,
            color=c.ink(color),
            center=True,
            name="Sequence number",
        )
        _launch_illustration(c, label, x + 0.27, top + 1.02, w - 0.54, color)
        display = re.sub(r"^(?:a|an|the)\s+", "", label, flags=re.I)
        c.fit(
            display[0].upper() + display[1:],
            x + 0.22,
            5.28,
            w - 0.44,
            0.72,
            24,
            bold=True,
            color=c.theme.primary,
            name="Workflow activity",
        )


def _growth(c, labels, top):
    start = len(c.slide.shapes)
    step = 12.0 / len(labels)
    cy = 4.25
    for i in range(len(labels) - 1):
        c.line(0.65 + (i + 0.5) * step, cy, 0.65 + (i + 1.5) * step, cy, color=c.color(i), width=4)
    for i, label in enumerate(labels):
        cx, color = 0.65 + (i + 0.5) * step, c.color(i)
        c.circle(cx, cy, 0.96, color, "growth-node")
        c.icon(label, cx, cy, 0.43, color=c.ink(color))
        c.circle(cx + 0.35, cy - 0.37, 0.34, c.theme.primary, "growth-order")
        c.text(
            str(i + 1),
            cx + 0.2,
            cy - 0.5,
            0.3,
            0.25,
            14,
            color=c.theme.on_primary,
            center=True,
            name="Sequence number",
        )
        yy = top if i % 2 == 0 else 5.0
        h, w = 1.18, step - 0.14
        c.tile(cx - w / 2, yy, w, h, mix(color, 0.82), "growth-activity")
        c.panel(cx - w / 2 + 0.12, yy + 0.12, w - 0.24, 0.045, color, "growth-rule")
        label_y = yy + 0.27
        c.fit(
            label[0].upper() + label[1:],
            cx - w / 2 + 0.13,
            label_y,
            w - 0.26,
            0.83,
            21,
            bold=True,
            color=c.theme.primary,
            center=True,
            name="Workflow activity",
        )
        stem_start, stem_end = (yy + h + 0.05, cy - 0.52) if i % 2 == 0 else (cy + 0.52, yy - 0.05)
        c.line(cx, stem_start, cx, stem_end, color=mix(color, 0.3), width=2)
    c.group(start, "growth-roadmap")


def _keywords(text):
    stop = {
        "about",
        "after",
        "before",
        "their",
        "there",
        "these",
        "those",
        "through",
        "which",
        "would",
        "should",
        "could",
        "proposed",
        "supplied",
        "details",
        "information",
        "verified",
        "presentation",
        "discuss",
        "point",
        "with",
        "from",
        "your",
        "that",
        "this",
        "next",
        "step",
    }
    words = re.findall(r"[A-Za-z][A-Za-z-]{3,17}", text)
    return list(dict.fromkeys(word.casefold() for word in words if word.casefold() not in stop))[:3]


def _points(c, points):
    if len(points) == 1:
        c.tile(0.65, 2.15, 7.45, 3.95, mix(c.color(0), 0.87), "statement-panel")
        c.fit(
            points[0], 1.0, 2.65, 6.75, 2.95, 34, bold=True, color=c.theme.primary, name="Statement"
        )
        terms = _keywords(points[0])
        for i, term in enumerate(terms):
            yy, color = 2.3 + i * 1.3, c.color(i + 1)
            c.circle(9.0, yy + 0.4, 0.76, color, "idea-node")
            c.icon(term, 9.0, yy + 0.4, 0.36, color=c.ink(color))
            c.fit(
                term.capitalize(),
                9.65,
                yy + 0.1,
                2.9,
                0.75,
                23,
                bold=True,
                color=c.theme.primary,
                name="Source keyword",
            )
        return
    cols = 2 if len(points) in {2, 4} else 3
    rows = (len(points) + cols - 1) // cols
    step, rh = 12.0 / cols, 4.0 / rows
    for i, point in enumerate(points):
        x, yy, w = 0.65 + (i % cols) * step, 2.12 + (i // cols) * rh, step - 0.22
        color = c.color(i)
        c.tile(x, yy, w, rh - 0.2, mix(color, 0.84), "source-panel")
        c.circle(x + 0.45, yy + 0.42, 0.48, color, "source-index")
        c.text(
            str(i + 1),
            x + 0.22,
            yy + 0.25,
            0.46,
            0.32,
            18,
            bold=True,
            color=c.ink(color),
            center=True,
        )
        c.fit(
            point,
            x + 0.26,
            yy + 0.9,
            w - 0.52,
            rh - 1.18,
            26 if rows == 1 else 20,
            color=c.theme.primary,
            name="Source point",
        )


def _data(c, points, chart):
    top = _source(c, points)
    if (
        len(chart.labels) > 3
        or any(len(label) > 20 for label in chart.labels)
        or any(len(f"{value:g}{chart.unit}") > 8 for value in chart.values)
    ):
        c.tile(0.65, top, 12.0, 6.12 - top, WHITE, "chart-panel")
        add_chart(c.slide, chart, c.theme, (0.85, top + 0.1, 11.6, 6.0 - top))
        return
    c.tile(0.65, top, 8.0, 6.12 - top, WHITE, "chart-panel")
    add_chart(c.slide, chart, c.theme, (0.85, top + 0.2, 7.55, 5.95 - top))
    rh = min(1.15, (6.12 - top) / len(chart.labels))
    for i, (label, value) in enumerate(zip(chart.labels, chart.values)):
        yy, color = top + i * rh, c.color(i)
        c.tile(8.9, yy, 3.75, rh - 0.12, mix(color, 0.86), "data-summary")
        c.fit(
            label,
            9.12,
            yy + 0.13,
            2.35,
            rh - 0.27,
            20,
            bold=True,
            color=c.theme.primary,
            name="Data category",
        )
        unit = chart.unit if chart.unit == "%" else " " + chart.unit if chart.unit else ""
        c.fit(
            f"{value:g}{unit}",
            11.5,
            yy + 0.13,
            0.95,
            rh - 0.27,
            22,
            bold=True,
            color=c.theme.primary,
            name="Supplied value",
        )


def _closing(c, data, groups, topics):
    c.slide.background.fill.fore_color.rgb = c.theme.primary
    c.gradient(0, 0, 13.333, 7.5, c.theme.primary, rgb("614279"), "closing-backdrop")
    c.panel(0.65, 1.2, 0.8, 0.08, mix(c.color(0), 0.4), "closing-accent")
    c.fit(
        data.title,
        0.65,
        1.72,
        6.4,
        1.65,
        46,
        bold=True,
        color=c.theme.on_primary,
        name="Slide title",
    )
    c.fit(data.subtitle or "", 0.65, 3.7, 6.0, 1.6, 28, color=c.theme.on_primary)
    if groups:
        c.tile(7.15, 1.0, 5.5, 5.1, WHITE, "scope-chart")
        c.text(
            "Items in the brief",
            7.5,
            1.35,
            4.8,
            0.45,
            22,
            bold=True,
            color=c.theme.primary,
            name="Chart explanation",
        )
        counts = SourceChart(
            tuple(name for name, _ in groups),
            tuple(float(len(values)) for _, values in groups),
            "items",
        )
        add_chart(c.slide, counts, c.theme, (7.45, 2.1, 4.8, 3.25), counts=True)
        c.text(
            "Counts of the supplied products and activities",
            7.5,
            5.45,
            4.8,
            0.5,
            14,
            color=c.theme.muted,
            name="Chart source",
        )
    else:
        for i, title in enumerate(topics[:3]):
            yy, color = 1.45 + i * 1.52, c.color(i)
            c.tile(7.35, yy, 5.1, 1.24, mix(color, 0.85), "closing-topic")
            c.circle(7.95, yy + 0.62, 0.7, color, "closing-topic-node")
            c.icon(title, 7.95, yy + 0.62, 0.34, color=c.ink(color))
            c.fit(
                title,
                8.55,
                yy + 0.2,
                3.5,
                0.85,
                23,
                bold=True,
                color=c.theme.primary,
                name="Source topic",
            )


def render(prs, data, theme, *, brand, number, deck=None):
    slide = new_slide(prs, theme)
    c = Canvas(slide, theme)
    groups = brief_groups(deck)
    topics = [slide.title for slide in deck.slides[1:-1]] if deck else [data.title]
    if data.layout == "title":
        kind = "cover"
        _cover(c, data, groups, brand, topics)
    elif data.layout == "closing":
        kind = "closing"
        _closing(c, data, groups, topics)
    else:
        chart = extract_chart(data.points)
        kind, labels = composition(data.points)
        if kind == "personalization" and len(labels) != 3:
            kind = "points"
        _header(c, data.title)
        if chart:
            kind = "data_chart"
            _data(c, data.points, chart)
        elif kind == "points":
            _points(c, data.points)
        else:
            top = _source(c, data.points)
            {
                "collection": _collection,
                "personalization": _personalization,
                "launch": _launch,
                "growth": _growth,
            }[kind](c, labels, top)
    slide._element.cSld.set("name", "studio_" + kind)
    footer_color = theme.on_primary if data.layout == "closing" else theme.muted
    c.line(0.65, 6.45, 12.65, 6.45, color=mix(theme.primary, 0.78), width=0.8)
    c.text(brand[:45], 0.65, 6.6, 6.5, 0.26, 14, color=footer_color, name="Footer:brand")
    c.text(
        f"Draft  |  {number:02d}",
        10.75,
        6.6,
        1.9,
        0.26,
        14,
        color=footer_color,
        name="Footer:number",
    )
    return slide
