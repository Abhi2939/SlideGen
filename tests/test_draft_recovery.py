import json
from io import BytesIO

import pytest
from pptx import Presentation
from pydantic import ValidationError
from fastapi.testclient import TestClient

from app.core.errors import DraftContentError, LLMProviderError, StructuredOutputError
from app.pipeline.draft import (
    DraftSection,
    _source_points,
    generate_draft_deck,
    planned_draft_model,
)
from app.pipeline.draft_recovery import source_only_content
from app.pipeline.grounding import is_supported_paraphrase
from app.qa.draft import build_draft
from app.qa.lint import lint_pptx
from app.jobs.store import JobStore
from app.main import create_app
from test_blush_prompt import FACT_TEXTS
from test_draft_delivery import content, generator, wait_status
from test_editorial_pipeline import PROMPT


def test_complete_source_sentences_can_span_section_boundaries():
    data = content()
    data["facts"] = [PROMPT.split(". ", 1)[0] + ".", PROMPT.split(". ", 1)[1]]
    result = planned_draft_model(PROMPT).model_validate(data)
    assert [section.title for section in result.sections] == [
        "Product range",
        "Personalization",
        "Launch plan",
        "Growth plan",
    ]
    assert "subscriptions" in " ".join(result.sections[-1].points)
    assert is_supported_paraphrase(
        "Subscriptions and new product lines", "subscriptions, new product lines"
    )
    assert not is_supported_paraphrase(
        "Subscriptions reached 1,000 customers", "subscriptions, new product lines"
    )


@pytest.mark.parametrize(
    "point",
    [
        "Blush is a custom beauty brand offering foundation, lipstick, and skincare blended exclusively for your skin tone, undertone, and style, so you never get the wrong shade.",
        " ".join(["a"] * 31),
    ],
)
def test_point_reflow_preserves_every_word_within_existing_slide_limits(point):
    result = DraftSection.model_validate({"title": "Source", "points": [point]})
    assert len(result.points) >= 2
    assert " ".join(result.points) == " ".join(point.split())
    assert all(len(value) <= 160 and len(value.split()) <= 30 for value in result.points)


def test_reflow_does_not_drop_copy_that_exceeds_capacity():
    with pytest.raises(ValidationError):
        DraftSection.model_validate({"title": "Source", "points": ["Long source wording " * 90]})


def test_long_generated_point_is_reflowed_without_another_provider_request():
    data = content()
    data["sections"][0]["points"] = [PROMPT.split(". ", 1)[0] + "."]
    fake, provider = generator([json.dumps(data)])
    deck = generate_draft_deck(PROMPT, generator=fake)
    assert len(deck.slides) == 6
    assert len(provider.calls) == 1


@pytest.mark.parametrize("count", [3, 6, 10])
def test_rejected_response_recovers_a_valid_editable_company_deck(tmp_path, count):
    prompt = PROMPT + f" Make a {count}-slide deck."
    fake, provider = generator(["not valid JSON"] * 2)
    deck = generate_draft_deck(prompt, generator=fake)
    assert len(provider.calls) == 2
    assert len(deck.slides) == count
    assert all("Source-only draft" in slide.speaker_notes for slide in deck.slides)
    path = build_draft(deck, tmp_path / "recovered.pptx")
    prs = Presentation(path)
    visible = " ".join(
        shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
    ).casefold()
    assert all(fact.casefold() in visible for fact in FACT_TEXTS)
    assert not any(
        "not valid JSON" in shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert lint_pptx(path, whitespace=False)["errors"] == []


def test_general_topic_has_no_company_source_recovery():
    fake, provider = generator(["not valid JSON"] * 2)
    with pytest.raises(DraftContentError):
        generate_draft_deck("Make a deck about dogs", generator=fake)
    assert len(provider.calls) == 2


def test_source_only_recovery_does_not_truncate_an_oversized_brief():
    prompt = "Acme is a company. " + "Our launch includes all the supplied details. " * 160
    assert source_only_content(prompt, planned_draft_model(prompt)) is None


def test_source_recovery_keeps_lowercase_presentation_directions_out_of_facts():
    prompt = PROMPT + " make a 6-slide deck."
    result = source_only_content(prompt, planned_draft_model(prompt))
    assert result is not None
    assert not any("make a" in fact for fact in result.facts)


def test_source_only_recovery_keeps_financial_numbers_from_another_company(tmp_path):
    prompt = "Acme Pay is a fintech company offering invoice tools. Marketing: 40%, Product: 35%, Operations: 25%. Make a 3-slide deck."
    fake, provider = generator(["{}"] * 2)
    deck = generate_draft_deck(prompt, generator=fake)
    visible = " ".join(point for slide in deck.slides for point in getattr(slide, "points", []))
    assert all(value in visible for value in ("40%", "35%", "25%"))
    assert len(provider.calls) == 2
    build_draft(deck, tmp_path / "finance.pptx")


def test_provider_access_failures_do_not_masquerade_as_source_recovery():
    fake, provider = generator([LLMProviderError("Access denied", status_code=401)])
    with pytest.raises(StructuredOutputError) as caught:
        generate_draft_deck(PROMPT, generator=fake)
    assert not isinstance(caught.value, DraftContentError)
    assert len(provider.calls) == 1


def test_source_chunking_is_lossless_for_supported_word_lengths():
    source = "All source facts, including subscriptions and product lines, remain visible."
    assert " ".join(_source_points(source, max_chars=30, max_words=6)) == source


def test_default_http_job_recovers_after_both_ai_responses_fail_validation(monkeypatch, tmp_path):
    base, provider = generator(["{}", "{}"])
    monkeypatch.setattr("app.pipeline.draft.create_llm_client", lambda *_: base.client)
    monkeypatch.setenv("SLIDEGEN_DELIVERY_MODE", "draft")
    store = JobStore(tmp_path)
    with TestClient(create_app(store=store)) as client:
        response = client.post("/api/jobs", json={"prompt": PROMPT})
        assert response.status_code == 202
        job_id = response.json()["id"]
        job = wait_status(client, job_id, "done")
        assert job["failure"] is None
        download = client.get(f"/api/jobs/{job_id}/download")
        assert download.status_code == 200
        deck = Presentation(BytesIO(download.content))
        assert len(deck.slides) == 6
        assert all(
            "Source-only draft" in slide.notes_slide.notes_text_frame.text for slide in deck.slides
        )
        assert client.get(f"/debug/{job_id}").json()["file_checks"]["errors"] == []
    assert len(provider.calls) == 2
