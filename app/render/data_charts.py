"""Editable native charts from explicit label/value pairs, never inferred metrics."""

from dataclasses import dataclass
from decimal import Decimal
import re

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.util import Inches, Pt

from app.render.theme import rgb


@dataclass(frozen=True)
class SourceChart:
    labels: tuple[str, ...]
    values: tuple[float, ...]
    unit: str


def extract_chart(points):
    """Require a homogeneous set of complete, explicitly labelled values."""
    source = "\n".join(points)
    # Do not split thousands separators, and do not turn narrative numbers into data.
    parts = re.split(r";|\n|,(?=\s*[A-Za-z])", source)
    pattern = r"\s*([A-Za-z][A-Za-z0-9 &'()/.-]{0,39})\s*[:=]\s*(-?\d[\d,]*(?:\.\d+)?)\s*(%|USD|INR|users|customers|units|items)?[.]?\s*"
    matches = [re.fullmatch(pattern, part, re.I) for part in parts]
    if not 2 <= len(matches) <= 8 or not all(matches):
        return None
    labels = tuple(m[1].strip() for m in matches)
    units = {m[3].casefold() if m[3] else "" for m in matches}
    if len(units) != 1 or len(set(label.casefold() for label in labels)) != len(labels):
        return None
    values = tuple(float(Decimal(m[2].replace(",", ""))) for m in matches)
    if not all(abs(value) < 1e12 for value in values):
        return None
    return SourceChart(labels, values, units.pop())


def add_chart(slide, source, theme, box, *, counts=False):
    x, y, w, h = box
    data = CategoryChartData()
    data.categories = source.labels
    data.add_series(
        "Items in the brief" if counts else (source.unit or "Supplied values"), source.values
    )
    shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, *map(Inches, (x, y, w, h)), data)
    shape.name = "Chart:brief-counts" if counts else "Chart:supplied-data"
    chart = shape.chart
    chart.has_legend = False
    chart.has_title = False
    chart.chart_style = 10
    chart.font.name = theme.font_body
    chart.font.size = Pt(14)
    chart.font.color.rgb = theme.text
    plot = chart.plots[0]
    plot.gap_width = 65
    plot.vary_by_categories = True
    for i, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = theme.chart_colors[i % len(theme.chart_colors)]
        point.format.line.fill.background()
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.name, labels.font.size = theme.font_body, Pt(16)
    labels.font.color.rgb = theme.text
    labels.number_format = '0.##"%"' if source.unit == "%" else "0.##"
    chart.category_axis.tick_labels.font.name = theme.font_body
    chart.category_axis.tick_labels.font.size = Pt(14)
    chart.category_axis.tick_labels.font.color.rgb = theme.text
    chart.category_axis.reverse_order = True
    for axis in (chart.category_axis, chart.value_axis):
        axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.fill.background()
    chart.value_axis.tick_labels.font.size = Pt(14)
    chart.value_axis.tick_labels.font.color.rgb = theme.muted
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb("E6DEE8")
    if all(value >= 0 for value in source.values):
        chart.value_axis.minimum_scale = 0
    return shape
