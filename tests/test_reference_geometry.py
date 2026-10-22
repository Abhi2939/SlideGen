"""Reference layouts at small, typical, and maximum fixture sizes, offline."""

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.core.deck_schema import Deck
from app.render.builder import build_deck
from app.render.geometry import check_slide_geometry


FIXTURE = Path(__file__).parent / "fixtures/decks/reference_layouts.json"


def reference_cases():
    slides = json.loads(FIXTURE.read_text(encoding="utf-8"))["slides"]
    yield from (
        s
        for s in slides
        if s["layout"] in {"ask_allocation", "market_snapshot", "flow_steps", "quadrants"}
    )
    for count in (2, 4, 7):
        yield {
            "layout": "ask_allocation",
            "title": "Allocation " * 7,
            "ask": "Rs. 14 Cr" if count < 7 else "Funding amount to be supplied",
            "center_label": "Allocation of the requested funding",
            "items": [
                {"label": "Contingency and other indirect expenses", "pct": None}
                for _ in range(count)
            ],
        }
    yield {
        "layout": "ask_allocation",
        "title": "Concentrated allocation",
        "ask": "Rs. 14 Cr",
        "items": [
            {"label": "Primary allocation", "pct": 98.0},
            *[{"label": "Other indirect expenses and contingency", "pct": 1 / 3} for _ in range(6)],
        ],
    }
    for count in (0, 6, 7):
        yield {
            "layout": "market_snapshot",
            "title": "Market snapshot",
            "start_value": None,
            "end_value": "USD 652.80 Bn",
            "segments": [
                {
                    "category": "Interface and service channels",
                    "lead_text": "Digital payments lead the supplied segment comparison",
                    "lead_value": "46.2%",
                    "next_text": "Next segment for a sourced growth comparison",
                    "next_value": "CAGR 18.7%",
                }
                for _ in range(4)
            ],
            "table_headers": [
                "Driver",
                "(+) Impact on CAGR forecast",
                "Geographic relevance",
                "Impact timeline",
            ]
            if count
            else [],
            "table_rows": [
                [
                    "Open banking and API standardization",
                    "+2.5%",
                    "North America, Europe, Asia-Pacific",
                    term,
                ]
                for term in (
                    "Short term",
                    "Medium term",
                    "Long term",
                    "Medium term",
                    "Long term",
                    "Short term",
                    "Long term",
                )[:count]
            ],
        }
    for count in (3, 5, 8):
        yield {
            "layout": "flow_steps",
            "title": "How it works",
            "tagline": "A conceptual workflow",
            "subtitle": "Confirm the operational details",
            "steps": [
                {
                    "title": "Review information and confirm your preferences",
                    "detail": "Consider the available choices and record the information needed to continue. Confirm the next step with the team.",
                }
                for _ in range(count)
            ],
        }
    for count in range(10):
        yield {
            "layout": "quadrants",
            "title": "Competitive landscape",
            "quadrants": [
                {
                    "heading": "Biometric and physical payment services",
                    "names": [
                        "International Digital Payments Network " + chr(65 + i)
                        for i in range(count)
                    ],
                }
                for _ in range(4)
            ],
        }


@pytest.mark.parametrize("content", list(reference_cases()), ids=lambda s: s["layout"])
def test_reference_content_stays_in_bounds_without_text_overlap(content, tmp_path):
    deck = Deck.model_validate(
        {
            "title": "Geometry",
            "slides": [
                {"layout": "title", "title": "Geometry"},
                content,
                {"layout": "closing", "title": "Close"},
            ],
        }
    )
    prs = Presentation(build_deck(deck, tmp_path / "check.pptx", strict=True))
    assert check_slide_geometry(prs.slides[1], prs.slide_width, prs.slide_height) == []


def test_geometry_reports_bounds_and_overlapping_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for x in (-0.1, 0.2):
        box = slide.shapes.add_textbox(Inches(x), Inches(1), Inches(2), Inches(1))
        box.text = "Text"
    errors = check_slide_geometry(slide, prs.slide_width, prs.slide_height)
    assert any("out of bounds" in e for e in errors)
    assert any("text overlap" in e for e in errors)
