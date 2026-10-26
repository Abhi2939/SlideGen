import json

import pytest
from pydantic import ValidationError
from pptx import Presentation

from app.core.deck_schema import Deck, StorySlide
from app.core.errors import (
    DeckReviewError,
    LLMProviderError,
    OutlineSourceError,
    SlideWriteError,
    StructuredOutputError,
)
from app.llm.client import LLMClient
from app.llm.schema import recovery_format
from app.llm.structured import StructuredGenerator
from app.pipeline.editorial import (
    ITEM_BOUNDS,
    StoryPlan,
    StoryOutline,
    normalize_outline,
    planned_slide_model,
    planned_review_model,
    source_copy_fallback,
    validate_copy,
    visible_copy,
)
from app.pipeline.orchestrator import generate_deck
from test_editorial_pipeline import FIXTURE, PROMPT, StoryGenerator
from test_llm_structured import FakeProvider, parse_error


@pytest.mark.parametrize("archetype,bounds", ITEM_BOUNDS.items())
def test_every_planned_archetype_has_exact_decoding_and_runtime_contract(archetype, bounds):
    model = planned_slide_model(archetype)
    prepared = recovery_format("openai/gpt-oss-120b", model.__name__, model.model_json_schema())
    schema = prepared["json_schema"]["schema"]
    assert schema["properties"]["archetype"]["enum"] == [archetype]
    assert schema["properties"]["items"]["minItems"] == bounds[0]
    assert schema["properties"]["items"]["maxItems"] == bounds[1]
    assert "items" in schema["required"]
    invalid_count = bounds[1] + 1
    with pytest.raises(ValidationError):
        model.model_validate(
            {
                "archetype": archetype,
                "title": "Title",
                "subtitle": "Tagline",
                "items": [{"title": "Item", "description": "Description"}] * invalid_count,
            }
        )


def test_slide_three_wrong_count_is_repaired_in_the_real_structured_loop():
    data = Deck.model_validate_json(FIXTURE.read_text()).slides[2].model_dump()
    data["items"] = data["items"][:2]
    fixed = Deck.model_validate_json(FIXTURE.read_text()).slides[2].model_dump()
    provider = FakeProvider([json.dumps(data), json.dumps(fixed)])
    result = StructuredGenerator(LLMClient(provider)).generate(
        planned_slide_model("split_image"),
        stage="write",
        model="openai/gpt-oss-120b",
        system_prompt="Write the three source dimensions",
        user_prompt=PROMPT,
        max_tokens=1800,
    )
    assert len(result.items) == 3
    assert len(provider.calls) == 2
    assert "at least 3 items" in provider.calls[1][0][-1]["content"]
    # The general Deck/PPTX contract remains unchanged.
    assert StorySlide.model_validate(result.model_dump()).archetype == "split_image"


class BrokenThirdSlide(StoryGenerator):
    def __init__(self, count, *, failure=None):
        super().__init__()
        self.count = count
        self.failure = failure
        self.third_requests = []

    def generate(self, output_model, *, stage, **kwargs):
        if stage == "write":
            request = json.loads(kwargs["user_prompt"])
            if request["plan"]["purpose"].startswith("Explain the source beauty proposition"):
                self.third_requests.append(kwargs)
                if len(self.third_requests) == 1:
                    self.original = super().generate(output_model, stage=stage, **kwargs)
                if self.failure:
                    raise self.failure
                items = [self.original.items[0].model_copy(deep=True) for _ in range(self.count)]
                return self.original.model_copy(update={"items": items})
        return super().generate(output_model, stage=stage, **kwargs)


@pytest.mark.parametrize("count", [0, 1, 2, 4, 5, 6])
def test_production_slide_three_failure_recovers_and_renders_whole_deck(
    count, tmp_path, monkeypatch
):
    from app.render.builder import build_deck
    from app.qa.lint import lint_pptx

    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = BrokenThirdSlide(count)
    deck = generate_deck(PROMPT, generator=fake, semantic_attempts=1)
    assert len(fake.third_requests) == 1
    assert fake.stages[-1] == "review"
    slide = deck.slides[2]
    assert [i.title for i in slide.items] == ["Skin tone", "Undertone", "Individual style"]
    request = json.loads(fake.third_requests[0]["user_prompt"])
    assert request["structure"] == {"archetype": "split_image", "min_items": 3, "max_items": 3}
    path = build_deck(deck, tmp_path / "recovered.pptx", strict=True)
    assert len(Presentation(path).slides) == 9
    # This offline fixture deliberately has no photos. Check for introduced
    # geometry/text errors against the same authored deck, not missing assets.
    baseline = lint_pptx(build_deck(fake.deck, tmp_path / "baseline.pptx", strict=True))
    existing = {(error["slide"], error["code"]) for error in baseline["errors"]}
    assert {(error["slide"], error["code"]) for error in lint_pptx(path)["errors"]} <= existing


def test_exhausted_malformed_slide_three_uses_source_and_still_reviews(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = BrokenThirdSlide(
        2, failure=StructuredOutputError("write", 2, "response did not match the required schema")
    )
    assert len(generate_deck(PROMPT, generator=fake).slides[2].items) == 3
    assert fake.stages[-1] == "review"


def test_source_recovery_does_not_hide_provider_access_failure():
    failure = StructuredOutputError("write", 1, "provider request failed (HTTP 401)")
    failure.__cause__ = LLMProviderError("provider request failed", status_code=401)
    with pytest.raises(StructuredOutputError) as caught:
        generate_deck(PROMPT, generator=BrokenThirdSlide(2, failure=failure))
    assert caught.value is failure


def test_no_source_recovery_for_unprovided_properties_or_unrelated_plans():
    plan = StoryPlan(
        archetype="split_image",
        title="Solution",
        purpose="Explain the source beauty proposition in three dimensions:",
    )
    assert source_copy_fallback(plan, "A brand selling foundation, lipstick and skincare.") is None
    assert source_copy_fallback(plan.model_copy(update={"purpose": "Other topic"}), PROMPT) is None
    fallback = source_copy_fallback(plan, PROMPT)
    validate_copy(fallback, plan, PROMPT)


def test_unsourced_count_failure_still_fails_without_fabricating_items(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.source_copy_fallback", lambda *a: None)
    with pytest.raises(SlideWriteError) as caught:
        generate_deck(PROMPT, generator=BrokenThirdSlide(2), semantic_attempts=1)
    assert caught.value.reason == "item_count_mismatch"
    assert "requires 3-3 items; received 2" in caught.value.detail


@pytest.mark.parametrize("archetype", ["split_image", "three_cards", "process_steps"])
def test_normalization_keeps_title_layout_and_evidence_in_the_same_source_role(archetype):
    fake = StoryGenerator()
    outline = fake.generate(StoryOutline, stage="outline")
    outline.facts = [PROMPT.split(". ")[0], "shade-matching quiz, waitlist, and creator-led demos"]
    for plan in outline.slides:
        plan.source_fact_ids = [0]
    outline.slides[2] = StoryPlan(
        archetype="split_image",
        title="Solution: Shade-Matching Quiz",
        purpose="Quiz solution",
        source_fact_ids=[1],
    )
    outline.slides[4] = StoryPlan(
        archetype=archetype,
        title="How It Works: Custom Blending",
        purpose="Draft unsupported production process",
        source_fact_ids=[0],
    )
    normalized = normalize_outline(outline, PROMPT)
    solution, workflow = normalized.slides[2], normalized.slides[4]
    assert solution.title == "Beauty made personal"
    assert solution.source_fact_ids == [0]
    assert workflow.archetype == "process_steps"
    assert workflow.purpose.startswith("Three touchpoints in order:")
    assert workflow.source_fact_ids == [1]
    assert source_copy_fallback(workflow, PROMPT) is not None


def test_review_references_are_bounded_by_actual_slides_and_fact_ids():
    model = planned_review_model(8, 3)
    with pytest.raises(ValidationError):
        model(passed=False, issues=[{"slide_index": 8, "fix": "Wrong reference"}])
    with pytest.raises(ValidationError):
        model(passed=False, missing_fact_ids=[3])
    assert model(passed=True).missing_fact_ids == []
    with pytest.raises(ValidationError):
        planned_review_model(8, 0)(passed=False, missing_fact_ids=[0])


@pytest.mark.parametrize("invented_fact", [False, True])
def test_outline_native_schema_recovery_preserves_source_checks_and_reaches_saved_deck(
    monkeypatch, tmp_path, invented_fact
):
    from app.render.builder import build_deck

    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fixture = StoryGenerator()
    calls = []

    class Provider:
        def complete(self, messages, **kwargs):
            calls.append(kwargs["response_format"]["type"])
            schema = json.loads(messages[0]["content"].split("JSON Schema:\n")[1])
            if len(calls) == 1:
                raise parse_error("json_validate_failed")
            if schema["title"] == "StoryOutline":
                result = fixture.generate(StoryOutline, stage="outline")
                if invented_fact:
                    result.facts.append("PRIVATE invented company revenue")
                return result.model_dump_json()
            if schema["title"] == "PlannedCopyReview":
                return '{"passed":true,"issues":[],"missing_fact_ids":[]}'
            return fixture.generate(
                StorySlide, stage="write", user_prompt=messages[1]["content"]
            ).model_dump_json()

    generator = StructuredGenerator(LLMClient(Provider()))
    if invented_fact:
        with pytest.raises(OutlineSourceError):
            generate_deck(PROMPT, generator=generator, semantic_attempts=1)
        assert fixture.stages == ["outline"]
    else:
        deck = generate_deck(PROMPT, generator=generator)
        path = build_deck(deck, tmp_path / "outline-recovery.pptx", strict=True)
        assert len(Presentation(path).slides) == 9
    assert calls[:2] == ["json_schema", "json_object"]


class DuplicateLaunchReview(StoryGenerator):
    def __init__(self, *, always_reject=False):
        super().__init__()
        self.reviews = []
        self.always_reject = always_reject

    def generate(self, output_model, *, stage, **kwargs):
        if stage == "review":
            request = json.loads(kwargs["user_prompt"])
            self.reviews.append(request)
            if len(self.reviews) == 1 or self.always_reject:
                return output_model(
                    passed=False,
                    issues=[
                        {
                            "slide_index": 5 if len(request["slides"]) == 9 else 4,
                            "fix": "Remove duplicate launch channels instead of rewriting the same content.",
                        }
                    ],
                )
            return output_model(passed=True)
        result = super().generate(output_model, stage=stage, **kwargs)
        if stage == "write":
            plan = StoryPlan.model_validate(json.loads(kwargs["user_prompt"])["plan"])
            if plan.purpose.startswith(
                ("Three touchpoints in order:", "Explain the three supplied launch channels")
            ):
                result = source_copy_fallback(plan, PROMPT)
        return result


def test_duplicate_source_workflow_merges_to_eight_slides_and_requires_another_review(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = DuplicateLaunchReview()
    deck = generate_deck(PROMPT, generator=fake, max_repair_rounds=0)
    assert len(deck.slides) == 8
    assert [len(review["slides"]) for review in fake.reviews] == [9, 8]
    assert [fact["id"] for fact in fake.reviews[0]["facts"]] == list(range(5))
    assert len({slide.archetype for slide in deck.slides}) >= 5
    assert fake.stages.count("write") == 9
    assert len([slide for slide in deck.slides if slide.title == "Launch touchpoints"]) == 1


def test_merging_never_skips_a_failed_final_review_or_drops_below_eight(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = DuplicateLaunchReview(always_reject=True)
    with pytest.raises(DeckReviewError) as caught:
        generate_deck(PROMPT, generator=fake, max_repair_rounds=0)
    assert caught.value.repair_rounds == 0
    assert caught.value.review.issues[0].slide_index == 4
    assert len(fake.reviews) == 2
    assert len(fake.reviews[-1]["slides"]) == 8


class MultipleFindingsReview(StoryGenerator):
    def __init__(self, *, always_reject=False, no_findings=False, missing_fact=False):
        super().__init__()
        self.always_reject = always_reject
        self.no_findings = no_findings
        self.missing_fact = missing_fact
        self.reviews = []
        self.repairs = []

    def generate(self, output_model, *, stage, **kwargs):
        if stage == "review":
            request = json.loads(kwargs["user_prompt"])
            self.stages.append(stage)
            self.reviews.append(request)
            item = request["slides"][1]["items"][0]
            passed = (
                item["title"] == "Individual style"
                and item["description"] == "Foundation, lipstick, and skincare."
                and not self.always_reject
            )
            return output_model(
                passed=passed,
                issues=[]
                if passed or self.no_findings
                else [
                    {"slide_index": 1, "fix": "PRIVATE_FINDING: Revise the heading."},
                    {"slide_index": 1, "fix": "PRIVATE_FINDING: Revise the description."},
                ],
                missing_fact_ids=[0] if self.missing_fact and not passed else [],
            )
        if stage == "repair":
            request = json.loads(kwargs["user_prompt"])
            self.stages.append(stage)
            self.repairs.append(request)
            slide = StorySlide.model_validate(request["failing_slide"])
            if "Revise the heading." in request["feedback"]:
                slide.items[0].title = "Individual style"
            if "Revise the description." in request["feedback"]:
                slide.items[0].description = "Foundation, lipstick, and skincare."
            return slide
        return super().generate(output_model, stage=stage, **kwargs)


def test_every_finding_for_a_slide_reaches_one_repair_and_passes_review(monkeypatch, caplog):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    caplog.set_level("INFO", logger="app.pipeline.editorial")
    fake = MultipleFindingsReview(missing_fact=True)
    deck = generate_deck(PROMPT, generator=fake)
    assert len(fake.reviews) == 2
    assert len(fake.repairs) == 2  # One issue slide and one assigned the missing fact.
    repair = next(r for r in fake.repairs if "Revise the heading." in r["feedback"])
    assert "Revise the description." in repair["feedback"]
    assert any("Include visibly: foundation" in r["feedback"] for r in fake.repairs)
    assert deck.slides[1].items[0].title == "Individual style"
    assert deck.slides[1].items[0].description == "Foundation, lipstick, and skincare."
    for index in range(2, len(deck.slides)):
        assert {"slide_index": index, **visible_copy(deck.slides[index])} == (
            fake.reviews[0]["slides"][index]
        )
    assert "issues=2 slide_indices=[1] missing_fact_ids=[0] repair_rounds=0" in caplog.text
    assert "PRIVATE_FINDING" not in caplog.text


def test_exhausted_editorial_review_preserves_findings_and_bounded_repairs(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = MultipleFindingsReview(always_reject=True)
    with pytest.raises(DeckReviewError) as caught:
        generate_deck(PROMPT, generator=fake, max_repair_rounds=2)
    assert len(fake.reviews) == 3
    assert len(fake.repairs) == 2
    assert caught.value.repair_rounds == 2
    assert caught.value.stage == "review"
    assert caught.value.reason == "source_review_failed"
    assert [issue.slide_index for issue in caught.value.review.issues] == [1, 1]


def test_failed_editorial_review_without_findings_retains_the_review(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = MultipleFindingsReview(no_findings=True)
    with pytest.raises(DeckReviewError) as caught:
        generate_deck(PROMPT, generator=fake)
    assert caught.value.review.passed is False
    assert caught.value.review.issues == []
    assert caught.value.repair_rounds == 0
    assert fake.repairs == []


@pytest.mark.parametrize("always_reject", [False, True])
def test_multiple_findings_after_merge_use_current_indices_and_bounded_repairs(
    monkeypatch, always_reject
):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)

    class FindingsAfterMerge(DuplicateLaunchReview):
        def __init__(self):
            super().__init__()
            self.repairs = []

        def generate(self, output_model, *, stage, **kwargs):
            request = json.loads(kwargs["user_prompt"]) if stage != "outline" else None
            if stage == "review" and len(request["slides"]) == 8:
                self.reviews.append(request)
                item = request["slides"][6]["items"][0]
                passed = (
                    item["title"] == "Measure interest"
                    and item["description"] == "Compare proposed signals over time."
                    and not always_reject
                )
                return output_model(
                    passed=passed,
                    issues=[]
                    if passed
                    else [
                        {"slide_index": 6, "fix": "Revise the heading."},
                        {"slide_index": 6, "fix": "Revise the description."},
                    ],
                )
            if stage == "repair":
                self.repairs.append(request)
                slide = StorySlide.model_validate(request["failing_slide"])
                if "Revise the heading." in request["feedback"]:
                    slide.items[0].title = "Measure interest"
                if "Revise the description." in request["feedback"]:
                    slide.items[0].description = "Compare proposed signals over time."
                return slide
            return super().generate(output_model, stage=stage, **kwargs)

    fake = FindingsAfterMerge()
    if always_reject:
        with pytest.raises(DeckReviewError) as caught:
            generate_deck(PROMPT, generator=fake, max_repair_rounds=2)
        assert caught.value.repair_rounds == 2
        assert [issue.slide_index for issue in caught.value.review.issues] == [6, 6]
        assert len(fake.repairs) == 2
    else:
        deck = generate_deck(PROMPT, generator=fake, max_repair_rounds=2)
        assert len(deck.slides) == 8
        assert deck.slides[6].items[0].title == "Measure interest"
        assert len(fake.repairs) == 1
    assert [len(review["slides"]) for review in fake.reviews] == [9] + [8] * (len(fake.repairs) + 1)
    for repair in fake.repairs:
        assert "Revise the heading." in repair["feedback"]
        assert "Revise the description." in repair["feedback"]
        assert repair["failing_slide"]["archetype"] == fake.reviews[1]["slides"][6]["archetype"]
