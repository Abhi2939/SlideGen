"""Editable renderers for the four layouts based on the supplied references."""

import math

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.render.common import add_text, new_slide
from app.render.assets import add_asset
from app.render.fit import text_height
from app.render.theme import Theme


REFERENCE_THEME = {
    "font": "Calibri",
    "navy": "1B365D",
    "blue": "263FB8",
    "green": "7DBE1E",
    "text": "333333",
    "muted": "6B7280",
    "line": "D1D5DB",
    "pale": "F3F4F6",
    "white": "FFFFFF",
}
FONT = REFERENCE_THEME["font"]
NAVY = RGBColor.from_string(REFERENCE_THEME["navy"])
BLUE = RGBColor.from_string(REFERENCE_THEME["blue"])
GREEN = RGBColor.from_string(REFERENCE_THEME["green"])
TEXT = RGBColor.from_string(REFERENCE_THEME["text"])
MUTED = RGBColor.from_string(REFERENCE_THEME["muted"])
LINE = RGBColor.from_string(REFERENCE_THEME["line"])
PALE = RGBColor.from_string(REFERENCE_THEME["pale"])
WHITE = RGBColor.from_string(REFERENCE_THEME["white"])
PALETTE = [
    "1B365D",
    "B8BCC2",
    "F5B800",
    "F26B1D",
    "2563EB",
    "9CA3AF",
    "7DBE1E",
]
PLACEHOLDER = "[Add value]"


def _fit_size(text, size, w, h):
    while size > 7 and text_height(text, size, Inches(w)) > Inches(h):
        size -= 0.5
    return size


def _shape(slide, x, y, w, h, fill, *, kind=MSO_SHAPE.RECTANGLE, line=None):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    for effect in shape._element.iter(qn("a:effectRef")):
        effect.set("idx", "0")
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = 0.16
    return shape


def _text(
    slide,
    x,
    y,
    w,
    h,
    value,
    *,
    size=14,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font=FONT,
):
    value = PLACEHOLDER if value is None else str(value)
    if value.startswith("[") or value.casefold().startswith(("add verified", "add supplied")):
        color = LINE if color == WHITE else MUTED
    size = _fit_size(value, size, w, h)
    box = add_text(
        slide,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        str(value),
        size_pt=size,
        color=color,
        font=font,
        bold=bold,
        align=align,
        anchor=valign,
    )
    if "\n" in value:
        # Native line breaks preserve paragraph centering in LibreOffice too.
        p = box.text_frame.paragraphs[0]
        p.runs[0].text = value.split("\n")[0]
        for part in value.split("\n")[1:]:
            p.add_line_break()
            run = p.add_run()
            run.text = part
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def _line(slide, x1, y1, x2, y2, *, color=LINE, width=1, dashed=False):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.shadow.inherit = False
    for effect in connector._element.iter(qn("a:effectRef")):
        effect.set("idx", "0")
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return connector


def _title(slide, text, *, subtitle=None, tagline=None, bg=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    _text(slide, 0.55, 0.28, 11.4, 0.70, text, size=27, color=NAVY, bold=True)
    if tagline:
        _text(slide, 0.58, 1.02, 12.1, 0.37, tagline, size=17, color=BLUE, bold=True)
    if subtitle:
        y = 1.42 if tagline else 1.02
        _text(slide, 0.58, y, 12.1, 0.32, subtitle, size=12, color=MUTED)


def render_ask_allocation(prs, data, theme: Theme):
    slide = new_slide(prs, theme)
    _title(slide, data.title)
    panel = _shape(slide, 0, 1.72, 3.65, 3.85, NAVY)
    tf = panel.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.28)
    tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = data.ask or PLACEHOLDER
    run.font.name = FONT
    run.font.size = Pt(_fit_size(run.text, 36, 3.17, 3.3))
    run.font.bold = True
    run.font.color.rgb = LINE if not data.ask else WHITE
    if add_asset(slide, "hero", 0.22, 4.55, 3.2, 1.75, crop=True):
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top = Inches(0.45)
    add_asset(slide, "edge", 11.55, 2.05, 1.55, 4.6, crop=True)

    items = data.items
    complete = all(item.pct is not None for item in items)
    if complete:
        values = [item.pct for item in items]
    else:
        known_total = sum(item.pct or 0 for item in items)
        unknown_share = (100 - known_total) / sum(item.pct is None for item in items)
        values = [item.pct if item.pct is not None else unknown_share for item in items]
    cx, cy, size = 7.25, 3.76, 2.55
    chart_data = CategoryChartData()
    chart_data.categories = [item.label for item in items] if complete else ["Allocation pending"]
    chart_data.add_series("Allocation", values if complete else [100])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(cx - size / 2),
        Inches(cy - size / 2),
        Inches(size),
        Inches(size),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot._element.find(qn("c:firstSliceAng")).set("val", "180")
    for index, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = (
            RGBColor.from_string(PALETTE[index % len(PALETTE)]) if complete else LINE
        )
    hole = plot._element.find(qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "66")
    has_icon = add_asset(slide, "centre", cx - 0.22, cy - 0.47, 0.44, 0.44)
    _text(
        slide,
        cx - 0.54,
        cy if has_icon else cy - 0.35,
        1.08,
        0.44 if has_icon else 0.7,
        data.center_label.replace("Fund Allocation", "Fund\nAllocation"),
        size=10 if has_icon else 12,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    total = sum(values) or 1
    accumulated = 0.0
    label_nodes = []
    for item, value in zip(items, values):
        midpoint = accumulated + value / total / 2
        accumulated += value / total
        theta = math.pi + 2 * math.pi * midpoint
        edge_x = cx + (size * 0.43) * math.sin(theta)
        edge_y = cy - (size * 0.43) * math.cos(theta)
        label_y = cy - (size / 2 + 0.42) * math.cos(theta)
        right = math.sin(theta) >= 0
        label_nodes.append((item, edge_x, edge_y, label_y, right))

    # Keep labels on their slice's side, in edge order. Elbows stay outside
    # the ring and terminate before the text column.
    for side in (False, True):
        nodes = sorted((node for node in label_nodes if node[4] is side), key=lambda node: node[2])
        slot_count = len(nodes)
        top, bottom = 1.85, 6.15
        gap = (bottom - top) / max(1, slot_count - 1)
        ys = [cy] if slot_count == 1 else [top + index * gap for index in range(slot_count)]
        track_gap = min(0.055, 0.32 / max(1, slot_count - 1))
        for track, (node, adjusted_y) in enumerate(zip(nodes, ys)):
            item, edge_x, edge_y, _, right = node
            end_x = 9.02 if right else 5.55
            elbow_x = cx + (size / 2 + 0.03) * (1 if right else -1)
            _line(slide, edge_x, edge_y, elbow_x, edge_y, color=MUTED, dashed=True)
            # Distinct nested tracks prevent concentrated slices sharing one
            # vertical leader line. Anchor order matches label order.
            track_x = elbow_x + track_gap * track * (1 if right else -1)
            _line(slide, elbow_x, edge_y, track_x, edge_y, color=MUTED, dashed=True)
            _line(slide, track_x, edge_y, track_x, adjusted_y, color=MUTED, dashed=True)
            _line(slide, track_x, adjusted_y, end_x, adjusted_y, color=MUTED, dashed=True)
            box_x = 9.08 if right else 3.78
            width = 2.13 if right else 1.70
            _text(
                slide,
                box_x,
                adjusted_y - 0.31,
                width,
                0.28,
                f"{item.pct:.2f}%" if item.pct is not None else "[ % ]",
                size=20,
                color=NAVY,
                bold=True,
                align=PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT,
            )
            _text(
                slide,
                box_x,
                adjusted_y + 0.015,
                width,
                0.5,
                item.label,
                size=10,
                color=TEXT,
                align=PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT,
            )
    if not complete:
        _text(
            slide,
            4.0,
            6.72,
            7.5,
            0.25,
            "Incomplete percentages: the grey ring is a placeholder.",
            size=9,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
    return slide


def _market_table(slide, data):
    if not data.table_headers or not data.table_rows:
        return
    rows = len(data.table_rows) + 1
    cols = len(data.table_headers)
    table = slide.shapes.add_table(
        rows, cols, Inches(0.48), Inches(4.46), Inches(12.35), Inches(min(2.62, 0.44 * rows))
    ).table
    table.first_row = False
    table.horz_banding = False
    widths = [0.37, 0.21, 0.25, 0.17] if cols == 4 else [1 / cols] * cols
    for column, fraction in zip(table.columns, widths):
        column.width = Inches(12.35 * fraction)

    def style(cell, text, cell_width, row_height, *, header=False, fill=WHITE, color=TEXT):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.025)
        cell.margin_bottom = Inches(0.015)
        tf = cell.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(
            _fit_size(text, 9 if not header else 10, cell_width - 0.14, row_height - 0.05)
        )
        run.font.bold = header
        run.font.color.rgb = (
            MUTED if text.startswith("[") or text.casefold().startswith("add verified") else color
        )

    for col, header in enumerate(data.table_headers):
        style(
            table.cell(0, col),
            header,
            table.columns[col].width / Inches(1),
            table.rows[0].height / Inches(1),
            header=True,
            fill=GREEN,
            color=WHITE,
        )
    for row_index, row in enumerate(data.table_rows, start=1):
        for col, value in enumerate(row):
            fill = WHITE if row_index % 2 else PALE
            if col == cols - 1:
                lower = value.casefold()
                if "short" in lower:
                    fill = RGBColor.from_string("FBE4E4")
                elif "medium" in lower:
                    fill = RGBColor.from_string("FDE9B8")
                elif "long" in lower:
                    fill = RGBColor.from_string("DDEFC4")
            style(
                table.cell(row_index, col),
                value,
                table.columns[col].width / Inches(1),
                table.rows[row_index].height / Inches(1),
                fill=fill,
            )
        row_y = 4.46 + sum(r.height for r in list(table.rows)[:row_index]) / Inches(1)
        if add_asset(slide, f"driver:{row_index - 1}", 0.55, row_y + 0.06, 0.20, 0.20):
            table.cell(row_index, 0).margin_left = Inches(0.32)


def render_market_snapshot(prs, data, theme: Theme):
    slide = new_slide(prs, theme)
    _title(slide, data.title, subtitle=data.subtitle, bg=PALE)

    _shape(slide, 0.5, 1.48, 4.18, 2.55, BLUE, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(slide, 0.78, 1.7, 3.55, 0.30, data.end_label, size=12, color=WHITE)
    _text(
        slide,
        0.78,
        2.03,
        3.55,
        0.74,
        data.end_value or PLACEHOLDER,
        size=25,
        color=WHITE,
        bold=True,
    )
    _shape(
        slide,
        0.76,
        2.86,
        3.25,
        1.03,
        WHITE,
        kind=MSO_SHAPE.ROUNDED_RECTANGLE,
        line=LINE,
    )
    _text(slide, 1.02, 2.97, 2.75, 0.26, data.start_label, size=11, color=MUTED)
    _text(
        slide,
        1.02,
        3.25,
        2.75,
        0.52,
        data.start_value or PLACEHOLDER,
        size=19,
        color=TEXT,
        bold=True,
    )

    _shape(slide, 4.93, 1.48, 7.88, 2.55, WHITE, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
    segments = data.segments
    if segments:
        column_width = 7.45 / len(segments)
        for index, segment in enumerate(segments):
            x = 5.14 + index * column_width
            if index:
                _line(slide, x - 0.12, 1.7, x - 0.12, 3.82)
            _text(slide, x, 1.68, column_width - 0.22, 0.34, segment.category, size=10, color=MUTED)
            _text(
                slide,
                x,
                2.08,
                column_width - 0.22,
                0.65,
                " ".join(filter(None, [segment.lead_text, segment.lead_value])),
                size=12,
                color=BLUE,
                bold=True,
            )
            _text(slide, x, 2.78, column_width - 0.22, 0.28, "↓", size=17, color=GREEN, bold=True)
            _text(slide, x, 3.08, column_width - 0.22, 0.43, segment.next_text, size=10, color=TEXT)
            if segment.next_value:
                _text(
                    slide,
                    x,
                    3.55,
                    column_width - 0.22,
                    0.3,
                    segment.next_value,
                    size=12,
                    color=BLUE,
                    bold=True,
                )
    else:
        _text(slide, 5.25, 2.45, 7.2, 0.55, "Add up to four market segments", size=14, color=MUTED)

    _market_table(slide, data)
    if not data.table_rows:
        _shape(slide, 0.5, 4.42, 12.31, 1.15, WHITE, kind=MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
        _text(
            slide,
            0.78,
            4.64,
            11.75,
            0.28,
            "Market inputs to complete",
            size=14,
            color=TEXT,
            bold=True,
        )
        _text(
            slide,
            0.78,
            5.02,
            11.75,
            0.3,
            "[Add sourced market size, growth rates, segment shares, and drivers]",
            size=12,
            color=MUTED,
        )
    return slide


def render_flow_steps(prs, data, theme: Theme):
    slide = new_slide(prs, theme)
    _title(slide, data.title, tagline=data.tagline, subtitle=data.subtitle, bg=PALE)
    steps = data.steps
    left = 1.4
    right = 13.33 - left
    gap = (right - left) / (len(steps) - 1)
    card_w = min(2.65, gap * 1.85)
    centers = []
    for index in range(len(steps)):
        t = index / (len(steps) - 1)
        centers.append((left + index * gap, 2.8 + 0.48 * (1 - math.sin(math.pi * t))))
    arc = [
        (left + (right - left) * i / 32, 2.8 + 0.48 * (1 - math.sin(math.pi * i / 32)))
        for i in range(33)
    ]
    for (x1, y1), (x2, y2) in zip(arc, arc[1:]):
        _line(slide, x1, y1, x2, y2, color=GREEN, width=2)
    for index, (step, (cx, cy)) in enumerate(zip(steps, centers), start=1):
        _shape(slide, cx - 0.3, cy - 0.3, 0.6, 0.6, GREEN, kind=MSO_SHAPE.OVAL)
        has_icon = add_asset(slide, f"step:{index - 1}", cx - 0.21, cy - 0.21, 0.42, 0.42)
        if has_icon:
            _shape(slide, cx + 0.22, cy - 0.13, 0.26, 0.26, BLUE, kind=MSO_SHAPE.OVAL)
        _text(
            slide,
            cx + 0.22 if has_icon else cx - 0.3,
            cy - 0.13 if has_icon else cy - 0.27,
            0.26 if has_icon else 0.6,
            0.24 if has_icon else 0.52,
            index,
            size=9 if has_icon else 15,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        card_y = 3.85 if index % 2 else 5.45
        _line(slide, cx, cy + 0.32, cx, card_y, color=LINE, dashed=True)
        card = _shape(
            slide,
            cx - card_w / 2,
            card_y,
            card_w,
            1.48,
            WHITE,
            kind=MSO_SHAPE.ROUNDED_RECTANGLE,
            line=LINE,
        )
        tf = card.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.12)
        tf.margin_top = tf.margin_bottom = Inches(0.08)
        title_p = tf.paragraphs[0]
        title_p.alignment = PP_ALIGN.CENTER
        title_run = title_p.add_run()
        title_run.text = step.title
        title_run.font.name = FONT
        title_size = max(9, _fit_size(step.title, 12, card_w - 0.24, 0.62 if step.detail else 1.25))
        title_run.font.size = Pt(title_size)
        title_run.font.bold = True
        title_run.font.color.rgb = MUTED if step.title.startswith("[") else TEXT
        if step.detail:
            detail_p = tf.add_paragraph()
            detail_p.alignment = PP_ALIGN.CENTER
            detail_run = detail_p.add_run()
            detail_run.text = step.detail
            detail_run.font.name = FONT
            detail_run.font.size = Pt(9)
            detail_run.font.color.rgb = MUTED
    return slide


def render_quadrants(prs, data, theme: Theme):
    slide = new_slide(prs, theme)
    _title(slide, data.title)
    _line(slide, 6.67, 1.4, 6.67, 7.12, color=LINE, dashed=True)
    _line(slide, 0.48, 4.22, 12.85, 4.22, color=LINE, dashed=True)
    for index, quadrant in enumerate(data.quadrants):
        col, row = index % 2, index // 2
        x = 0.56 + col * 6.3
        heading_y = 1.5 if row == 0 else 6.62
        pill = _shape(
            slide,
            x + 0.48,
            heading_y,
            5.05,
            0.48,
            GREEN,
            kind=MSO_SHAPE.ROUNDED_RECTANGLE,
        )
        pill.adjustments[0] = 0.5
        _text(
            slide,
            x + 0.58,
            heading_y + 0.03,
            4.85,
            0.4,
            quadrant.heading,
            size=11,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        chip_top = 2.18 if row == 0 else 4.58
        for item_index, name in enumerate(quadrant.names):
            chip_x = x + 0.34 + (item_index % 3) * 1.88
            chip_y = chip_top + (item_index // 3) * 0.69
            if add_asset(slide, name, chip_x + 0.08, chip_y + 0.04, 1.56, 0.52):
                continue
            chip = _shape(
                slide,
                chip_x,
                chip_y,
                1.72,
                0.60,
                WHITE,
                kind=MSO_SHAPE.ROUNDED_RECTANGLE,
                line=LINE,
            )
            tf = chip.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.06)
            tf.margin_top = tf.margin_bottom = Inches(0.04)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = name
            run.font.name = FONT
            run.font.size = Pt(_fit_size(name, 10, 1.60, 0.50))
            run.font.bold = True
            run.font.color.rgb = MUTED if name.startswith("[") else TEXT
        if not quadrant.names:
            _text(
                slide,
                x + 0.42,
                chip_top + 0.4,
                5.18,
                1.15,
                "[Add verified competitor names or examples]",
                size=13,
                color=MUTED,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
    return slide


RENDERERS = {
    "ask_allocation": render_ask_allocation,
    "market_snapshot": render_market_snapshot,
    "flow_steps": render_flow_steps,
    "quadrants": render_quadrants,
}
