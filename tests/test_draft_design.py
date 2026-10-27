import re
from pathlib import Path

import pytest
from pptx import Presentation

from app.core.deck_schema import Deck
from app.qa.draft import build_draft
from app.qa.lint import lint_pptx
from app.render.draft import composition


FIXTURE = Path("tests/fixtures/decks/blush_reference_draft.json")


def test_reference_design_preserves_source_in_distinct_editable_diagrams(tmp_path):
    deck = Deck.model_validate_json(FIXTURE.read_text())
    path = build_draft(deck, tmp_path / "blush.pptx")
    prs = Presentation(path)
    assert [s._element.cSld.get("name") for s in prs.slides] == [
        "draft_cover",
        "draft_collection",
        "draft_personalization",
        "draft_launch",
        "draft_growth",
        "draft_closing",
    ]
    for data, slide in zip(deck.slides, prs.slides):
        visible = " ".join(s.text for s in slide.shapes if s.has_text_frame)
        for point in getattr(data, "points", []):
            assert point in visible
        assert not any(s.shape_type == 13 for s in slide.shapes)
        assert all(
            run.font.name == "Calibri"
            for shape in slide.shapes
            if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
        )
    growth = prs.slides[4]
    labels = [s.text for s in growth.shapes if s.name == "Workflow activity"]
    assert labels == ["UGC", "Subscriptions", "New product lines", "Pop-ups", "Festive scaling"]
    assert not re.search(r"\bQ[1-4]\b", " ".join(s.text for s in growth.shapes if s.has_text_frame))
    assert lint_pptx(path, whitespace=False)["errors"] == []


@pytest.mark.parametrize(
    "points,kind,labels",
    [
        (
            [
                "Blended exclusively for your skin tone, undertone, and style, so you never get the wrong shade."
            ],
            "personalization",
            ["skin tone", "undertone", "style"],
        ),
        (
            ["Our launch will focus on a shade-matching quiz, waitlist, and creator-led demos,"],
            "launch",
            ["a shade-matching quiz", "waitlist", "creator-led demos"],
        ),
        (
            [
                "Our launch will focus on a shade-matching quiz, waitlist, and creator-led demos, followed by quarterly growth through UGC, subscriptions, new product lines, pop-ups, and festive scaling."
            ],
            "growth",
            ["UGC", "subscriptions", "new product lines", "pop-ups", "festive scaling"],
        ),
        (["A conceptual workflow with no supplied steps."], "points", []),
    ],
)
def test_diagrams_use_only_labels_extracted_from_source(points, kind, labels):
    assert composition(points) == (kind, labels)


@pytest.mark.parametrize("count", [1, 2, 3, 4])
def test_generic_and_dense_drafts_keep_their_text_and_pass_export(tmp_path, count):
    deck = Deck.model_validate_json(FIXTURE.read_text())
    value = "Discuss the supplied product details with the team and agree on the next planning priorities."
    points = [value] * count
    slide = {"layout": "draft_content", "title": "Planning priorities", "points": points}
    deck = Deck.model_validate(
        {
            **deck.model_dump(),
            "slides": [deck.slides[0].model_dump(), slide, deck.slides[-1].model_dump()],
        }
    )
    path = build_draft(deck, tmp_path / "generic.pptx")
    visible = " ".join(s.text for s in Presentation(path).slides[1].shapes if s.has_text_frame)
    assert visible.count(value) == count
    assert lint_pptx(path, whitespace=False)["errors"] == []
