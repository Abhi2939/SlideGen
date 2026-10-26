import json
from pathlib import Path

import pytest
from pptx import Presentation

from app.core.deck_schema import Deck, StorySlide
from app.core.errors import DeckReviewError, OutlineSourceError, ProviderLimitError, SlideWriteError
from app.pipeline.editorial import (
    clean_copy,
    ground_roadmap,
    source_launch_fallback,
    validate_copy,
    StoryPlan,
)
from app.pipeline.orchestrator import generate_deck

PROMPT = "Blush is a custom beauty brand offering foundation, lipstick, and skincare blended exclusively for your skin tone, undertone, and style, so you never get the wrong shade. Our launch will focus on a shade-matching quiz, waitlist, and creator-led demos, followed by quarterly growth through UGC, subscriptions, new product lines, pop-ups, and festive scaling."
FIXTURE = Path("tests/fixtures/decks/blush_story.json")


class StoryGenerator:
    def __init__(self, deck=None):
        self.deck = deck or Deck.model_validate_json(FIXTURE.read_text())
        self.index = 0
        self.stages = []

    def generate(self, output_model, *, stage, **kwargs):
        self.stages.append(stage)
        if stage == "outline":
            return output_model.model_validate(
                {
                    "brand": self.deck.title,
                    "audience": "Customers",
                    "tone": "Personal",
                    "theme": self.deck.theme,
                    "facts": ["foundation", "lipstick", "skincare", "UGC", "festive scaling"],
                    "slides": [
                        {
                            "archetype": s.archetype,
                            "title": s.title,
                            "purpose": s.title,
                            "source_fact_ids": list(range(5)),
                        }
                        for s in self.deck.slides
                    ],
                }
            )
        if stage == "write":
            slide = self.deck.slides[self.index]
            self.index += 1
            plan = json.loads(kwargs["user_prompt"])["plan"]
            return slide.model_copy(update={"archetype": plan["archetype"], "title": plan["title"]})
        if stage == "review":
            return output_model(passed=True, issues=[], missing_fact_ids=[])
        raise AssertionError(stage)


class SourceCorrectionGenerator(StoryGenerator):
    def __init__(self, rejected_fact, *, recover=True):
        super().__init__()
        self.rejected_fact = rejected_fact
        self.recover = recover
        self.requests = []
        self.outline_calls = 0

    def generate(self, output_model, *, stage, **kwargs):
        self.requests.append((stage, kwargs))
        result = super().generate(output_model, stage=stage, **kwargs)
        if stage == "outline":
            self.outline_calls += 1
            if not self.recover or self.outline_calls == 1:
                result.facts.append(self.rejected_fact)
                result.slides[0].source_fact_ids.append(len(result.facts) - 1)
        return result


@pytest.mark.parametrize(
    "rejected_fact",
    [
        "Blush offers personalized cosmetics for every customer",
        "PRIVATE_COMPLETION revenue reached $1m",
    ],
)
def test_outline_source_mismatch_is_corrected_before_writing(monkeypatch, rejected_fact):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = SourceCorrectionGenerator(rejected_fact)
    progress = []
    deck = generate_deck(PROMPT, generator=fake, progress_callback=progress.append)
    assert len(deck.slides) == 9
    assert fake.stages == ["outline", "outline"] + ["write"] * 9 + ["review"]
    correction = json.loads(fake.requests[1][1]["user_prompt"])
    assert correction["brief"] == PROMPT
    assert correction["rejected_facts"] == [{"id": 5, "text": rejected_fact}]
    assert "exact contiguous excerpt" in correction["correction"]
    assert "source_fact_ids" in correction["correction"]
    assert "Correcting source facts" in progress
    for stage, kwargs in fake.requests[2:]:
        assert rejected_fact not in kwargs["user_prompt"]


@pytest.mark.parametrize("attempts", [1, 2, 3])
def test_persistent_source_mismatch_is_bounded_and_never_written(attempts, caplog):
    fake = SourceCorrectionGenerator("PRIVATE_COMPLETION unsupported revenue", recover=False)
    with pytest.raises(OutlineSourceError) as caught:
        generate_deck(PROMPT, generator=fake, semantic_attempts=attempts)
    assert caught.value.stage == "outline"
    assert caught.value.reason == "outline_fact_not_in_source"
    assert caught.value.fact_ids == [5]
    assert fake.stages == ["outline"] * attempts
    assert "outline_fact_not_in_source" in caplog.text
    assert "PRIVATE" not in caplog.text
    assert "PRIVATE" not in str(caught.value)


def test_provider_limit_during_source_correction_stops_immediately():
    class LimitedGenerator(SourceCorrectionGenerator):
        def generate(self, *args, **kwargs):
            if self.outline_calls:
                raise ProviderLimitError()
            return super().generate(*args, **kwargs)

    fake = LimitedGenerator("PRIVATE_COMPLETION unsupported revenue")
    with pytest.raises(ProviderLimitError):
        generate_deck(PROMPT, generator=fake, semantic_attempts=3)
    assert fake.stages == ["outline"]


def test_source_correction_respects_job_deadline(monkeypatch):
    now = [100]
    monkeypatch.setattr("app.pipeline.editorial.time.monotonic", lambda: now[0])

    class SlowGenerator(SourceCorrectionGenerator):
        def generate(self, *args, **kwargs):
            result = super().generate(*args, **kwargs)
            now[0] = 102
            return result

    fake = SlowGenerator("PRIVATE_COMPLETION unsupported revenue")
    with pytest.raises(TimeoutError):
        generate_deck(PROMPT, generator=fake, deadline=101)
    assert fake.stages == ["outline"]


def test_new_default_uses_two_pass_story_and_preserves_editable_contract(monkeypatch, tmp_path):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = StoryGenerator()
    deck = generate_deck(PROMPT, generator=fake)
    assert deck.design_version == "gamma-v1" and deck.theme == "beauty"
    assert fake.stages == ["outline"] + ["write"] * 9 + ["review"]
    assert len({s.archetype for s in deck.slides}) >= 5
    assert [
        i.quarter for s in deck.slides if s.archetype == "timeline_horizontal" for i in s.items
    ] == ["Q1", "Q2", "Q3", "Q4", "Q4"]
    from app.render.builder import build_deck

    path = build_deck(deck, tmp_path / "story.pptx", strict=True)
    prs = Presentation(path)
    assert len(prs.slides) == 9
    assert all(slide.slide_layout.name != "Blank" for slide in prs.slides)
    assert all(slide.shapes.title is not None for slide in prs.slides)


def test_internal_copy_filter_moves_notes_off_canvas():
    slide = StorySlide(
        archetype="three_cards",
        title="Collection [internal note]",
        subtitle="Conceptual workflow details; confirm before use",
        items=[{"title": "Item", "description": "A useful description with enough detail."}] * 3,
    )
    result = clean_copy(slide)
    assert result.title == "Collection"
    assert result.subtitle is None
    assert "confirm before use" in result.notes


def test_exhausted_copy_validation_has_safe_reason_and_slide_index(monkeypatch):
    fake = StoryGenerator()
    monkeypatch.setattr("app.pipeline.editorial.clean_copy", lambda slide, prompt: slide)

    def reject_copy(*args):
        raise ValueError("unprovided numeric claim: PRIVATE_COMPLETION")

    monkeypatch.setattr("app.pipeline.editorial.validate_copy", reject_copy)
    with pytest.raises(SlideWriteError) as caught:
        generate_deck(PROMPT, generator=fake)
    assert caught.value.reason == "unprovided_numeric_data"
    assert caught.value.slide_index == 0
    assert caught.value.stage == "write"
    assert fake.stages == ["outline", "write", "write"]


def test_illustrative_numbers_require_visible_and_notes_disclosure():
    plan = StoryPlan(
        archetype="three_cards", title="Metrics", purpose="Illustrative planning targets"
    )
    slide = StorySlide(
        archetype="three_cards",
        title="Targets",
        illustrative=True,
        items=[{"title": "70%", "description": "A proposed quiz completion planning target."}] * 3,
    )
    with pytest.raises(ValueError, match="disclosure"):
        validate_copy(slide, plan, PROMPT)


def test_unprovided_company_values_are_rejected():
    plan = StoryPlan(archetype="big_stat", title="Traction", purpose="Results")
    slide = StorySlide(
        archetype="big_stat",
        title="Traction",
        items=[{"title": "99% conversion", "description": "A fabricated commercial result."}],
    )
    with pytest.raises(ValueError, match="numeric"):
        validate_copy(slide, plan, PROMPT)


def test_growth_requires_every_source_activity_separately():
    deck = Deck.model_validate_json(FIXTURE.read_text())
    timeline = next(s for s in deck.slides if s.archetype == "timeline_horizontal")
    plan = StoryPlan(
        archetype="timeline_horizontal", title="Growth roadmap", purpose="Quarterly growth"
    )
    validate_copy(timeline, plan, PROMPT)
    timeline.items[2].title = "Creator demos"
    with pytest.raises(ValueError, match="five separate activities"):
        validate_copy(timeline, plan, PROMPT)


@pytest.mark.parametrize("variation", ["reordered", "missing", "duplicate", "quarters", "aliases"])
def test_grounded_roadmap_recovers_source_structure_without_another_model_call(variation):
    deck = Deck.model_validate_json(FIXTURE.read_text())
    slide = next(s for s in deck.slides if s.archetype == "timeline_horizontal")
    plan = StoryPlan(archetype=slide.archetype, title=slide.title, purpose="Quarterly growth")
    original = slide.model_dump_json()
    if variation == "reordered":
        slide.items = list(reversed(slide.items))
    elif variation == "missing":
        slide.items = slide.items[:3]
    elif variation == "duplicate":
        slide.items[3] = slide.items[0].model_copy(deep=True)
    elif variation == "quarters":
        for item in slide.items:
            item.quarter = "Q1"
    else:
        for item, title in zip(
            slide.items,
            ["User-generated content", "Subscriptions", "New lines", "Pop-up stores", "Festive"],
        ):
            item.title = title
    grounded = ground_roadmap(slide, plan, PROMPT)
    validate_copy(grounded, plan, PROMPT)
    assert [i.title for i in grounded.items] == [
        "UGC",
        "subscriptions",
        "new product lines",
        "pop-ups",
        "festive scaling",
    ]
    assert [i.quarter for i in grounded.items] == ["Q1", "Q2", "Q3", "Q4", "Q4"]
    assert grounded.subtitle.startswith("Proposed")
    assert "not supplied dates" in grounded.notes
    if variation == "reordered":
        expected = StorySlide.model_validate_json(original)
        assert [i.description for i in grounded.items] == [i.description for i in expected.items]


def test_roadmap_discards_unrelated_activities_without_reassigning_their_claims():
    deck = Deck.model_validate_json(FIXTURE.read_text())
    slide = next(s for s in deck.slides if s.archetype == "timeline_horizontal")
    plan = StoryPlan(archetype=slide.archetype, title=slide.title, purpose="Quarterly growth")
    slide.items[2].title = "Creator demos"
    slide.items[2].description = "PRIVATE_GENERATION raised $99m with guaranteed shipping."
    grounded = ground_roadmap(slide, plan, PROMPT)
    validate_copy(grounded, plan, PROMPT)
    assert grounded.items[2].title == "new product lines"
    assert "PRIVATE_GENERATION" not in grounded.model_dump_json()
    assert "raised" not in grounded.items[2].description


def test_grounded_roadmap_still_rejects_unsupported_descriptions():
    deck = Deck.model_validate_json(FIXTURE.read_text())
    slide = next(s for s in deck.slides if s.archetype == "timeline_horizontal")
    plan = StoryPlan(archetype=slide.archetype, title=slide.title, purpose="Quarterly growth")
    slide.items[0].description = "Revenue reached $99m."
    grounded = ground_roadmap(slide, plan, PROMPT)
    with pytest.raises(ValueError, match="numeric"):
        validate_copy(grounded, plan, PROMPT)


def test_roadmap_without_explicit_activity_list_is_not_replaced():
    slide = Deck.model_validate_json(FIXTURE.read_text()).slides[6]
    plan = StoryPlan(archetype=slide.archetype, title=slide.title, purpose="Growth")
    assert ground_roadmap(slide, plan, "Create a general growth roadmap.") is slide


def test_roadmap_alias_comparison_is_symmetric_for_source_and_generated_labels():
    slide = Deck.model_validate_json(FIXTURE.read_text()).slides[6]
    plan = StoryPlan(archetype=slide.archetype, title=slide.title, purpose="Growth")
    validate_copy(slide, plan, PROMPT.replace("UGC,", "user-generated content,"))


class MismatchedRoadmapGenerator(StoryGenerator):
    def generate(self, output_model, *, stage, **kwargs):
        slide = super().generate(output_model, stage=stage, **kwargs)
        if stage == "write" and slide.archetype == "timeline_horizontal":
            slide = slide.model_copy(deep=True)
            slide.items = list(reversed(slide.items[:4]))
            for item in slide.items:
                item.quarter = "Q1"
        return slide


class UnsupportedLaunchGenerator(StoryGenerator):
    def __init__(self, *, correct_on_retry=False, limit_on_retry=False):
        super().__init__()
        self.launch_requests = []
        self.correct_on_retry = correct_on_retry
        self.limit_on_retry = limit_on_retry

    def generate(self, output_model, *, stage, **kwargs):
        if stage == "write":
            request = json.loads(kwargs["user_prompt"])
            if request["plan"]["purpose"].startswith("Three touchpoints in order:"):
                self.launch_requests.append(request)
                if len(self.launch_requests) == 1:
                    self.launch_copy = super().generate(output_model, stage=stage, **kwargs)
                else:
                    self.stages.append(stage)
                    if self.limit_on_retry:
                        raise ProviderLimitError()
                    if self.correct_on_retry:
                        return self.launch_copy.model_copy(deep=True)
                rejected = self.launch_copy.model_copy(deep=True)
                rejected.items[1].title = "Join the waitlist"
                rejected.items[
                    1
                ].description = "Secure early access and receive exclusive launch updates."
                return rejected
        return super().generate(output_model, stage=stage, **kwargs)


def test_semantic_retry_edits_the_rejected_copy_instead_of_starting_blind(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = UnsupportedLaunchGenerator(correct_on_retry=True)
    deck = generate_deck(PROMPT, generator=fake)
    first, correction = fake.launch_requests
    assert "rejected_copy" not in first
    assert "early access" in correction["rejected_copy"]["items"][1]["description"]
    assert "unsupported" in correction["feedback"]
    assert "untrusted model output" in correction["correction"]
    assert "notes" not in correction["rejected_copy"]
    assert correction["brief"] == PROMPT
    assert "early access" not in deck.slides[4].model_dump_json()


def test_exhausted_launch_claims_use_validated_source_copy_and_still_receive_review(
    monkeypatch, tmp_path
):
    from app.render.builder import build_deck

    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = UnsupportedLaunchGenerator()
    deck = generate_deck(PROMPT, generator=fake)
    assert len(fake.launch_requests) == 2
    assert fake.stages[-1] == "review"
    launch = deck.slides[4]
    assert [i.title for i in launch.items] == [
        "shade-matching quiz",
        "waitlist",
        "creator-led demos",
    ]
    assert "early access" not in launch.model_dump_json()
    assert "exclusive" not in launch.model_dump_json()
    assert launch.subtitle.startswith("Proposed")
    path = build_deck(deck, tmp_path / "source-launch.pptx", strict=True)
    assert len(Presentation(path).slides) == 9


def test_launch_source_fallback_never_masks_provider_limits(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = UnsupportedLaunchGenerator(limit_on_retry=True)
    with pytest.raises(ProviderLimitError):
        generate_deck(PROMPT, generator=fake)
    assert len(fake.launch_requests) == 2
    assert "review" not in fake.stages


def test_no_source_launch_fallback_when_explicit_supported_clause_is_absent(monkeypatch):
    fake = UnsupportedLaunchGenerator()
    prompt = PROMPT.replace("launch will focus on", "launch might eventually consider")
    with pytest.raises(SlideWriteError) as caught:
        generate_deck(prompt, generator=fake)
    assert caught.value.reason == "unsupported_claim"
    assert len(fake.launch_requests) == 2


def test_source_launch_fallback_is_limited_to_the_source_channel_plan():
    plan = StoryPlan(archetype="process_steps", title="Manufacturing", purpose="Explain production")
    assert source_launch_fallback(plan, PROMPT) is None
    plan = plan.model_copy(update={"purpose": "Three touchpoints in order: supplied launch"})
    assert (
        source_launch_fallback(plan, PROMPT.replace("creator-led demos", "AI recommendations"))
        is None
    )


def test_source_fallback_does_not_bypass_final_content_review(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)

    class RejectedReviewGenerator(UnsupportedLaunchGenerator):
        def generate(self, output_model, *, stage, **kwargs):
            if stage == "review":
                self.stages.append(stage)
                return output_model(
                    passed=False, issues=[{"slide_index": 4, "fix": "Preserve source wording"}]
                )
            return super().generate(output_model, stage=stage, **kwargs)

    fake = RejectedReviewGenerator()
    with pytest.raises(DeckReviewError):
        generate_deck(PROMPT, generator=fake, max_repair_rounds=0)
    assert len(fake.launch_requests) == 2
    assert fake.stages[-1] == "review"


def test_full_pipeline_grounds_slide_seven_before_review(monkeypatch, tmp_path):
    from app.render.builder import build_deck

    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    fake = MismatchedRoadmapGenerator()
    deck = generate_deck(PROMPT, generator=fake, semantic_attempts=1)
    assert fake.stages == ["outline"] + ["write"] * 9 + ["review"]
    validate_copy(
        deck.slides[6],
        StoryPlan(archetype="timeline_horizontal", title=deck.slides[6].title, purpose="Growth"),
        PROMPT,
    )
    path = build_deck(deck, tmp_path / "corrected-roadmap.pptx", strict=True)
    timeline = Presentation(path).slides[6]

    def visible_text(shapes):
        result = []
        for shape in shapes:
            if shape.has_text_frame:
                result.append(shape.text)
            if shape.shape_type == 6:
                result.extend(visible_text(shape.shapes))
        return result

    text = "\n".join(visible_text(timeline.shapes))
    for label in (
        "Q1  UGC",
        "Q2  subscriptions",
        "Q3  new product lines",
        "Q4  pop-ups",
        "Q4  festive scaling",
    ):
        assert label in text


def test_inferred_mechanisms_are_labelled_as_proposals():
    slide = StorySlide(
        archetype="three_cards",
        title="Product priorities",
        items=[
            {
                "title": "Dashboard",
                "description": "A dashboard bringing invoice information into a shared view.",
            }
        ]
        * 3,
    )
    result = clean_copy(slide, "Acme Pay helps SMEs organize invoice payments.")
    assert result.subtitle.startswith("Proposed concept")
    assert "not supplied current product facts" in result.notes


def test_placeholder_outline_entries_never_become_source_evidence():
    from app.pipeline.editorial import source_outline, StoryOutline

    deck = Deck.model_validate_json(FIXTURE.read_text())
    outline = StoryOutline(
        brand="Blush",
        audience="Customers",
        tone="Warm",
        theme="beauty",
        facts=["foundation", "Placeholder: market size"],
        slides=[
            {
                "archetype": s.archetype,
                "title": s.title,
                "purpose": s.title,
                "source_fact_ids": [0, 1],
            }
            for s in deck.slides
        ],
    )
    cleaned = source_outline(outline, PROMPT)
    assert cleaned.facts == ["foundation"]
    assert all(p.source_fact_ids == [0] for p in cleaned.slides)


def test_same_content_changes_theme_for_other_topics(monkeypatch):
    monkeypatch.setattr("app.pipeline.editorial.prepare_assets", lambda deck, *a, **k: deck)
    for topic, expected in [
        ("edtech learning platform", "edtech"),
        ("fintech payments pitch", "fintech"),
    ]:
        fixture = Deck.model_validate_json(FIXTURE.read_text()).model_copy(
            update={"theme": expected}
        )
        fake = StoryGenerator(fixture)
        # The fake preserves source evidence deliberately; no live provider calls.
        deck = generate_deck(PROMPT + " " + topic, generator=fake)
        assert deck.design_version == "gamma-v1"
        assert json.loads(deck.model_dump_json())["slides"][0]["archetype"] == "cover"
