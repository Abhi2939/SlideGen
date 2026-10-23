import pytest

from app.pipeline.analyze import analyze_prompt
from app.pipeline.grounding import unsupported_claim_categories
from app.pipeline.models import AnalysisOutput, OutlineOutput
from app.pipeline.outline import create_outline


@pytest.mark.parametrize(
    "claim",
    [
        "Follow us on socials",
        "Confidence in every purchase",
        "Book a demo",
        "Join our community",
        "Shop now",
        "Guaranteed results",
        "Guarantees you never get the wrong shade",
        "Stay tuned",
        "Join the waitlist",
        "Ready to launch",
    ],
)
def test_closing_claims_and_ctas_need_source_evidence(claim):
    assert unsupported_claim_categories(claim, "Blush offers custom foundation")
    assert unsupported_claim_categories(claim, claim) == []


def test_one_supplied_cta_does_not_authorize_other_ctas():
    assert unsupported_claim_categories("Follow us on socials", "Visit our website")


@pytest.mark.parametrize(
    "claim",
    [
        "Skincare formulated for your specific skin needs",
        "Build excitement with exclusive early access",
    ],
)
def test_refined_live_blush_additions_need_their_own_source_evidence(claim):
    assert unsupported_claim_categories(
        claim, "Custom products matched to skin tone and style. Launch uses a waitlist."
    )
    assert unsupported_claim_categories(claim, claim) == []


def test_one_product_specification_does_not_authorize_unrelated_capabilities():
    assert unsupported_claim_categories(
        "Skincare for specific skin needs", "User-friendly interface"
    )


def test_delivery_inflections_are_not_an_unsupported_claim_escape():
    assert unsupported_claim_categories("Delivering custom products", "Custom beauty products")


@pytest.mark.parametrize(
    "claim",
    [
        "Instant shade recommendation",
        "Instant results",
        "Priority access",
        "Exclusive launch updates",
    ],
)
def test_owner_deck_performance_and_access_promises_need_evidence(claim):
    source = (
        "Blush offers custom products. Launch plans include a quiz, waitlist and creator demos."
    )
    assert unsupported_claim_categories(claim, source, allow_structural_inference=True)
    assert unsupported_claim_categories(claim, source + " " + claim) == []


def test_generic_closing_actions_cannot_be_added_to_a_source_plan(tmp_path):
    from app.core.deck_schema import ClosingSlide, Deck
    from app.pipeline.write_slides import _write_one
    from app.render.builder import build_deck
    from pptx import Presentation

    analysis = AnalysisOutput(
        deck_title="Blush",
        deck_type="product",
        slide_count=3,
        required_facts=[{"id": 0, "text": "Blush offers custom foundation"}],
    )
    outline = OutlineOutput(
        deck_title="Blush",
        slides=[
            {"index": 0, "layout": "title", "title": "Blush", "purpose": "Introduce"},
            {"index": 1, "layout": "bullets", "title": "Products", "purpose": "Describe"},
            {
                "index": 2,
                "layout": "closing",
                "title": "Summary",
                "purpose": "Conclude",
                "required_fact_ids": [0],
            },
        ],
    )

    class Writer:
        calls = 0

        def generate(self, output_model, **kwargs):
            self.calls += 1
            if self.calls == 1:
                return ClosingSlide(
                    layout="closing", title="Summary", subtitle="Explore your personalized shade"
                )
            assert "closing action not supplied" in kwargs["user_prompt"]
            return ClosingSlide(
                layout="closing", title="Summary", subtitle="Blush offers custom foundation"
            )

    writer = Writer()
    closing = _write_one(
        "Blush is a beauty brand. Blush offers custom foundation.",
        analysis,
        outline,
        2,
        writer,
        "fake",
        semantic_attempts=2,
        max_tokens=1600,
    )
    assert writer.calls == 2
    deck = Deck(
        title="Blush",
        slides=[
            {"layout": "title", "title": "Blush"},
            {
                "layout": "bullets",
                "title": "Products",
                "bullets": [{"text": "Custom products"}, {"text": "Beauty"}],
            },
            closing,
        ],
    )
    presentation = Presentation(build_deck(deck, tmp_path / "closing.pptx", strict=True))
    assert "Explore your" not in "\n".join(
        s.text for s in presentation.slides[2].shapes if s.has_text_frame
    )


@pytest.mark.parametrize(
    "claim",
    [
        "Automated risk scoring",
        "Instant fund disbursement",
        "User-friendly interface",
        "Analyze undertone through a brief questionnaire and visual guide",
        "Large addressable market",
        "Significant SME base seeking financing",
        "Growing demand for digital credit",
        "Increasing adoption of fintech solutions",
        "Sign up for early access",
        "Contact sales for demo",
    ],
)
def test_live_discovered_unsupplied_mechanisms_market_claims_and_ctas_are_guarded(claim):
    assert unsupported_claim_categories(
        claim, "Fintech for SMEs. Include market.", allow_problem_framing=True
    )
    assert unsupported_claim_categories(claim, claim, allow_problem_framing=True) == []


def test_duplicate_bullet_preview_is_removed_when_visual_plan_covers_all_facts():
    from app.pipeline.outline import _compact_duplicate_fact_plans

    outline = OutlineOutput(
        deck_title="Blush",
        slides=[
            {"index": 0, "layout": "title", "title": "Blush", "purpose": "Introduce"},
            {
                "index": 1,
                "layout": "bullets",
                "title": "Products",
                "purpose": "Describe",
                "required_fact_ids": [0],
            },
            {
                "index": 2,
                "layout": "process",
                "title": "Launch",
                "purpose": "Plans",
                "required_fact_ids": [1, 2, 3],
            },
            {
                "index": 3,
                "layout": "bullets",
                "title": "Launch focus",
                "purpose": "Plans",
                "required_fact_ids": [1, 2, 3],
            },
            {"index": 4, "layout": "closing", "title": "Summary", "purpose": "Close"},
        ],
    )
    compact = _compact_duplicate_fact_plans("Blush launch and growth plans", outline)
    assert len(compact.slides) == 4
    assert [p.index for p in compact.slides] == [0, 1, 2, 3]
    assert {i for p in compact.slides for i in p.required_fact_ids} == {0, 1, 2, 3}
    assert _compact_duplicate_fact_plans("Make a 5-slide deck", outline) is outline


class Generator:
    def __init__(self, result):
        self.result = result

    def generate(self, output_model, **kwargs):
        return output_model.model_validate(self.result)


def test_sparse_educational_prompt_targets_six_substantive_slides():
    analysis = analyze_prompt(
        "Make a deck about dogs",
        Generator(
            {
                "deck_title": "Dogs",
                "deck_type": "education",
                "slide_count": 3,
            }
        ),
        "fake",
    )
    assert analysis.slide_count == 6


@pytest.mark.parametrize("count_requested", [False, True])
def test_analyzer_cannot_turn_inferred_target_into_a_hard_review_constraint(count_requested):
    prompt = "Blush is a beauty brand with launch and growth plans."
    if count_requested:
        prompt += " Make a 5-slide deck."
    result = analyze_prompt(
        prompt,
        Generator(
            {
                "deck_title": "Blush",
                "deck_type": "product",
                "slide_count": 5,
                "constraints": ["5 slides", "Slide count: five", "Use clear language"],
            }
        ),
        "fake",
    )
    assert result.slide_count == 5
    assert ("5 slides" in result.constraints) == count_requested
    assert ("Slide count: five" in result.constraints) == count_requested
    assert "Use clear language" in result.constraints


def test_outline_can_use_fewer_fuller_slides_without_losing_facts():
    analysis = AnalysisOutput(
        deck_title="Blush",
        deck_type="product",
        slide_count=8,
        required_facts=[
            {"id": 0, "text": "Blush makes custom foundation"},
            {"id": 1, "text": "Plans: UGC, subscriptions and pop-ups"},
        ],
    )
    result = {
        "deck_title": "Blush",
        "slides": [
            {"index": 0, "title": "Blush", "purpose": "Introduce", "layout": "title"},
            {
                "index": 1,
                "title": "Products",
                "purpose": "Describe",
                "layout": "bullets",
                "required_fact_ids": [0],
            },
            {
                "index": 2,
                "title": "Growth Roadmap",
                "purpose": "Future activities",
                "layout": "bullets",
                "required_fact_ids": [1],
            },
        ],
    }
    outline = create_outline(
        "Blush makes custom foundation. Growth through UGC, subscriptions and pop-ups.",
        analysis,
        Generator(result),
        "fake",
    )
    assert len(outline.slides) == 3
    assert outline.slides[2].layout == "flow_steps"
    assert set(fact for slide in outline.slides for fact in slide.required_fact_ids) == {0, 1}
    with pytest.raises(ValueError, match="analysis requested 8"):
        create_outline("Make an 8-slide Blush deck", analysis, Generator(result), "fake")


def test_fewer_slides_never_allows_dropping_source_facts():
    analysis = AnalysisOutput(
        deck_title="Blush",
        deck_type="product",
        slide_count=8,
        required_facts=[{"id": 0, "text": "Creator-led demos"}],
    )
    outline = OutlineOutput(
        deck_title="Blush",
        slides=[
            {
                "index": i,
                "title": "Blush" if i == 0 else "Products",
                "purpose": "Describe",
                "layout": "title" if i == 0 else "bullets",
            }
            for i in range(3)
        ],
    )
    with pytest.raises(ValueError, match="omits required fact"):
        create_outline("Blush creator-led demos", analysis, Generator(outline.model_dump()), "fake")


@pytest.mark.parametrize(
    "extra, expected_count", [("", 4), (" Include stats.", 5), (" Make a 5-slide deck.", 5)]
)
def test_live_blush_plan_avoids_unrequested_empty_stat_and_groups_product_attributes(
    extra, expected_count
):
    prompt = "Blush is a beauty brand offering foundation. Products are blended for your skin tone and undertone. Our launch uses a quiz, waitlist and creator demos. Growth through UGC and subscriptions."
    analysis = AnalysisOutput(
        deck_title="Blush",
        deck_type="product",
        slide_count=5,
        required_facts=[
            {"id": 0, "text": "Blush"},
            {"id": 1, "text": "foundation"},
            {"id": 2, "text": "Products are blended for your skin tone and undertone"},
            {"id": 3, "text": "Launch uses a quiz, waitlist and creator demos"},
            {"id": 4, "text": "Growth through UGC and subscriptions"},
        ],
    )
    response = {
        "deck_title": "Blush",
        "slides": [
            {
                "index": 0,
                "layout": "title",
                "title": "Blush",
                "purpose": "Introduce",
                "required_fact_ids": [0],
            },
            {
                "index": 1,
                "layout": "bullets",
                "title": "Our Products",
                "purpose": "Products",
                "required_fact_ids": [1],
            },
            {
                "index": 2,
                "layout": "bullets",
                "title": "Launch Strategy",
                "purpose": "Launch plans",
                "required_fact_ids": [2, 3],
            },
            {
                "index": 3,
                "layout": "process",
                "title": "Growth Plan",
                "purpose": "Future plans",
                "required_fact_ids": [4],
            },
            {
                "index": 4,
                "layout": "stat",
                "title": "Product snapshot",
                "purpose": "Summary",
                "key_points": ["3 categories"],
                "required_fact_ids": [],
            },
        ],
    }
    outline = create_outline(prompt + extra, analysis, Generator(response), "fake")
    assert len(outline.slides) == expected_count
    assert outline.slides[1].required_fact_ids == [1, 2]
    assert outline.slides[2].required_fact_ids == [3]
    assert outline.slides[3].layout == "flow_steps"
    assert {i for plan in outline.slides for i in plan.required_fact_ids} == {0, 1, 2, 3, 4}


def test_business_owners_as_audience_does_not_force_company_placeholders():
    from app.pipeline.write_slides import _can_use_sparse_company_placeholder

    analysis = AnalysisOutput(deck_title="Launch", deck_type="launch plan", slide_count=6)
    assert not _can_use_sparse_company_placeholder(
        "Create a 6-slide product launch presentation for small business owners", analysis
    )
    assert _can_use_sparse_company_placeholder("Make a deck for my electronics company", analysis)


def test_market_periods_need_evidence_and_defaults_do_not_invent_years():
    from app.core.deck_schema import MarketSnapshotSlide
    from app.pipeline.write_slides import _numeric_data_issue

    slide = MarketSnapshotSlide(layout="market_snapshot", title="Market")
    assert slide.start_label == "[Base period]"
    assert slide.end_label == "[Forecast period]"
    analysis = AnalysisOutput(deck_title="Market", deck_type="overview", slide_count=3)
    invented = slide.model_copy(update={"start_label": "2025", "end_label": "2030"})
    assert _numeric_data_issue(invented, analysis) is not None
    supplied = AnalysisOutput.model_validate(
        {**analysis.model_dump(), "required_facts": [{"id": 0, "text": "2025 to 2030"}]}
    )
    assert _numeric_data_issue(invented, supplied) is None


def test_review_aligns_exact_safe_scaffolds_without_dropping_findings():
    import json
    from app.pipeline.review import review_deck
    from app.pipeline.models import ReviewOutput
    from app.pipeline.write_slides import _placeholder_slide

    analysis = AnalysisOutput(deck_title="Launch", deck_type="company", slide_count=3)
    outline = OutlineOutput(
        deck_title="Launch",
        slides=[
            {"index": 0, "layout": "title", "title": "Launch", "purpose": "Introduce"},
            {
                "index": 1,
                "layout": "bullets",
                "title": "Unsupported product features",
                "purpose": "Explain",
                "key_points": ["Affordable pricing", "24/7 support"],
            },
            {"index": 2, "layout": "closing", "title": "Follow us", "purpose": "Close"},
        ],
    )
    slides = [_placeholder_slide(plan) for plan in outline.slides]

    class Reviewer:
        def generate(self, output_model, **kwargs):
            plans = json.loads(
                kwargs["user_prompt"].split("SLIDE OUTLINE:\n")[1].split("\n\nGENERATED DECK")[0]
            )
            assert plans["slides"][1]["title"] == slides[1].title
            assert plans["slides"][1]["key_points"] == []
            return ReviewOutput(
                passed=False,
                issues=[
                    {
                        "category": "overflow",
                        "severity": "high",
                        "detail": "Actual copy is too long",
                        "slide_indices": [1],
                    }
                ],
                regenerate_slide_indices=[1],
            )

    review = review_deck("Make a company deck", analysis, outline, slides, Reviewer(), "fake")
    assert not review.passed
    assert review.issues[0].category == "overflow"


def test_unsupplied_chart_numbers_are_placeholders_until_example_data_is_requested(tmp_path):
    from app.core.deck_schema import ChartSlide, Deck
    from app.pipeline.write_slides import _guard_illustrative_data, _slide_text_content
    from app.render.builder import build_deck
    from pptx import Presentation

    chart = ChartSlide(
        layout="chart",
        title="Launch metrics",
        chart_type="bar",
        categories=["Awareness", "Interest"],
        series=[{"name": "Example", "values": [42, 87]}],
        provenance="illustrative",
    )
    guarded = _guard_illustrative_data(chart, "Make a launch presentation")
    assert guarded.placeholder
    assert guarded.series[0].values == [0, 0]
    assert "42" not in _slide_text_content(guarded)
    deck = Deck(
        title="Launch",
        slides=[
            {"layout": "title", "title": "Launch"},
            guarded,
            {"layout": "closing", "title": "Summary"},
        ],
    )
    prs = Presentation(build_deck(deck, tmp_path / "placeholder.pptx", strict=True))
    assert not any(s.has_chart for s in prs.slides[1].shapes)
    assert any(s.has_text_frame and "no example figures" in s.text for s in prs.slides[1].shapes)
    assert _guard_illustrative_data(chart, "Use illustrative data") is chart
    dates = chart.model_copy(update={"categories": ["2025", "2030"]})
    assert _guard_illustrative_data(dates, "Make a launch deck").categories == [
        "[Add category]",
        "[Add category]",
    ]


def test_incomplete_allocation_uses_neutral_placeholder_ring(tmp_path):
    from app.core.deck_schema import Deck
    from app.render.builder import build_deck
    from pptx import Presentation

    deck = Deck(
        title="Funding",
        slides=[
            {"layout": "title", "title": "Funding"},
            {
                "layout": "ask_allocation",
                "title": "Allocation",
                "items": [{"label": "Product", "pct": None}, {"label": "Team", "pct": None}],
            },
            {"layout": "closing", "title": "Summary"},
        ],
    )
    prs = Presentation(build_deck(deck, tmp_path / "allocation.pptx", strict=True))
    chart = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
    assert list(chart.series[0].values) == [100]
    assert len(chart.series[0].points) == 1


def test_inferred_company_workflow_details_are_visibly_labelled():
    from app.core.deck_schema import FlowStepsSlide
    from app.pipeline.write_slides import _label_inferred_workflow

    slide = FlowStepsSlide(
        layout="flow_steps",
        title="Launch",
        steps=[
            {"title": "Quiz", "detail": "An online form to explore preferences"},
            {"title": "Waitlist", "detail": "Capture email addresses for notifications"},
            {"title": "Demos"},
        ],
    )
    result = _label_inferred_workflow(slide, "Our brand plans a quiz, waitlist and demos")
    assert "Conceptual" in result.subtitle
    assert _label_inferred_workflow(slide, "An educational presentation") is slide


def test_exhausted_workflow_json_can_complete_from_source_phrases_without_inventing_steps(tmp_path):
    import json
    from app.config import Settings
    from app.core.errors import StructuredOutputError
    from app.pipeline.orchestrator import generate_deck
    from app.render.builder import build_deck
    from pptx import Presentation

    fact = "shade-matching quiz, waitlist, and creator-led demos"

    class Stages:
        def generate(self, output_model, **kwargs):
            stage = kwargs["stage"]
            if stage == "analyze":
                return output_model(
                    deck_title="Blush",
                    deck_type="launch",
                    slide_count=3,
                    required_facts=[{"id": 0, "text": fact, "verbatim": True}],
                )
            if stage == "outline":
                return output_model(
                    deck_title="Blush",
                    slides=[
                        {"index": 0, "layout": "title", "title": "Blush", "purpose": "Introduce"},
                        {
                            "index": 1,
                            "layout": "process",
                            "title": "Launch Plan",
                            "purpose": "Present supplied launch activities",
                            "required_fact_ids": [0],
                            "key_points": [
                                "Shade-matching quiz",
                                "Waitlist",
                                "Creator-led demos",
                                "Invented sales channel",
                            ],
                        },
                        {
                            "index": 2,
                            "layout": "closing",
                            "title": "Summary",
                            "purpose": "Conclude",
                        },
                    ],
                )
            if stage == "write":
                plan = json.loads(
                    kwargs["user_prompt"].split("SLIDE PLAN:\n")[1].split("\n\nFACTS ASSIGNED")[0]
                )
                if plan["layout"] == "process":
                    raise StructuredOutputError("write", 2, "response was not valid JSON")
                return output_model(layout=plan["layout"], title=plan["title"])
            if stage == "review":
                outline = json.loads(
                    kwargs["user_prompt"]
                    .split("SLIDE OUTLINE:\n")[1]
                    .split("\n\nGENERATED DECK")[0]
                )
                assert outline["slides"][1]["layout"] == "flow_steps"
                return output_model(passed=True)
            raise AssertionError(stage)

    deck = generate_deck(
        f"Make a 3-slide Blush launch deck. Our launch: {fact}.",
        settings=Settings("groq", None, "fake", None, None, None, None, 2),
        generator=Stages(),
    )
    assert deck.slides[1].layout == "flow_steps"
    assert len(deck.slides[1].steps) == 3
    assert "Invented" not in deck.model_dump_json()
    presentation = Presentation(build_deck(deck, tmp_path / "source-workflow.pptx", strict=True))
    visible = "\n".join(s.text for s in presentation.slides[1].shapes if s.has_text_frame)
    assert fact in visible


@pytest.mark.parametrize("kind", ["auth", "quota"])
def test_source_workflow_never_masks_authentication_or_quota_failures(kind):
    from app.core.errors import LLMProviderError, ProviderLimitError, StructuredOutputError
    from app.pipeline.models import OutlineSlidePlan
    from app.pipeline.write_slides import _source_workflow, _write_one

    fact = "Quiz, waitlist, demos"
    plan = OutlineSlidePlan(
        index=1,
        layout="process",
        title="Launch Plan",
        purpose="Source launch",
        required_fact_ids=[0],
        key_points=["Quiz", "waitlist", "demos"],
    )
    analysis = AnalysisOutput(
        deck_title="Blush",
        deck_type="launch",
        slide_count=3,
        required_facts=[{"id": 0, "text": fact}],
    )
    outline = OutlineOutput(
        deck_title="Blush",
        slides=[
            {"index": 0, "layout": "title", "title": "Blush", "purpose": "Introduce"},
            plan,
            {"index": 2, "layout": "closing", "title": "Summary", "purpose": "Conclude"},
        ],
    )
    assert _source_workflow(plan, analysis.required_facts) is not None

    class Failure:
        def generate(self, *args, **kwargs):
            if kind == "quota":
                raise ProviderLimitError()
            raise StructuredOutputError(
                "write", 1, "provider request failed"
            ) from LLMProviderError("provider request failed", status_code=401)

    with pytest.raises(ProviderLimitError if kind == "quota" else StructuredOutputError):
        _write_one(
            "Our launch: " + fact,
            analysis,
            outline,
            1,
            Failure(),
            "fake",
            semantic_attempts=2,
            max_tokens=1600,
        )
