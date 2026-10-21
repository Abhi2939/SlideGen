import json
from pathlib import Path

import pytest
from pptx import Presentation

from app.core.deck_schema import Deck
from app.core.deck_schema import LAYOUTS
from app.render.builder import build_deck
from app.render.layout import LAYOUT_RENDERERS

FIXTURES = Path(__file__).parent


@pytest.mark.parametrize("path", sorted(FIXTURES.glob("*.json")), ids=lambda p: p.stem)
def test_fixture_renders_every_slide(path, tmp_path):
    deck = Deck.model_validate(json.loads(path.read_text(encoding="utf-8")))
    out = build_deck(deck, tmp_path / "out.pptx")
    assert len(Presentation(str(out)).slides) == len(deck.slides)


def test_each_schema_layout_has_a_renderer():
    assert set(LAYOUT_RENDERERS) == set(LAYOUTS)


def test_fixture_uses_no_placeholder_layouts(tmp_path):
    path = FIXTURES / "all_layouts.json"
    deck = Deck.model_validate(json.loads(path.read_text(encoding="utf-8")))
    out = build_deck(deck, tmp_path / "all-layouts.pptx", strict=True)
    presentation = Presentation(str(out))
    slide_text = [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert not any("layout not implemented yet" in text for text in slide_text)


def test_rendered_chart_shows_values_and_unit(tmp_path):
    path = FIXTURES / "all_layouts.json"
    deck = Deck.model_validate(json.loads(path.read_text(encoding="utf-8")))
    out = build_deck(deck, tmp_path / "charts.pptx", strict=True)
    presentation = Presentation(str(out))
    chart = next(
        shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart
    )
    assert chart.plots[0].has_data_labels
    assert chart.plots[0].data_labels.show_value


def test_blush_brand_theme_renders_editable_cosmetic_art(tmp_path):
    deck = Deck.model_validate(
        {
            "title": "Blush",
            "theme": "blush",
            "slides": [
                {"layout": "title", "title": "Blush", "subtitle": "Makeup made only for YOU."},
                {
                    "layout": "bullets",
                    "title": "The collection",
                    "bullets": [
                        {"text": "Custom-blended foundation"},
                        {"text": "Custom-blended lipstick and skincare"},
                    ],
                },
                {"layout": "closing", "title": "No more wrong shades"},
            ],
        }
    )

    out = build_deck(deck, tmp_path / "blush.pptx", strict=True)

    presentation = Presentation(str(out))
    assert len(presentation.slides) == 3
    for slide in (presentation.slides[0], presentation.slides[2]):
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
        assert {"FOUNDATION", "LIPSTICK", "SKINCARE"} <= set(texts)


def test_reference_layout_fixture_renders_charts_tables_and_groups(tmp_path):
    deck = Deck.model_validate(
        json.loads((FIXTURES / "reference_layouts.json").read_text(encoding="utf-8"))
    )
    out = build_deck(deck, tmp_path / "reference-layouts.pptx", strict=True)
    presentation = Presentation(str(out))

    assert len(presentation.slides) == 6
    allocation = presentation.slides[1]
    assert any(shape.has_chart for shape in allocation.shapes)
    assert any("85.30%" in shape.text for shape in allocation.shapes if shape.has_text_frame)
    market = presentation.slides[2]
    assert any(shape.has_table for shape in market.shapes)
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "How it works" in text
    assert "Mobile wallets" in text
    assert "Visa" in text
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left >= 0
            assert shape.top >= 0
            assert shape.left + shape.width <= presentation.slide_width
            assert shape.top + shape.height <= presentation.slide_height


def test_allocation_labels_do_not_overlap(tmp_path):
    deck = Deck.model_validate(
        {
            "title": "Allocation check",
            "slides": [
                {"layout": "title", "title": "Allocation check"},
                {
                    "layout": "ask_allocation",
                    "title": "Ask and Allocation",
                    "ask": "Rs. 14 Cr",
                    "items": [
                        {"label": "Contingency / Buffer", "pct": 85.3},
                        {"label": "Other indirect expenses", "pct": 6.94},
                        {"label": "Team", "pct": 3.02},
                        {"label": "Marketing", "pct": 2.57},
                        {"label": "Direct expenses", "pct": 1.95},
                        {"label": "Capex", "pct": 0.22},
                    ],
                },
                {"layout": "closing", "title": "Done"},
            ],
        }
    )
    out = build_deck(deck, tmp_path / "allocation.pptx", strict=True)
    slide = Presentation(str(out)).slides[1]
    label_shapes = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame
        and ("%" in shape.text or shape.text in {item.label for item in deck.slides[1].items})
    ]

    for index, shape in enumerate(label_shapes):
        for other in label_shapes[index + 1 :]:
            horizontal_overlap = (
                shape.left < other.left + other.width and other.left < shape.left + shape.width
            )
            vertical_overlap = (
                shape.top < other.top + other.height and other.top < shape.top + shape.height
            )
            assert not (horizontal_overlap and vertical_overlap), (shape.text, other.text)


def test_empty_quadrants_show_explicit_name_placeholders(tmp_path):
    deck = Deck.model_validate(
        {
            "title": "Landscape",
            "slides": [
                {"layout": "title", "title": "Landscape"},
                {
                    "layout": "quadrants",
                    "title": "Competitive Landscape",
                    "quadrants": [
                        {"heading": "Wallets"},
                        {"heading": "Cards"},
                        {"heading": "Networks"},
                        {"heading": "Other"},
                    ],
                },
                {"layout": "closing", "title": "Done"},
            ],
        }
    )
    out = build_deck(deck, tmp_path / "quadrants.pptx", strict=True)
    slide = Presentation(str(out)).slides[1]
    placeholders = [
        shape.text
        for shape in slide.shapes
        if shape.has_text_frame and "Add verified competitor names" in shape.text
    ]

    assert len(placeholders) == 4


def test_flow_reference_renders_eight_stages_and_labels_missing_steps(tmp_path):
    deck = Deck.model_validate(
        {
            "title": "Workflow",
            "slides": [
                {"layout": "title", "title": "Workflow"},
                {
                    "layout": "flow_steps",
                    "title": "Customer Journey",
                    "steps": [
                        {"title": "Sign up"},
                        {"title": "Verify KYC"},
                        {"title": "Get funded"},
                    ],
                },
                {"layout": "closing", "title": "Next"},
            ],
        }
    )
    out = build_deck(deck, tmp_path / "workflow.pptx", strict=True)
    slide = Presentation(str(out)).slides[1]
    texts = [shape.text for shape in slide.shapes if shape.has_text_frame]

    assert sum(text.isdigit() for text in texts) == 8
    assert texts.count("[Add verified step]") == 5
